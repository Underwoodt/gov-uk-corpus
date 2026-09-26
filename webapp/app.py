"""GOV.UK-styled Shortlist Builder — FastAPI + Jinja front end.

Renders the Categories list and the Create/Edit + Preview pages in the Defra /
GOV.UK Design System style, reusing the existing data layer (categories, facets,
shortlist). Backend is selected by env (SQLite locally, Postgres on the server).

Run locally:
    CORPUS_DB=data/pilot.db uvicorn webapp.app:app --reload --port 8600
On the server (Postgres):
    set -a; . ~/gov-uk-corpus.env; set +a
    uvicorn webapp.app:app --host 0.0.0.0 --port 8600
"""
from __future__ import annotations

import csv
import hashlib
import hmac
import html
import concurrent.futures
import gzip
import io
import json
import logging
import os
import re
import threading
import secrets
import time
import urllib.request
from typing import Dict, List, Optional
from urllib.parse import quote, urlencode

from fastapi import Depends, FastAPI, Form, HTTPException, Query, Request
from fastapi.responses import (HTMLResponse, JSONResponse, PlainTextResponse,
                               RedirectResponse, Response)
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.concurrency import run_in_threadpool

from govuk_corpus import accounts, ai_models, audit
from govuk_corpus import categories as cat
from govuk_corpus import category_counts, category_transfer, feedback, guardrails, sessions
from govuk_corpus import (audit_stats, category_interview, evaluate, evaluate_driver, extract,
                          gold, keyword_explain, llm, orgs, peak_schedule, pricing, prompts, readability,
                          reporting, roles, search_augment, settings, shortlist, stage_align,
                          sustainability, view_counts)
from govuk_corpus import orgs as orgs_mod   # stable module handle (some routes take an `orgs` param)
from govuk_corpus.backend import db

HERE = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.getenv("CORPUS_DB", "data/pilot.db")
PASSWORD = os.getenv("DASHBOARD_PASSWORD")
COOKIE = "sb_auth"

# ---- auth mode (accounts programme, phase 2) ----------------------------
# "shared" = the existing single DASHBOARD_PASSWORD gate (default; live behaviour
# is unchanged). "accounts" = per-user session login (built behind this flag).
AUTH_MODE = os.getenv("AUTH_MODE", "shared").strip().lower()

# ---- CSRF (double-submit token) -----------------------------------------
CSRF_COOKIE = "sb_csrf"


def _csrf_token(request: Request) -> str:
    """This browser's CSRF token: the existing cookie, or a fresh one stashed on
    request.state for the middleware to set on the response."""
    tok = request.cookies.get(CSRF_COOKIE)
    if tok:
        return tok
    tok = getattr(request.state, "_csrf_new", None)
    if not tok:
        tok = secrets.token_urlsafe(32)
        request.state._csrf_new = tok
    return tok


async def _csrf_guard(request: Request) -> None:
    """Reject state-changing requests without a matching CSRF token. Double-submit:
    the token must arrive as the X-CSRF-Token header (fetch) or a `csrf` form field
    and equal the sb_csrf cookie. Safe methods pass through."""
    if request.method in ("GET", "HEAD", "OPTIONS", "TRACE"):
        return
    cookie = request.cookies.get(CSRF_COOKIE)
    token = request.headers.get("x-csrf-token")
    if token is None and "form" in request.headers.get("content-type", "").lower():
        token = (await request.form()).get("csrf")
    if not cookie or not token or not hmac.compare_digest(str(cookie), str(token)):
        raise HTTPException(status_code=403, detail="CSRF check failed")


app = FastAPI(title="Shortlist Builder", dependencies=[Depends(_csrf_guard)])
app.mount("/static", StaticFiles(directory=os.path.join(HERE, "static")), name="static")
templates = Jinja2Templates(directory=os.path.join(HERE, "templates"))


# ---- security headers + CSRF cookie -------------------------------------
# Additive, safe headers on every response; HSTS is opt-in (ENABLE_HSTS=1, HTTPS
# only). CSP is intentionally NOT set yet (inline scripts/styles). The CSRF cookie
# is issued here when absent so forms/fetches always have a token to echo back.
@app.middleware("http")
async def _security_headers(request: Request, call_next):
    resp = await call_next(request)
    resp.headers.setdefault("X-Content-Type-Options", "nosniff")
    resp.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    resp.headers.setdefault("X-Frame-Options", "SAMEORIGIN")
    if os.getenv("ENABLE_HSTS") == "1":
        resp.headers.setdefault("Strict-Transport-Security", "max-age=31536000; includeSubDomains")
    if CSRF_COOKIE not in request.cookies:
        # Readable by JS (the fetch wrapper echoes it in a header) — not a secret on
        # its own; it only has to match the copy the browser also sends.
        resp.set_cookie(CSRF_COOKIE, _csrf_token(request), samesite="lax",
                        secure=_secure_cookies(), max_age=31536000, path="/")
    return resp


# ---- forced password change gate ----------------------------------------
# When an admin has "dirtied" a record (must_change_password), a signed-in user is confined to
# the set-password page (and logout) until they choose a new one. current_user caches on
# request.state, so the handler doesn't resolve the session twice.
_PW_GATE_ALLOW = {"/account/set-password", "/logout", "/login", "/forgot-password"}


@app.middleware("http")
async def _force_password_change(request: Request, call_next):
    if AUTH_MODE == "accounts":
        path = request.url.path
        if not (path.startswith("/static") or path.startswith("/reset-password")
                or path in _PW_GATE_ALLOW):
            try:
                u = current_user(request)
            except Exception:
                u = None
            if u and u.get("must_change_password"):
                return RedirectResponse(url=str(request.url_for("set_password_page")), status_code=303)
    return await call_next(request)


# ---- last-resort error page ---------------------------------------------
# Best practice: never surface a traceback, SQL, or stack detail to the user —
# it leaks internals and is a standard pentest finding. Instead we log the full
# traceback server-side under a short reference and show the user a plain page
# carrying only that reference, which they can quote when reporting the problem
# (an operator then greps the logs for `ref=<code>`). HTTPException (403/404/…)
# keeps Starlette's own handling — this only catches unhandled 500s.
_error_log = logging.getLogger("webapp.error")


@app.exception_handler(Exception)
async def _unhandled_error(request: Request, exc: Exception):
    ref = secrets.token_hex(4)
    _error_log.exception("unhandled error ref=%s method=%s path=%s",
                         ref, request.method, request.url.path)
    body = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Something went wrong</title>
<link rel="stylesheet" href="/static/govuk.css"></head>
<body><div class="container" style="max-width:640px;margin:60px auto;padding:0 20px;">
<h1>Sorry, something went wrong</h1>
<p>The page couldn't be displayed because of a problem on our side. This has been
logged and we can look into it.</p>
<p>If you report this, please quote the reference below so we can find the exact
error in the logs:</p>
<p><strong>Reference:</strong> <code style="font-size:1.1em;">{ref}</code></p>
<p style="margin-top:28px;"><a href="/">Return to the start page</a></p>
</div></body></html>"""
    return HTMLResponse(body, status_code=500)

_META_CACHE: Dict[str, str] = {}

# ---- count cache ---------------------------------------------------------
# The corpus changes ~daily but the funnel is viewed constantly, so cache each
# distinct count for COUNT_CACHE_TTL seconds (0 disables). Keyed by the filter
# set, not the connection, since the underlying data is the same DB.
_COUNT_CACHE: Dict[tuple, tuple] = {}
_COUNT_TTL = int(os.getenv("COUNT_CACHE_TTL", "300"))
_COUNT_LOCK = threading.Lock()

# Doc-type options for the Page Types picker are org-scoped and change ~daily; cache the
# result per (org set, children flag) so repeat edits load instantly. Same TTL as counts.
_DOCTYPE_CACHE: Dict[tuple, tuple] = {}
_DOCTYPE_LOCK = threading.Lock()


def _count_key(filters: dict) -> tuple:
    return (tuple(sorted(filters.get("organisations") or [])),
            tuple(sorted(filters.get("document_types") or [])),
            tuple(sorted(filters.get("keywords") or [])),
            filters.get("match", "all"))


def cached_count(conn, **filters) -> int:
    """shortlist.count with a short TTL cache keyed by the filter set."""
    if _COUNT_TTL <= 0:
        return shortlist.count(conn, **filters)
    key = _count_key(filters)
    now = time.time()
    with _COUNT_LOCK:
        hit = _COUNT_CACHE.get(key)
        if hit and hit[1] > now:
            return hit[0]
    n = shortlist.count(conn, **filters)
    with _COUNT_LOCK:
        _COUNT_CACHE[key] = (n, now + _COUNT_TTL)
    return n

# Prefilled into the Page Types box on a new category (edit/clear as needed).
DEFAULT_DOC_TYPES = "\n".join(category_interview.MAIN_DOCUMENT_TYPES)


# ---- helpers -------------------------------------------------------------
# Per-connection Postgres statement timeout (ms). A query that exceeds it errors
# (the funnel shows "error") rather than hanging. Configurable so it can be raised
# on a slow box without a code change; 0 disables the timeout entirely.
try:
    DB_STATEMENT_TIMEOUT_MS = int(os.getenv("DB_STATEMENT_TIMEOUT_MS", "15000"))
except ValueError:
    DB_STATEMENT_TIMEOUT_MS = 15000


def connect():
    if not _is_pg():
        return db.connect(DB_PATH)
    return db.connect(DB_PATH, statement_timeout_ms=DB_STATEMENT_TIMEOUT_MS or None)


def _is_pg() -> bool:
    return db.__name__.endswith("db_pg")


def corpus_total(conn) -> Optional[int]:
    """Total usable pages — computed once and cached (also drives the top banner)."""
    if "n" not in _META_CACHE:
        try:
            _META_CACHE["n"] = conn.execute(
                "SELECT COUNT(*) AS n FROM content WHERE is_redirect = 0 AND content_hash IS NOT NULL"
            ).fetchone()["n"]
        except Exception:
            _META_CACHE["n"] = None
    return _META_CACHE["n"]


def corpus_meta(conn) -> str:
    n = corpus_total(conn)
    return f"Newest snapshot · {n:,} pages" if n is not None else "Newest snapshot"


def ctx(conn, request: Request, **extra) -> dict:
    spent = _daily_spend(conn)
    budget = _budget(conn)
    pct = round(spent / budget * 100, 1) if budget > 0 else None
    user = current_user(request)
    # In accounts mode the effective role is the logged-in user's own role; in
    # shared mode it's the single global role set on Settings.
    role = user["role"] if (AUTH_MODE == "accounts" and user) else roles.get_role(conn)
    base = {"request": request, "corpus_meta": corpus_meta(conn), "active_nav": "categories",
            "budget_bar": {"spent": round(spent, 4), "budget": budget, "pct": pct},
            # Effective role + a gate helper for templates:
            #   {% if can_use('Administrator') %}…{% endif %}
            "role": role, "can_use": lambda required=None: roles.allows(role, required),
            # Accounts mode: the logged-in user (None in shared-password mode).
            "auth_mode": AUTH_MODE, "user": user,
            "csrf_token": _csrf_token(request)}
    base.update(extra)
    return base


# ---- AI Assistant (temporary prototype) ---------------------------------
# Provider is chosen in the UI (Settings page) and stored in app_settings.
# API keys still come from env; the toggle just picks which provider to use.
# Providers, credentials, budget and the model call itself live in govuk_corpus.llm (so the
# evaluation driver and the benchmark CLI run without FastAPI). The module-level names below
# are kept because routes reference them and the mocked tests patch them.
PROVIDERS = llm.PROVIDERS
DEFAULT_PROVIDER = llm.DEFAULT_PROVIDER
_provider_key = llm.provider_key
_bedrock_creds = llm.bedrock_creds


def _provider_configured(provider: str) -> bool:
    """Whether a provider has usable credentials — an API key, or AWS creds for Bedrock
    (routes through the patchable hooks above)."""
    if provider == "bedrock":
        return _bedrock_creds() is not None
    return _provider_key(provider) is not None


DEFAULT_DAILY_BUDGET = llm.DEFAULT_DAILY_BUDGET     # USD/day
DEFAULT_MAX_DOCS = 600          # documents per evaluation run
MAX_CONSEC_EVAL_ERRORS = 6      # consecutive AI errors before a run bails (provider likely down)


_budget = llm.budget


def _max_docs(conn) -> int:
    try:
        return int(settings.get_setting(conn, "ai_max_docs_per_run", str(DEFAULT_MAX_DOCS)))
    except (TypeError, ValueError):
        return DEFAULT_MAX_DOCS


_daily_spend = llm.daily_spend
_log_ai_usage = llm.log_usage


# Phase → settings key holding the chosen model id for that phase.
PHASE_MODEL_KEYS = {
    evaluate.PHASE_INCLUSION: "phase_model_inclusion",
    evaluate.PHASE_EXCLUSION: "phase_model_exclusion",
    evaluate.PHASE_ADJUDICATION: "phase_model_adjudication",
}

# Phase → settings key holding the execution mode (synchronous | batch) for that phase.
PHASE_MODE_KEYS = {
    evaluate.PHASE_INCLUSION: "phase_mode_inclusion",
    evaluate.PHASE_EXCLUSION: "phase_mode_exclusion",
    evaluate.PHASE_ADJUDICATION: "phase_mode_adjudication",
}


def _phase_mode(conn, phase: str) -> str:
    """The stored execution mode for a phase (defaults to synchronous)."""
    return evaluate.normalise_mode(settings.get_setting(conn, PHASE_MODE_KEYS.get(phase, ""), ""))


def _config_from_model(conn, m: dict) -> dict:
    """Build an AI config dict from an ai_models row."""
    provider = m["provider"]
    p = PROVIDERS.get(provider) or PROVIDERS[DEFAULT_PROVIDER]
    return {"provider": provider, "label": p["label"], "base_url": p["base_url"],
            "model": m["model_id"], "key": _provider_key(provider),
            "has_key": _provider_configured(provider),
            "price_in": m["input_per_m"], "price_out": m["output_per_m"],
            "grid": dict(m), "peak_bitmap": peak_schedule.get_bitmap(conn, provider)}


def _default_config(conn) -> dict:
    p = PROVIDERS[DEFAULT_PROVIDER]
    return {"provider": DEFAULT_PROVIDER, "label": p["label"], "base_url": p["base_url"],
            "model": p["model"], "key": _provider_key(DEFAULT_PROVIDER),
            "has_key": _provider_configured(DEFAULT_PROVIDER),
            "price_in": p["price_in"], "price_out": p["price_out"]}


def _ai_config(conn) -> dict:
    """Resolve the active AI model (from the ai_models table) + provider env."""
    ai_models.seed_defaults(conn)
    active = settings.get_setting(conn, "active_model_id", "")
    m = ai_models.get_model(conn, active) if active else None
    if not m:
        models = ai_models.list_models(conn)
        m = models[0] if models else None
    return _config_from_model(conn, m) if m else _default_config(conn)


def _ai_config_for_phase(conn, phase: str) -> dict:
    """Resolve the model configured for a phase (Settings → Model per phase), falling
    back to the active model when the phase has no specific model set."""
    ai_models.seed_defaults(conn)
    mid = settings.get_setting(conn, PHASE_MODEL_KEYS.get(phase, ""), "")
    m = ai_models.get_model(conn, mid) if mid else None
    return _config_from_model(conn, m) if m else _ai_config(conn)


def _ai_reply(config: dict, system: str, prompt: str, *, sampling: Optional[dict] = None) -> dict:
    """Single-turn convenience wrapper over _ai_chat (used by the page evaluator). `sampling` =
    {temperature, thinking, effort} or None for the provider defaults. Kept as a def (not an
    alias) so it routes through the patchable `_ai_chat` hook."""
    sm = sampling or {}
    return _ai_chat(config, system, [{"role": "user", "content": prompt}],
                    max_tokens=llm.eval_max_tokens(), temperature=sm.get("temperature"),
                    thinking=sm.get("thinking"), effort=sm.get("effort"))


_ai_chat = llm.chat


def _secure_cookies() -> bool:
    """Send the Secure flag once we're on confirmed HTTPS (same signal as HSTS)."""
    return os.getenv("ENABLE_HSTS") == "1" or os.getenv("SECURE_COOKIES") == "1"


def _client_ip(request: Request) -> Optional[str]:
    xff = request.headers.get("x-forwarded-for", "")
    if xff:
        return xff.split(",")[0].strip()[:64]
    return request.client.host if request.client else None


# Simple in-process sliding-window rate limiter (per key). Fine for a single
# uvicorn worker; a shared/DB-backed limiter is a later hardening step.
_RATE: Dict[str, list] = {}


def _rate_limit(key: str, max_hits: int, window_s: int) -> bool:
    """Record a hit for `key`; return False if it exceeds max_hits in the window."""
    now = time.time()
    hits = [t for t in _RATE.get(key, []) if t > now - window_s]
    if len(hits) >= max_hits:
        _RATE[key] = hits
        return False
    hits.append(now)
    _RATE[key] = hits
    return True


def _set_session_cookie(resp, cookie: str) -> None:
    resp.set_cookie(sessions.COOKIE_NAME, cookie, httponly=True, samesite="lax",
                    secure=_secure_cookies(), max_age=sessions.SESSION_TTL_HOURS * 3600, path="/")


def current_user(request: Request) -> Optional[dict]:
    """The logged-in user (accounts mode) or None. Resolves the session cookie once
    per request and caches it on request.state."""
    if AUTH_MODE != "accounts":
        return None
    cached = getattr(request.state, "_user", "unset")
    if cached != "unset":
        return cached
    user = None
    cookie = request.cookies.get(sessions.COOKIE_NAME, "")
    if cookie:
        conn = connect()
        try:
            user = sessions.resolve(conn, cookie)
        except Exception:
            user = None
        finally:
            conn.close()
    request.state._user = user
    return user


def authed(request: Request) -> bool:
    if AUTH_MODE == "accounts":
        return current_user(request) is not None
    if not PASSWORD:
        return True
    token = request.cookies.get(COOKIE, "")
    want = hmac.new(PASSWORD.encode(), b"ok", hashlib.sha256).hexdigest()
    return hmac.compare_digest(token, want)


def login_redirect(request: Request) -> RedirectResponse:
    return RedirectResponse(url=str(request.url_for("login")), status_code=303)


def form_values(form) -> dict:
    """Flatten a submitted form into a category data dict."""
    d = {k: (form.get(k) or "").strip() for k in
         ("owner_email", "keywords", "inclusion_context", "exclusion_context",
          "adjudication_hints_keep", "adjudication_hints_drop")}
    # Organisations and Page Types are multi-selects — collect every selected slug.
    d["dept_slugs"] = "\n".join(s.strip() for s in form.getlist("dept_slugs") if s.strip())
    d["document_type_slugs"] = "\n".join(s.strip() for s in form.getlist("document_type_slugs") if s.strip())
    d["include_child_orgs"] = form.get("include_child_orgs")  # checkbox: "on" or absent
    d["hybrid_on_save"] = "1"   # GOV.UK hybrid search now always runs on save (no longer opt-in)
    # Name is a free-text field (stored in `description`). The slug is derived from it and used
    # only for export filenames. Fall back to a submitted slug (the assistant sends one) if the
    # Name is empty, deriving a readable Name from it.
    name = (form.get("description") or "").strip()
    slug_in = (form.get("slug") or "").strip()
    if not name and slug_in:
        name = cat.prettify(slug_in)
    d["description"] = name
    d["slug"] = cat.slugify(name) or slug_in
    return d


# ---- auth ---------------------------------------------------------------
_LOGIN_HEAD = ("""<!doctype html><meta charset=utf-8>
    <link rel="icon" href="data:image/svg+xml,<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 16 16'><path d='M4 1h5l3 3v11H4z' fill='%231d70b8'/><path d='M9 1v3h3z' fill='%23003078'/><g fill='none' stroke='%23fff' stroke-width='1' stroke-linecap='round'><path d='M6 7h5'/><path d='M6 9h5'/><path d='M6 11h4'/></g></svg>">
    <link rel=stylesheet href='/static/govuk.css'>
    <div class='masthead'><div class='wrap'><img class='brand-logo' src='/static/notgovuk.svg' alt='NOT.GOV.UK — Defra supplier, operated with GOV.UK API-supplied data'><span class='brand'>Content Shortlist Builder</span></div></div>
    <script>
    // Give every password field a show/hide eye toggle (auth pages only).
    document.addEventListener('DOMContentLoaded', function () {
      var EYE = "<svg viewBox='0 0 24 24' fill='none' stroke='currentColor' stroke-width='2' stroke-linecap='round' stroke-linejoin='round'><path d='M1 12s4-7 11-7 11 7 11 7-4 7-11 7-11-7-11-7z'/><circle cx='12' cy='12' r='3'/></svg>";
      var OFF = "<svg viewBox='0 0 24 24' fill='none' stroke='currentColor' stroke-width='2' stroke-linecap='round' stroke-linejoin='round'><path d='M17.94 17.94A10.07 10.07 0 0 1 12 20C5 20 1 12 1 12a18.45 18.45 0 0 1 5.06-5.94M9.9 4.24A9.12 9.12 0 0 1 12 4c7 0 11 8 11 8a18.5 18.5 0 0 1-2.16 3.19'/><path d='M14.12 14.12a3 3 0 1 1-4.24-4.24'/><line x1='1' y1='1' x2='23' y2='23'/></svg>";
      document.querySelectorAll('input[type=password]').forEach(function (inp) {
        var wrap = document.createElement('span');
        wrap.className = 'pw-wrap';
        inp.parentNode.insertBefore(wrap, inp);
        wrap.appendChild(inp);
        var btn = document.createElement('button');
        btn.type = 'button'; btn.className = 'pw-toggle';
        btn.setAttribute('aria-label', 'Show password');
        btn.innerHTML = EYE;
        wrap.appendChild(btn);
        btn.addEventListener('click', function () {
          var show = inp.type === 'password';
          inp.type = show ? 'text' : 'password';
          btn.innerHTML = show ? OFF : EYE;
          btn.setAttribute('aria-label', show ? 'Hide password' : 'Show password');
          inp.focus();
        });
      });
    });
    </script>""")


@app.get("/login", response_class=HTMLResponse)
def login(request: Request, bad: int = 0, locked: int = 0, reset: int = 0):
    err = ""
    if reset:
        err = "<div class='success'>Your password has been changed. Sign in with your new password.</div>"
    if locked:
        err = f"<div class='error-summary'><h2>There is a problem</h2><ul><li>{accounts.LOGIN_LOCKED_MESSAGE}</li></ul></div>"
    elif bad:
        msg = accounts.LOGIN_FAILED_MESSAGE if AUTH_MODE == "accounts" else "Incorrect password."
        err = f"<div class='error-summary'><h2>There is a problem</h2><ul><li>{msg}</li></ul></div>"
    if AUTH_MODE == "accounts":
        fields = ("<div class='field'><label class='q' for='e'>Email address</label>"
                  "<input id='e' name='email' type='email' autocomplete='username'></div>"
                  "<div class='field'><label class='q' for='p'>Password</label>"
                  "<input id='p' name='password' type='password' autocomplete='current-password'></div>")
    else:
        fields = ("<div class='field'><label class='q' for='p'>Password</label>"
                  "<input id='p' name='password' type='password'></div>")
    reg = (f"<p style='margin-top:16px;'>"
           f"<a href='{request.url_for('forgot_password')}'>Forgotten your password?</a> · "
           f"<a href='{request.url_for('register')}'>Create an account</a></p>"
           if AUTH_MODE == "accounts" else "")
    return HTMLResponse(
        f"""{_LOGIN_HEAD}
        <div class='wrap body'><h1>Sign in</h1>{err}
        <form class='authform' method=post action='{request.url_for('do_login')}'>
        <input type=hidden name=csrf value='{_csrf_token(request)}'>{fields}
        <button class='btn' type=submit>Sign in</button></form>{reg}</div>
        <div class='page-id'>guc-0014</div>""")


@app.post("/login")
def do_login(request: Request, password: str = Form(""), email: str = Form("")):
    if AUTH_MODE == "accounts":
        conn = connect()
        try:
            user, reason = accounts.authenticate(conn, email, password, ip=_client_ip(request))
            if user is None:
                where = "?locked=1" if reason == "locked" else "?bad=1"
                return RedirectResponse(url=str(request.url_for("login")) + where, status_code=303)
            cookie = sessions.create_session(conn, user["id"], ip=_client_ip(request),
                                             user_agent=request.headers.get("user-agent"))
        finally:
            conn.close()
        # Admin "dirtied" the record (or a reset was required) → force a new password first.
        dest = ("set_password_page" if user.get("must_change_password") else "list_categories_page")
        resp = RedirectResponse(url=str(request.url_for(dest)), status_code=303)
        _set_session_cookie(resp, cookie)
        return resp
    # shared-password mode
    if PASSWORD and not hmac.compare_digest(password, PASSWORD):
        return RedirectResponse(url=str(request.url_for("login")) + "?bad=1", status_code=303)
    resp = RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    token = hmac.new((PASSWORD or "").encode(), b"ok", hashlib.sha256).hexdigest()
    resp.set_cookie(COOKIE, token, httponly=True, samesite="lax", max_age=86400)
    return resp


@app.post("/logout")
def do_logout(request: Request):
    resp = RedirectResponse(url=str(request.url_for("login")), status_code=303)
    if AUTH_MODE == "accounts":
        cookie = request.cookies.get(sessions.COOKIE_NAME, "")
        if cookie:
            conn = connect()
            try:
                user = current_user(request)
                sessions.revoke(conn, cookie)
                if user:
                    accounts.audit(conn, "logout", user_id=user["id"], ip=_client_ip(request))
            finally:
                conn.close()
        resp.delete_cookie(sessions.COOKIE_NAME, path="/")
    else:
        resp.delete_cookie(COOKIE)
    return resp


# ---- forgotten password / reset / forced change (accounts mode) ----------
# Iteration 1: no email yet — the reset link is shown on screen (self-serve) or handed over by
# an admin. A later iteration emails it. Tokens are single-use, hashed at rest, and expiring.
def _auth_page(request: Request, body: str, pageid: str) -> HTMLResponse:
    return HTMLResponse(f"{_LOGIN_HEAD}\n<div class='wrap body'>{body}</div>"
                        f"<div class='page-id'>{pageid}</div>")


@app.get("/forgot-password", response_class=HTMLResponse)
def forgot_password(request: Request, sent: int = 0):
    if AUTH_MODE != "accounts":
        return RedirectResponse(url=str(request.url_for("login")), status_code=303)
    action = request.url_for("forgot_password_submit")
    body = (f"<h1>Reset your password</h1>"
            f"<p>Enter your email address and we'll create a reset link.</p>"
            f"<form class='authform' method=post action='{action}'>"
            f"<input type=hidden name=csrf value='{_csrf_token(request)}'>"
            f"<div class='field'><label class='q' for='e'>Email address</label>"
            f"<input id='e' name='email' type='email' autocomplete='username'></div>"
            f"<button class='btn' type=submit>Create reset link</button></form>"
            f"<p style='margin-top:16px;'><a href='{request.url_for('login')}'>Back to sign in</a></p>")
    return _auth_page(request, body, "guc-0024")


@app.post("/forgot-password", response_class=HTMLResponse)
async def forgot_password_submit(request: Request):
    if AUTH_MODE != "accounts":
        return RedirectResponse(url=str(request.url_for("login")), status_code=303)
    form = await request.form()
    email = (form.get("email") or "").strip()
    link_html = ""
    conn = connect()
    try:
        u = accounts.get_user_by_email(conn, email) if email else None
        if u:
            token = accounts.create_reset_token(conn, u["id"])
            url = str(request.url_for("reset_password_page", token=token))
            accounts.audit(conn, "password_reset_requested", user_id=u["id"], ip=_client_ip(request))
            link_html = (f"<div class='success' style='margin-top:14px;'>Reset link created — open it to set a "
                         f"new password (valid {accounts.RESET_TTL_HOURS}h, single use):<br>"
                         f"<a href='{url}' style='word-break:break-all;'>{url}</a></div>")
    finally:
        conn.close()
    # Same generic wording whether or not the account exists (no email enumeration); the link only
    # appears when it matched. Iteration 1 shows it here; later this is emailed instead.
    body = (f"<h1>Reset your password</h1>"
            f"<p>If an account exists for <strong>{html.escape(email)}</strong>, a reset link has been created."
            f" (In this version the link is shown below rather than emailed.)</p>{link_html}"
            f"<p style='margin-top:16px;'><a href='{request.url_for('login')}'>Back to sign in</a></p>")
    return _auth_page(request, body, "guc-0024")


@app.get("/reset-password/{token}", response_class=HTMLResponse)
def reset_password_page(request: Request, token: str, bad: int = 0):
    if AUTH_MODE != "accounts":
        return RedirectResponse(url=str(request.url_for("login")), status_code=303)
    conn = connect()
    try:
        u = accounts.user_for_reset_token(conn, token)
    finally:
        conn.close()
    if not u:
        body = ("<h1>Reset link expired</h1><p>This reset link is invalid, already used, or expired. "
                f"<a href='{request.url_for('forgot_password')}'>Request a new one</a>.</p>")
        return _auth_page(request, body, "guc-0025")
    err = ("<div class='error-summary'><h2>There is a problem</h2><ul><li>"
           "The password didn't meet the policy or the two entries differed.</li></ul></div>" if bad else "")
    action = request.url_for("reset_password_submit", token=token)
    body = (f"<h1>Set a new password</h1><p>For <strong>{html.escape(u['email'])}</strong>.</p>{err}"
            f"<form class='authform' method=post action='{action}'>"
            f"<input type=hidden name=csrf value='{_csrf_token(request)}'>"
            f"<div class='field'><label class='q' for='p'>New password</label>"
            f"<input id='p' name='password' type='password' autocomplete='new-password'></div>"
            f"<div class='field'><label class='q' for='p2'>Confirm new password</label>"
            f"<input id='p2' name='password2' type='password' autocomplete='new-password'></div>"
            f"<button class='btn' type=submit>Set password</button></form>")
    return _auth_page(request, body, "guc-0025")


@app.post("/reset-password/{token}")
async def reset_password_submit(request: Request, token: str):
    if AUTH_MODE != "accounts":
        return RedirectResponse(url=str(request.url_for("login")), status_code=303)
    form = await request.form()
    pw, pw2 = (form.get("password") or ""), (form.get("password2") or "")
    back = str(request.url_for("reset_password_page", token=token))
    if pw != pw2:
        return RedirectResponse(url=back + "?bad=1", status_code=303)
    conn = connect()
    try:
        try:
            ok = accounts.reset_password_with_token(conn, token, pw)
        except ValueError:
            return RedirectResponse(url=back + "?bad=1", status_code=303)
        if not ok:
            return RedirectResponse(url=back, status_code=303)   # token went stale
    finally:
        conn.close()
    return RedirectResponse(url=str(request.url_for("login")) + "?reset=1", status_code=303)


@app.get("/account/set-password", response_class=HTMLResponse)
def set_password_page(request: Request, bad: int = 0):
    """Forced password change: the user is signed in but must_change_password is set."""
    if not authed(request):
        return login_redirect(request)
    u = current_user(request)
    if not u or not u.get("must_change_password"):
        return RedirectResponse(url=str(request.url_for("profile_page")), status_code=303)
    err = ("<div class='error-summary'><h2>There is a problem</h2><ul><li>"
           "The password didn't meet the policy or the two entries differed.</li></ul></div>" if bad else "")
    action = request.url_for("set_password_submit")
    body = (f"<h1>Set a new password</h1>"
            f"<p>An administrator has asked you to choose a new password before continuing.</p>{err}"
            f"<form class='authform' method=post action='{action}'>"
            f"<input type=hidden name=csrf value='{_csrf_token(request)}'>"
            f"<div class='field'><label class='q' for='p'>New password</label>"
            f"<input id='p' name='password' type='password' autocomplete='new-password'></div>"
            f"<div class='field'><label class='q' for='p2'>Confirm new password</label>"
            f"<input id='p2' name='password2' type='password' autocomplete='new-password'></div>"
            f"<button class='btn' type=submit>Set password</button></form>")
    return _auth_page(request, body, "guc-0026")


@app.post("/account/set-password")
async def set_password_submit(request: Request):
    if not authed(request):
        return login_redirect(request)
    u = current_user(request)
    if not u:
        return login_redirect(request)
    form = await request.form()
    pw, pw2 = (form.get("password") or ""), (form.get("password2") or "")
    back = str(request.url_for("set_password_page"))
    if pw != pw2:
        return RedirectResponse(url=back + "?bad=1", status_code=303)
    conn = connect()
    try:
        try:
            accounts.set_password(conn, u["id"], pw)
        except ValueError:
            return RedirectResponse(url=back + "?bad=1", status_code=303)
        accounts.audit(conn, "password_changed_forced", user_id=u["id"], ip=_client_ip(request))
    finally:
        conn.close()
    return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)


# ---- registration (accounts mode only) ----------------------------------
def _register_page(request: Request, *, error: str = "", values: Optional[dict] = None) -> HTMLResponse:
    v = values or {}
    def field(k):
        return html.escape((v.get(k) or "").strip(), quote=True)
    err = (f"<div class='error-summary'><h2>There is a problem</h2><ul><li>{html.escape(error)}</li></ul></div>"
           if error else "")
    suggestion = accounts.suggest_passphrase()
    rules_li = "".join(f"<li>{html.escape(r)}</li>" for r in accounts.PASSWORD_RULES)
    return HTMLResponse(
        f"""{_LOGIN_HEAD}
        <div class='wrap body'><h1>Create an account</h1>{err}
        <p class='secondary'>Use your DEFRA or Equal Experts email address.</p>
        <form class='authform' method=post action='{request.url_for('do_register')}'>
        <input type=hidden name=csrf value='{_csrf_token(request)}'>
        <div class='field'><label class='q' for='fn'>First name</label>
        <input id='fn' name='first_name' type='text' autocomplete='given-name' value='{field("first_name")}'></div>
        <div class='field'><label class='q' for='ln'>Last name</label>
        <input id='ln' name='last_name' type='text' autocomplete='family-name' value='{field("last_name")}'></div>
        <div class='field'><label class='q' for='e'>Email address</label>
        <input id='e' name='email' type='email' autocomplete='username' value='{field("email")}'></div>
        <div class='pw-suggest'>
        Suggested password: <code id='pwsg'>{suggestion}</code>
        <button type='button' class='btn secondary-btn' style='padding:4px 10px;font-size:14px;margin-left:8px;'
          onclick="var p=document.getElementById('pwsg').textContent;document.getElementById('p').value=p;document.getElementById('p2').value=p;">Use this</button>
        <div class='secondary' style='margin-top:6px;'>Three random words — easy to remember, hard to guess. Or choose your own.</div></div>
        <div class='field'><label class='q' for='p'>Password</label>
        <input id='p' name='password' type='password' autocomplete='new-password'></div>
        <div class='field'><label class='q' for='p2'>Confirm password</label>
        <input id='p2' name='confirm' type='password' autocomplete='new-password'></div>
        <details><summary>Password rules</summary>
        <div class='detail'><ul style='margin:0;padding-left:20px;'>{rules_li}</ul></div></details>
        <button class='btn' type=submit>Create account</button></form>
        <p style='margin-top:16px;'><a href='{request.url_for('login')}'>Already have an account? Sign in</a></p></div>
        <div class='page-id'>guc-0015</div>""")


@app.get("/register", response_class=HTMLResponse)
def register(request: Request):
    if AUTH_MODE != "accounts":
        return RedirectResponse(url=str(request.url_for("login")), status_code=303)
    return _register_page(request)


@app.post("/register")
def do_register(request: Request, first_name: str = Form(""), last_name: str = Form(""),
                email: str = Form(""), password: str = Form(""), confirm: str = Form("")):
    if AUTH_MODE != "accounts":
        return RedirectResponse(url=str(request.url_for("login")), status_code=303)
    vals = {"first_name": first_name, "last_name": last_name, "email": email}
    ip = _client_ip(request)
    if not _rate_limit(f"register:{ip}", 5, 3600):
        return _register_page(request, error="Too many attempts. Please try again later.", values=vals)
    if password != confirm:
        return _register_page(request, error="The passwords do not match.", values=vals)
    conn = connect()
    try:
        # role/account_status are NOT taken from the form (mass-assignment guard):
        # create_user defaults them to User / active.
        user = accounts.create_user(conn, email=email, first_name=first_name,
                                    last_name=last_name, password=password)
    except accounts.EmailTakenError:
        conn.close()
        return _register_page(request, error="An account with that email address already exists.", values=vals)
    except ValueError as e:
        conn.close()
        return _register_page(request, error=str(e), values=vals)
    accounts.audit(conn, "account_created", user_id=user["id"], ip=ip)
    cookie = sessions.create_session(conn, user["id"], ip=ip,
                                     user_agent=request.headers.get("user-agent"))
    conn.close()
    resp = RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    _set_session_cookie(resp, cookie)
    return resp


# ---- list ---------------------------------------------------------------
@app.get("/", response_class=HTMLResponse)
def list_categories_page(request: Request, flash: str = ""):
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    rows = cat.list_categories(conn)
    # Read the precomputed input-shortlist size per category (refreshed by the
    # nightly cycle's tidy-up and on every create/edit) in one query, instead of
    # running a live corpus COUNT per row on this hot page.
    counts = category_counts.get_counts(conn)
    for r in rows:
        r["display_name"] = cat.display_name(r)
        r["updated"] = (r.get("updated_at") or r.get("created_at") or "")[:10] or "—"
        hit = counts.get(int(r["id"]))
        r["pages_kept"] = "{:,}".format(hit["pages_kept"]) if hit and hit["pages_kept"] is not None else "—"
        r["pages_kept_at"] = (hit["computed_at"] or "")[:10] if hit else ""
    lsql, lparams = cat.list_categories_query()
    return templates.TemplateResponse("list.html", ctx(
        conn, request, categories=rows, flash=flash, list_sql=_display_sql(lsql, lparams)))


# ---- create -------------------------------------------------------------
@app.get("/categories/new", response_class=HTMLResponse)
def new_category_page(request: Request):
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    return templates.TemplateResponse("form.html", _form_ctx(conn, request, None, {}, []))


# ---- category assistant (guided interview -> pre-fill the create form) ---
@app.get("/categories/new/assistant", response_class=HTMLResponse)
def category_assistant_page(request: Request):
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    cats = cat.list_categories(conn)
    categories = [{"id": c["id"],
                   "name": cat.prettify(c.get("slug")) or (c.get("description") or "Untitled")}
                  for c in cats]
    resp = templates.TemplateResponse("category_assistant.html", ctx(
        conn, request, greeting=category_interview.GREETING, categories=categories,
        main_doc_types=list(category_interview.MAIN_DOCUMENT_TYPES)))
    conn.close()
    return resp


@app.get("/api/categories/{cid}/definition")
def api_category_definition(request: Request, cid: int):
    """The category's current field values, for the assistant to import and refine."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    try:
        c = cat.get_category(conn, cid)
        if not c:
            return JSONResponse({"error": "not found"}, status_code=404)
        name = cat.prettify(c.get("slug")) or (c.get("description") or "Untitled")
        fields = {k: c.get(k) for k in category_interview.FIELD_KEYS
                  if c.get(k) is not None and c.get(k) != ""}
        if "include_child_orgs" in c:
            fields["include_child_orgs"] = bool(c.get("include_child_orgs"))
        return JSONResponse({"name": name, "fields": fields})
    finally:
        conn.close()


def _slug_reference(conn, text: str) -> str:
    """A note of REAL organisation / document-type slugs matching words in the user's
    latest message, appended to the interview prompt so the assistant recommends valid
    slugs (widely) rather than inventing them — and re-checks whenever a field changes."""
    org_matches = orgs.search(conn, text, limit=20)
    dt_matches = category_interview.match_document_types(text)
    if not org_matches and not dt_matches:
        return ""
    lines = ["SLUG REFERENCE — real slugs from the corpus that match words in the user's "
             "latest message. When the user names organisations or document types that are "
             "NOT exact slugs, recommend from these (widely — offer any that plausibly "
             "match) and never invent a slug. If a word matches none here, say so and offer "
             "the closest options."]
    if org_matches:
        lines.append("Organisation slugs: " + ", ".join(m["slug"] for m in org_matches))
    if dt_matches:
        lines.append("Document-type slugs: " + ", ".join(dt_matches))
    return "\n".join(lines)


@app.post("/api/categories/assistant")
async def api_category_assistant(request: Request):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    body = await request.json()
    messages = body.get("messages") or []
    if not isinstance(messages, list) or not messages:
        return JSONResponse({"error": "no messages"}, status_code=400)
    edit_fields = body.get("edit_fields") if isinstance(body.get("edit_fields"), dict) else None
    # Keep only role/content and cap history length to bound cost.
    clean = [{"role": m.get("role"), "content": str(m.get("content") or "")}
             for m in messages[-24:] if m.get("role") in ("user", "assistant")]
    # Guardrail: block the user's latest message if it carries personal data or
    # prohibited language, before it reaches the model. Pause on the same question.
    last_user = next((m["content"] for m in reversed(clean) if m["role"] == "user"), "")
    finding = guardrails.check(last_user)
    if finding:
        return JSONResponse({"reply": guardrails.refusal_message(finding),
                             "fields": None, "suggestion": None})
    conn = connect()
    try:
        budget = _budget(conn)
        spent = _daily_spend(conn)
        if budget > 0 and spent >= budget:
            return JSONResponse({"error": f"Daily AI budget of ${budget:.2f} reached "
                                 f"(${spent:.4f} spent today)."}, status_code=429)
        cfg = _ai_config(conn)
        # Re-check the user's latest message against real slugs and hand the model the
        # matches, so a changed organisation/doc-type field gets valid-slug recommendations.
        system = category_interview.system_prompt(edit_fields, base=prompts.default_text("builder"))
        ref = _slug_reference(conn, last_user)
        if ref:
            system = system + "\n\n" + ref
        res = await run_in_threadpool(_ai_chat, cfg, system, clean)
        if res.get("error"):
            return JSONResponse({"error": res["error"]}, status_code=502)
        _log_ai_usage(conn, res.get("cost_usd"), res.get("input_tokens"),
                      res.get("output_tokens"), "assistant")
        reply = res.get("reply", "")
        fields = category_interview.parse_fields(reply)
        suggestion = category_interview.parse_suggestion(reply)
        return JSONResponse({"reply": reply, "fields": fields, "suggestion": suggestion})
    finally:
        conn.close()


def _slug_errors(conn, data: dict) -> list:
    """Server-side slug validation — the hard backstop behind the assistant's guidance.
    Reject organisation / document-type slugs that don't exist (plurals and typos won't),
    suggesting the closest real slugs. Returns human-readable errors ([] if all valid)."""
    errors: list = []
    known_dt = set(category_interview.DOCUMENT_TYPES)
    for s in cat.parse_list(data.get("document_type_slugs")):
        if s not in known_dt:
            near = category_interview.match_document_types(s)
            hint = f" Did you mean: {', '.join(near)}?" if near else ""
            errors.append(f"Unknown document type '{s}' — not a real slug.{hint}")
    for s in cat.parse_list(data.get("dept_slugs")):
        if orgs.get_org(conn, s) is None:
            near = [m["slug"] for m in orgs.search(conn, s, limit=5)]
            hint = f" Did you mean: {', '.join(near)}?" if near else ""
            errors.append(f"Unknown organisation '{s}' — not a real slug.{hint}")
    return errors


@app.post("/categories/new")
async def create_category(request: Request):
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    form = await request.form()
    data = form_values(form)
    # The owner is whoever is logged in — take their email (accounts mode), not the form.
    cu = current_user(request)
    if cu and cu.get("email"):
        data["owner_email"] = cu["email"]
    errors = cat.validate(data) + _slug_errors(conn, data)
    if errors:
        return templates.TemplateResponse("form.html", _form_ctx(conn, request, None, data, errors))
    cid = cat.create_category(conn, data)
    conn.close()
    # The shortlist rebuild + eval reconcile run on the interstitial (guc-0021), which shows
    # progress and then moves on to the Keyword Matching view (guc-0003a).
    return RedirectResponse(url=str(request.url_for("rebuild_category_page", cid=cid)), status_code=303)


@app.post("/categories/{cid}/copy")
async def copy_category_route(request: Request, cid: int):
    """Duplicate a category into a new draft and open it for editing."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    cu = current_user(request)
    owner = cu.get("email") if cu else None
    new_id = cat.copy_category(conn, cid, owner_email=owner)
    if new_id is None:
        conn.close()
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    category_counts.refresh_one(conn, new_id)  # give the copy its own stored count
    conn.close()
    return RedirectResponse(url=str(request.url_for("edit_category_page", cid=new_id)), status_code=303)


@app.post("/categories/{cid}/delete")
async def delete_category_route(request: Request, cid: int):
    """Permanently delete a shortlist and everything scoped to it, then return to the list."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    try:
        cat.delete_category(conn, cid)              # rows across the 6 category-scoped tables
        settings.set_setting(conn, f"funnel_cache_{cid}", "")   # + its per-shortlist settings keys
        settings.set_setting(conn, f"active_run_{cid}", "")
    finally:
        conn.close()
    return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)


# ---- edit ---------------------------------------------------------------
@app.get("/categories/{cid}/edit", response_class=HTMLResponse)
def edit_category_page(request: Request, cid: int):
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    return templates.TemplateResponse("form.html", _form_ctx(conn, request, category, category, []))


@app.post("/categories/{cid}/edit")
async def update_category(request: Request, cid: int):
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    form = await request.form()
    data = form_values(form)   # Name (and its derived slug) are editable now
    errors = cat.validate(data) + _slug_errors(conn, data)
    if errors:
        merged = {**category, **data}
        return templates.TemplateResponse("form.html", _form_ctx(conn, request, category, merged, errors))
    cat.update_category(conn, cid, data)
    conn.close()
    # The shortlist rebuild + eval reconcile run on the interstitial (guc-0021), which shows
    # progress and then moves on to the Keyword Matching view (guc-0003a).
    return RedirectResponse(url=str(request.url_for("rebuild_category_page", cid=cid)), status_code=303)


# ---- post-save rebuild interstitial (guc-0021) --------------------------
@app.get("/categories/{cid}/rebuilding", response_class=HTMLResponse)
def rebuild_category_page(request: Request, cid: int):
    """Progress page shown after a definition is saved: it drives the shortlist rebuild and
    the evaluation reconcile (as staged POSTs), then moves on to guc-0003a."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    category["display_name"] = cat.display_name(category)
    resp = templates.TemplateResponse("rebuilding.html", ctx(
        conn, request, category=category, hybrid_on_save=True))  # GOV.UK hybrid search always runs on save
    conn.close()
    return resp


@app.post("/api/categories/{cid}/refresh-shortlist")
def api_refresh_shortlist(request: Request, cid: int):
    """Recompute a category's stored count + materialised shortlist membership."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    if not cat.get_category(conn, cid):
        conn.close()
        return JSONResponse({"error": "not found"}, status_code=404)
    category_counts.refresh_one(conn, cid)
    n = reporting.membership_count(conn, cid)
    # Shortlist is now finished — collect GOV.UK Search view_count for its pages and stamp the
    # collection date onto `content`. Best-effort: a GOV.UK hiccup must not fail the rebuild.
    try:
        vc = view_counts.collect_for_category(conn, cid)
    except Exception as e:
        vc = {"error": f"{type(e).__name__}: {e}"}
    conn.close()
    return JSONResponse({"ok": True, "membership_count": n, "view_counts": vc})


@app.post("/api/categories/{cid}/reconcile-eval")
def api_reconcile_eval(request: Request, cid: int):
    """Align the category's LLM evaluation results with its refreshed shortlist."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    if not cat.get_category(conn, cid):
        conn.close()
        return JSONResponse({"error": "not found"}, status_code=404)
    evaluate.reconcile_to_shortlist(conn, cid)
    conn.close()
    return JSONResponse({"ok": True})


def _form_ctx(conn, request, category, values, errors) -> dict:
    org_options = orgs.all_orgs(conn)
    org_tree = orgs.hierarchy_forest(conn)          # nested forest, ordered by page count
    selected_orgs = cat.parse_list((values or {}).get("dept_slugs"))
    return ctx(
        conn, request,
        is_edit=category is not None,
        category=category,
        # When editing, title the page with the shortlist's own name; new shortlists get the generic label.
        form_title=(cat.display_name(category) if category else "Build a New Shortlist"),
        action=(str(request.url_for("update_category", cid=category["id"])) if category
                else str(request.url_for("create_category"))),
        v=values or {},
        org_options=org_options,
        org_option_slugs=[o["slug"] for o in org_options],
        org_tree=org_tree,
        org_counts_computed_at=orgs.counts_computed_at(conn),
        selected_orgs=selected_orgs,
        selected_doc_types=cat.parse_list((values or {}).get("document_type_slugs")),
        errors=errors,
    )


# ---- preview / run ------------------------------------------------------
def _filters(category) -> dict:
    return dict(
        organisations=cat.parse_list(category.get("dept_slugs")),
        document_types=cat.parse_list(category.get("document_type_slugs")),
        keywords=cat.parse_list(category.get("keywords")),
        match="any",
    )


def _effective_filters(conn, category) -> dict:
    """Filters actually applied: when 'include child organisations' is set, expand
    the chosen orgs to include their child departments (all descendants)."""
    f = _filters(category)
    if category.get("include_child_orgs") and f["organisations"]:
        f["organisations"] = orgs.expand_with_children(conn, f["organisations"], recursive=True)
    return f


_FUNNEL_STAGES = {  # stage -> (label, which filters apply)
    "all":     ("All pages", ()),
    "org":     ("After organisation filter", ("organisations",)),
    "doctype": ("After document-type filter", ("organisations", "document_types")),
    "keyword": ("After keyword filter", ("organisations", "document_types", "keywords")),
}


# ---- persistent, version-keyed funnel cache -----------------------------
# Funnel counts are deterministic given (definition, corpus), so they can be cached
# and re-served without a query until the definition or the corpus changes.
def _def_version(category: dict) -> str:
    """A hash of the funnel-relevant fields — changes only when the filters change
    (editing the AI context or owner doesn't invalidate)."""
    parts = "|".join([
        category.get("dept_slugs") or "", category.get("document_type_slugs") or "",
        category.get("keywords") or "", str(category.get("include_child_orgs") or "")])
    return hashlib.md5(parts.encode("utf-8")).hexdigest()[:12]


def _corpus_version(conn) -> str:
    """A cheap token that changes when the corpus is (re)crawled. The crawler stamps
    runs.finished_at on each import; org-hierarchy changes ride along with it."""
    try:
        row = conn.execute("SELECT COALESCE(MAX(finished_at), '') AS v FROM runs").fetchone()
        return str(row["v"] or "")
    except Exception:
        return ""


def _funnel_cache_read(conn, cid: int, def_v: str, corpus_v: str) -> dict:
    """The cached stage counts for this category, or {} if stale/absent."""
    raw = settings.get_setting(conn, f"funnel_cache_{cid}", "")
    if not raw:
        return {}
    try:
        blob = json.loads(raw)
    except (ValueError, TypeError):
        return {}
    if blob.get("def") == def_v and blob.get("corpus") == corpus_v:
        return blob.get("stages") or {}
    return {}


def _funnel_cache_write(conn, cid: int, def_v: str, corpus_v: str, stage: str, count) -> None:
    raw = settings.get_setting(conn, f"funnel_cache_{cid}", "")
    blob = {}
    try:
        blob = json.loads(raw) if raw else {}
    except (ValueError, TypeError):
        blob = {}
    if blob.get("def") != def_v or blob.get("corpus") != corpus_v:
        blob = {"def": def_v, "corpus": corpus_v, "stages": {}}   # versions moved on -> reset
    blob.setdefault("stages", {})[stage] = count
    settings.set_setting(conn, f"funnel_cache_{cid}", json.dumps(blob))


@app.get("/categories/{cid}", response_class=HTMLResponse)
def preview_category_page(request: Request, cid: int):
    """Fast shell — the funnel and results load progressively via the JSON API below."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    category["display_name"] = cat.display_name(category)
    filters = _effective_filters(conn, category)
    # The specific filter values applied at each funnel step (for the collapsed rows).
    stage_value_key = {"all": None, "org": "organisations",
                       "doctype": "document_types", "keyword": "keywords"}
    stages = [{"stage": k, "label": v[0],
               "vals": filters[stage_value_key[k]] if stage_value_key[k] else []}
              for k, v in _FUNNEL_STAGES.items()]
    # SQL preview is cheap (no DB hit) — render the per-stage COUNT queries the funnel
    # runs, plus the shortlist SELECT, with comment separators.
    def _pretty(sql_, params_):
        return shortlist.pretty_sql(shortlist.interpolate_sql(sql_, params_))

    stage_specs = [
        ("All pages", {}),
        ("After organisation filter", {"organisations": filters["organisations"]}),
        ("After document-type filter",
         {"organisations": filters["organisations"], "document_types": filters["document_types"]}),
        ("After keyword filter",
         {"organisations": filters["organisations"], "document_types": filters["document_types"],
          "keywords": filters["keywords"], "match": "any"}),
    ]
    blocks = ["-- ===== Funnel counts: one COUNT(*) per stage (the numbers in the table above) ====="]
    for label, kw in stage_specs:
        cs, cp = shortlist.build_query(count_only=True, **kw)
        blocks.append(f"-- {label}\n{_pretty(cs, cp)};")
    sel_sql, sel_params = shortlist.build_query(include_title=True, limit=10000, **filters)
    blocks.append("-- ===== Shortlist rows: the pages returned (results table / download) =====\n"
                  + f"{_pretty(sel_sql, sel_params)};")
    pretty = "\n\n".join(blocks)
    eval_max_docs = _max_docs(conn)
    # Total AI spend accrued to this shortlist, split by phase for the two AI funnel rows
    # (deterministic stages have no AI cost). Sums every run, not just the active one.
    ai_incl = ai_excl = 0.0
    for r in conn.execute(
            f"SELECT phase, COALESCE(SUM(cost), 0) AS c FROM evaluation_runs "
            f"WHERE category_id = {shortlist._P} GROUP BY phase", (cid,)).fetchall():
        d = dict(r)
        c = float(d.get("c") or 0)
        if "exclusion" in (d.get("phase") or "").lower():
            ai_excl += c
        else:
            ai_incl += c
    ai_cost = {"incl": ai_incl, "excl": ai_excl, "total": ai_incl + ai_excl}
    conn.close()
    return templates.TemplateResponse("preview.html", ctx(
        connect(), request, category=category, sql=pretty, stages=stages,
        eval_max_docs=eval_max_docs, ai_cost=ai_cost, filters=filters))


def _incl_run_options(conn, cid: int) -> list:
    """The category's inclusion runs (chain heads) for the Run selector dropdowns, newest first:
    [{run_id, label}]. Label is the run name + its date."""
    out = []
    for r in conn.execute(
            f"SELECT run_id, name, started_at FROM evaluation_runs "
            f"WHERE category_id = {shortlist._P} AND source_run_id IS NULL ORDER BY started_at DESC",
            (cid,)).fetchall():
        d = dict(r)
        dt = (d.get("started_at") or "")[:10]
        label = (d.get("name") or d["run_id"]) + (f" · {dt}" if dt else "")
        out.append({"run_id": d["run_id"], "label": label})
    return out


@app.get("/categories/{cid}/shortlist", response_class=HTMLResponse)
def shortlist_page(request: Request, cid: int, stage: str = "final", run: str = ""):
    """Audit Results page — the shortlist pages list. `stage` pre-selects the shortlist
    stage dropdown, `run` the Run dropdown. Data loads from /api/categories/{cid}/audit-shortlist.
    (GDS Compliance is now its own page — see gds_compliance_page.)"""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    category["display_name"] = cat.display_name(category)
    if stage not in _AUDIT_STAGE_KEYS:
        stage = "final"
    return templates.TemplateResponse("audit_shortlist.html", ctx(
        conn, request, category=category, stage=stage, run=run,
        runs=_incl_run_options(conn, cid),
        audit_stages=_AUDIT_STAGES, audit_sections=_DOWNLOAD_SECTIONS))


def _funnel_table_ctx(conn, category, cid: int):
    """(stages, ai_cost) for the Selection-funnel table: the per-stage labels + the filter
    values applied at each step, and total AI spend split by phase for the two AI rows."""
    filters = _effective_filters(conn, category)
    stage_value_key = {"all": None, "org": "organisations",
                       "doctype": "document_types", "keyword": "keywords"}
    stages = [{"stage": k, "label": v[0],
               "vals": filters[stage_value_key[k]] if stage_value_key[k] else []}
              for k, v in _FUNNEL_STAGES.items()]
    ai_incl = ai_excl = 0.0
    for r in conn.execute(
            f"SELECT phase, COALESCE(SUM(cost), 0) AS c FROM evaluation_runs "
            f"WHERE category_id = {shortlist._P} GROUP BY phase", (cid,)).fetchall():
        d = dict(r)
        c = float(d.get("c") or 0)
        if "exclusion" in (d.get("phase") or "").lower():
            ai_excl += c
        else:
            ai_incl += c
    return stages, {"incl": ai_incl, "excl": ai_excl, "total": ai_incl + ai_excl}


@app.get("/categories/{cid}/ai-pipeline", response_class=HTMLResponse)
def ai_pipeline_page(request: Request, cid: int):
    """AI evaluation pipeline for a shortlist (guc-0004c) — its own top-level page
    (Active Run / Run History), with the Selection funnel table above. Behaviour lives
    in /static/ai_pipeline.js plus this page's inline funnel script."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    category["display_name"] = cat.display_name(category)
    stages, ai_cost = _funnel_table_ctx(conn, category, cid)
    resp = templates.TemplateResponse("ai_pipeline_page.html", ctx(
        conn, request, category=category, eval_max_docs=_max_docs(conn),
        stages=stages, ai_cost=ai_cost))
    conn.close()
    return resp


@app.get("/categories/{cid}/govuk-additions", response_class=HTMLResponse)
def govuk_additions_page(request: Request, cid: int, only: str = ""):
    """Review the GOV.UK-Search-only pages the hybrid step adds to this shortlist's AI input
    (guc-0027): pages the keyword filter missed but GOV.UK Search returned — their URL,
    document type, publishing organisations, and the GOV.UK keyword(s) that caught them.
    `only=forwarded` narrows the list to just the pages actually forwarded to the AI
    (in the corpus and at/above the relevance floor) — the 'GOV.UK only' count on the funnel."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    category["display_name"] = cat.display_name(category)
    P = shortlist._P
    rows = [dict(r) for r in conn.execute(
        f"SELECT sp.url AS url, sp.title AS title, sp.document_type AS document_type, "
        f"sp.phrases AS phrases, sp.es_score AS es_score, "
        f"(CASE WHEN c.content_hash IS NOT NULL AND c.is_redirect = 0 THEN 1 ELSE 0 END) AS in_corpus "
        f"FROM category_search_pages sp LEFT JOIN content c ON c.url = sp.url "
        f"WHERE sp.category_id = {P} AND sp.source = 'search' "
        f"ORDER BY (CASE WHEN c.content_hash IS NOT NULL AND c.is_redirect = 0 THEN 0 ELSE 1 END), sp.url",
        (cid,)).fetchall()]
    # Publishing organisations per page (title, falling back to slug), deduped, order-preserving.
    urls = [r["url"] for r in rows if r.get("url")]
    org_map = {}
    for i in range(0, len(urls), 500):
        ch = urls[i:i + 500]
        ph = ",".join([P] * len(ch))
        for orow in conn.execute(
                f"SELECT po.page_url AS url, COALESCE(o.title, po.organisation_slug) AS org "
                f"FROM page_organisations po LEFT JOIN organisations o ON o.slug = po.organisation_slug "
                f"WHERE po.page_url IN ({ph})", tuple(ch)).fetchall():
            d = dict(orow)
            org_map.setdefault(d["url"], []).append(d["org"])
    floor = _GOVUK_MIN_ES_SCORE
    for r in rows:
        r["organisations"] = ", ".join(dict.fromkeys(org_map.get(r["url"], [])))
        r["in_corpus"] = bool(r.get("in_corpus"))
        es = r.get("es_score")
        # A page is fed to the AI only if it's in the corpus and above the relevance floor
        # (a page with no es_score is kept — unknown, not low).
        r["evaluated"] = r["in_corpus"] and (es is None or not floor or es >= floor)
    forwarded_only = (only == "forwarded")
    if forwarded_only:                       # the 'GOV.UK only' funnel row links here
        rows = [r for r in rows if r["evaluated"]]
    resp = templates.TemplateResponse("govuk_additions.html", ctx(
        conn, request, category=category, rows=rows, govuk_min_es_score=floor,
        forwarded_only=forwarded_only))
    conn.close()
    return resp


@app.get("/categories/{cid}/gds-compliance", response_class=HTMLResponse)
def gds_compliance_page(request: Request, cid: int):
    """GDS Compliance (Non-LLM) quality/freshness dashboard for a shortlist (guc-0004b) —
    promoted from a Shortlist sub-tab to its own page. Data via /api/.../audit-stats."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    category["display_name"] = cat.display_name(category)
    stages = [(s, _FUNNEL_STAGES[s][0]) for s in ("all", "org", "doctype", "keyword")]
    gds_check_meta = [{"name": c.name, "weight": c.weight, "reason": c.reason}
                      for c in readability.CHECKS]
    resp = templates.TemplateResponse("gds_compliance_page.html", ctx(
        conn, request, category=category, stages=stages, gds_check_meta=gds_check_meta))
    conn.close()
    return resp


@app.get("/api/categories/{cid}/funnel")
def api_funnel(request: Request, cid: int, stage: str = "all"):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    if stage not in _FUNNEL_STAGES:
        return JSONResponse({"error": "unknown stage"}, status_code=404)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        return JSONResponse({"error": "not found"}, status_code=404)
    label, applies = _FUNNEL_STAGES[stage]
    # Serve from the persistent cache unless the definition or corpus has changed. The
    # keyword stage is the exception: it is cheap (~40ms) and is the count people reconcile
    # against (Both + Shortlister only), so it is ALWAYS recomputed live and never read from
    # or written to the persistent cache — that cache only invalidates on a full crawl, so an
    # incremental corpus change (e.g. a GOV.UK fetch) could otherwise leave it stale.
    live_only = (stage == "keyword")
    def_v, corpus_v = _def_version(category), _corpus_version(conn)
    cached = _funnel_cache_read(conn, cid, def_v, corpus_v)
    if not live_only and stage in cached:
        conn.close()
        return JSONResponse({"stage": stage, "label": label, "count": cached[stage], "cached": True})
    if stage == "all":
        # Same number as the top banner — reuse the cached corpus total, no query.
        n = corpus_total(conn)
        if n is None:
            n = cached_count(conn)
    else:
        filters = _effective_filters(conn, category)
        kw = {k: filters[k] for k in applies}
        if "keywords" in kw:
            kw["match"] = "any"
        # keyword: hit the DB directly (~40ms) so it bypasses the 5-min in-memory TTL cache
        # too and is always exactly current; other stages use the short-TTL count cache.
        n = shortlist.count(conn, **kw) if live_only else cached_count(conn, **kw)
    if not live_only:
        _funnel_cache_write(conn, cid, def_v, corpus_v, stage, n)
    conn.close()
    return JSONResponse({"stage": stage, "label": label, "count": n, "cached": False})


@app.post("/api/categories/{cid}/funnel/refresh")
def api_funnel_refresh(request: Request, cid: int):
    """Clear this category's cached funnel counts (and the in-memory count/total caches)
    so the next funnel load re-queries from scratch."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    if not cat.get_category(conn, cid):
        conn.close()
        return JSONResponse({"error": "not found"}, status_code=404)
    settings.set_setting(conn, f"funnel_cache_{cid}", "")   # persistent funnel cache
    conn.close()
    with _COUNT_LOCK:
        _COUNT_CACHE.clear()   # in-memory TTL counts (funnel + org/keyword breakdowns)
    _META_CACHE.clear()        # cached corpus total ("All pages")
    return JSONResponse({"ok": True})


@app.get("/api/categories/{cid}/audit-stats")
def api_audit_stats(request: Request, cid: int, stage: str = "keyword"):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    if stage not in _FUNNEL_STAGES:
        return JSONResponse({"error": "unknown stage"}, status_code=404)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return JSONResponse({"error": "not found"}, status_code=404)
    filters = _effective_filters(conn, category)
    applies = _FUNNEL_STAGES[stage][1]     # which filters this stage applies
    kw = {k: filters[k] for k in applies}
    kw.setdefault("match", "any")
    try:
        out = audit_stats.stats(conn, **kw)
    except Exception as e:
        conn.close()
        return JSONResponse({"error": f"{type(e).__name__}: {e}"}, status_code=500)
    conn.close()
    return JSONResponse(out)


def _display_sql(sql: str, params) -> str:
    """Pretty, parameter-inlined SQL for the advanced-only 'Show SQL' links under lists.
    Display only — the executed query still uses safe parameter binding."""
    return shortlist.pretty_sql(shortlist.interpolate_sql(sql, list(params or [])))


# ---- reporting (read-only, over the persisted shortlist membership) ------
@app.get("/api/reporting/categories")
def api_reporting_overview(request: Request):
    """Every category with its stored input-count and materialised shortlist size."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    try:
        return JSONResponse({"categories": reporting.overview(conn)})
    finally:
        conn.close()


@app.get("/api/reporting/categories/{cid}/pages")
def api_reporting_pages(request: Request, cid: int, limit: int = 100, offset: int = 0):
    """A category's materialised shortlist pages (joined to content), paginated —
    no filter re-execution."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    limit = max(1, min(limit, 500))
    offset = max(0, offset)
    conn = connect()
    try:
        return JSONResponse(reporting.category_pages(conn, cid, limit=limit, offset=offset))
    finally:
        conn.close()


@app.get("/api/reporting/categories/{cid}/summary")
def api_reporting_summary(request: Request, cid: int):
    """Aggregates over a category's materialised shortlist — for dashboard charts."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    try:
        return JSONResponse(reporting.category_summary(conn, cid))
    finally:
        conn.close()


# ---- GOV.UK Search coverage (augment the shortlist with provenance) ------
@app.post("/api/categories/{cid}/govuk-compare")
async def api_govuk_compare(request: Request, cid: int):
    """Run an org-scoped GOV.UK Search of the category's keywords, compare with its stored
    shortlist, and persist the tagged union (Shortlister / Both / GOV.UK Search only)."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return JSONResponse({"error": "not found"}, status_code=404)
    filters = _effective_filters(conn, category)   # orgs (expanded) + keywords

    def work():
        return search_augment.compare(
            conn, cid, filters["keywords"], filters["organisations"],
            lambda phrases, org_slugs, doctypes, progress=None: _govuk_search_multi(phrases, org_slugs, doctypes),
            document_types=filters["document_types"])
    try:
        out = await run_in_threadpool(work)
    except Exception as e:
        conn.close()
        logging.getLogger("govuk_compare").warning("compare failed for %s: %s", cid, e)
        return JSONResponse({"error": "Couldn't reach GOV.UK Search — try again."}, status_code=502)
    conn.close()
    return JSONResponse(out)


@app.get("/api/categories/{cid}/augmented-pages")
def api_augmented_pages(request: Request, cid: int, source: str = "", limit: int = 100,
                        offset: int = 0, q: str = "", loaded: str = "", min_score: float = 0.0):
    """The stored GOV.UK-coverage augmented shortlist, tagged by source, paginated; `q`
    filters by title across the whole set."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    limit = max(1, min(limit, 500))
    offset = max(0, offset)
    conn = connect()
    try:
        out = search_augment.augmented_pages(conn, cid, source=source, limit=limit, offset=offset,
                                             q=q, loaded=loaded, min_score=min_score)
        summary = search_augment.summary(conn, cid)   # query_urls, thresholds, computed_at, global total
        # The top-of-page sums track the narrowing thresholds (min relevance + title search),
        # recomputed each load; total stays global so the "not run yet" check still works.
        fs = search_augment.filtered_summary(conn, cid, q=q, min_score=min_score)
        summary["counts"] = fs["counts"]
        summary["evaluable"] = fs["evaluable"]
        summary["pending_fetch"] = fs["pending_fetch"]
        out["summary"] = summary
        return JSONResponse(out)
    finally:
        conn.close()


@app.get("/api/categories/{cid}/ai-input-breakdown")
def api_ai_input_breakdown(request: Request, cid: int):
    """Model-1 breakdown of what a fresh run forwards to the AI (guc-0004c):
      A = Both              — corpus keyword match that GOV.UK Search also returned
      B = Shortlister only  — corpus keyword match GOV.UK did not return
      C = GOV.UK only       — in the corpus and at/above the relevance floor
    A and B are corpus keyword matches (always in corpus, floor NOT applied); C is the
    GOV.UK-only top-up gated by the floor. total = A + B + C = the pages a fresh run evaluates."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    floor = _GOVUK_MIN_ES_SCORE
    conn = connect()
    try:
        P = shortlist._P
        J = "category_search_pages sp LEFT JOIN content c ON c.url = sp.url"
        incorp = "c.is_redirect = 0 AND c.content_hash IS NOT NULL"

        def cnt(where, args):
            return dict(conn.execute(
                f"SELECT COUNT(*) AS n FROM {J} WHERE {where}", tuple(args)).fetchone())["n"]

        both = cnt(f"sp.category_id = {P} AND sp.source = {P} AND {incorp}", (cid, "both"))
        shortlister = cnt(f"sp.category_id = {P} AND sp.source = {P} AND {incorp}", (cid, "shortlister"))
        if floor and floor > 0:
            govuk = cnt(f"sp.category_id = {P} AND sp.source = {P} AND {incorp} "
                        f"AND (sp.es_score IS NULL OR sp.es_score >= {P})", (cid, "search", floor))
        else:
            govuk = cnt(f"sp.category_id = {P} AND sp.source = {P} AND {incorp}", (cid, "search"))
        return JSONResponse({"both": both, "shortlister": shortlister, "govuk_only": govuk,
                             "total": both + shortlister + govuk, "floor": floor})
    finally:
        conn.close()


_GOVUK_FETCH_CAP = 300   # bound one fetch batch; re-run to continue if more remain


@app.post("/api/categories/{cid}/govuk-fetch")
async def api_govuk_fetch(request: Request, cid: int):
    """Fetch this category's GOV.UK-Search-only pages that aren't in the corpus yet, so they
    gain content and become eligible for LLM evaluation. Bounded per call; re-run for more."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    if not cat.get_category(conn, cid):
        conn.close()
        return JSONResponse({"error": "not found"}, status_code=404)
    urls = search_augment.pending_fetch_urls(conn, cid)[:_GOVUK_FETCH_CAP]

    def work():
        counters = {}
        if urls:
            run_id = db.start_run(conn, stage="govuk-augment", scope=str(cid))
            counters = stage_align.align_urls(conn, run_id, urls, source="govuk-search", stage="govuk-augment")
            db.finish_run(conn, run_id, counters)
        updated = search_augment.refresh_content_ids(conn, cid)
        # Drop container rows whose html_publication attachment we already hold (re-tag it).
        search_augment.resolve_attachment_containers(conn, cid)
        return counters, updated
    try:
        counters, updated = await run_in_threadpool(work)
    except Exception as e:
        conn.close()
        logging.getLogger("govuk_fetch").warning("fetch failed for %s: %s", cid, e)
        return JSONResponse({"error": "Couldn't fetch the pages — try again."}, status_code=502)
    evaluable = search_augment.evaluable_search_only(conn, cid)
    pending = len(search_augment.pending_fetch_urls(conn, cid))
    conn.close()
    fetched = (counters.get("new", 0) + counters.get("changed", 0) + counters.get("unchanged", 0))
    return JSONResponse({"fetched": fetched, "updated": updated, "pending": pending, "evaluable": evaluable})


@app.get("/categories/{cid}/augmented/download")
def download_augmented(request: Request, cid: int):
    """CSV of the augmented shortlist with the provenance source column."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    rows = search_augment.export_rows(conn, cid)
    conn.close()
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["source", "url", "content_id", "title", "document_type",
                "corpus_keywords", "govuk_keywords", "es_score"])
    for r in rows:
        w.writerow([r.get("source"), r.get("url"), r.get("content_id"),
                    r.get("title"), r.get("document_type"),
                    r.get("corpus_keywords"), r.get("govuk_keywords"), r.get("es_score")])
    return Response(buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="govuk-coverage-{cid}-{_dl_stamp()}.csv"'})


@app.get("/api/categories/{cid}/org-breakdown")
def api_org_breakdown(request: Request, cid: int):
    """Pages contributed by EACH organisation, within the document-type + keyword filters
    — the 'which organisations matched' breakdown. Organisations overlap (a page can have
    several), so they don't sum to the shortlist total."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return JSONResponse({"error": "not found"}, status_code=404)
    filters = _effective_filters(conn, category)
    try:
        # Read from the materialised shortlist (category_shortlist_pages) — instant, and
        # it can't time out the way the old whole-corpus scan could.
        orgs_out = reporting.org_breakdown(conn, cid, filters["organisations"])
        mcount = reporting.membership_count(conn, cid)
        computed_at = reporting.membership_computed_at(conn, cid)
    except Exception as e:
        conn.close()
        return JSONResponse({"error": f"{type(e).__name__}: {e}"}, status_code=500)
    conn.close()
    osql, oparams = reporting.org_breakdown_query(cid, filters["organisations"])
    sql = _display_sql(osql, oparams) if osql else "-- No organisations set for this category."
    # 'pending' = the category has organisations but its shortlist hasn't been materialised
    # yet (no refresh has run); the UI shows a "not computed yet" note rather than "none".
    pending = bool(filters["organisations"]) and mcount == 0
    return JSONResponse({"orgs": orgs_out, "sql": sql, "pending": pending, "computed_at": computed_at})


@app.get("/api/categories/{cid}/doctype-breakdown")
def api_doctype_breakdown(request: Request, cid: int):
    """Pages in the materialised shortlist by effective document type — the 'which
    document types matched' breakdown. Reads category_shortlist_pages (instant)."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return JSONResponse({"error": "not found"}, status_code=404)
    filters = _effective_filters(conn, category)
    try:
        types_out = reporting.doctype_breakdown(conn, cid)
        mcount = reporting.membership_count(conn, cid)
        computed_at = reporting.membership_computed_at(conn, cid)
    except Exception as e:
        conn.close()
        return JSONResponse({"error": f"{type(e).__name__}: {e}"}, status_code=500)
    conn.close()
    dsql, dparams = reporting.doctype_breakdown_query(cid)
    sql = _display_sql(dsql, dparams)
    has_filter = bool(filters["organisations"] or filters["document_types"] or filters["keywords"])
    pending = has_filter and mcount == 0
    return JSONResponse({"types": types_out, "sql": sql, "pending": pending, "computed_at": computed_at})


@app.get("/api/categories/{cid}/keyword-breakdown")
def api_keyword_breakdown(request: Request, cid: int):
    """Pages matching EACH keyword individually, within the org + document-type set —
    the 'which terms matched' breakdown. Counts are independent, so they overlap and
    do not sum to the keyword-stage total (a page can match several terms)."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return JSONResponse({"error": "not found"}, status_code=404)
    filters = _effective_filters(conn, category)
    # Count each keyword within the materialised shortlist (a few thousand rows), not the
    # whole corpus — instant, and can't time out on common terms (import/export/transit).
    terms = reporting.keyword_breakdown(conn, cid, filters["keywords"])
    mcount = reporting.membership_count(conn, cid)
    computed_at = reporting.membership_computed_at(conn, cid)
    if reporting.has_matched_keywords(conn, cid):
        # Counts come from the stored per-page hits (same source as the row lozenges).
        sql = reporting.matched_keywords_sql(cid) if filters["keywords"] else "-- No keywords set for this category."
    else:
        blocks = []
        for kw in filters["keywords"]:
            ksql, kp = reporting.keyword_count_query(cid, kw)
            blocks.append(f"-- Count for keyword: {kw}\n{_display_sql(ksql, kp)};")
        sql = "\n\n".join(blocks) if blocks else "-- No keywords set for this category."
    conn.close()
    pending = bool(filters["keywords"]) and mcount == 0
    return JSONResponse({"terms": terms, "sql": sql, "pending": pending, "computed_at": computed_at})


@app.get("/api/categories/{cid}/keyword-overlap")
def api_keyword_overlap(request: Request, cid: int, n: int = 5):
    """Overlap of the top-N keywords within the materialised shortlist — region counts for
    a Venn diagram. `n` (2-5) selects how many of the top keywords to include."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    n = max(2, min(int(n or 5), 5))
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return JSONResponse({"error": "not found"}, status_code=404)
    filters = _effective_filters(conn, category)
    mcount = reporting.membership_count(conn, cid)
    # Top-N keywords by their in-shortlist count (keyword_breakdown is sorted desc).
    top = [t["keyword"] for t in reporting.keyword_breakdown(conn, cid, filters["keywords"])
           if t["count"]][:n]
    try:
        data = reporting.keyword_overlap(conn, cid, top)
    except Exception as e:
        conn.close()
        return JSONResponse({"error": f"{type(e).__name__}: {e}"}, status_code=500)
    computed_at = reporting.membership_computed_at(conn, cid)
    conn.close()
    raw = data.pop("sql", None)
    data["sql"] = _display_sql(raw[0], raw[1]) if isinstance(raw, tuple) else (raw or "")
    data["available_keywords"] = len([t for t in filters["keywords"]])
    data["pending"] = bool(filters["keywords"]) and mcount == 0
    data["computed_at"] = computed_at
    return JSONResponse(data)


@app.post("/api/keyword-explain")
async def api_keyword_explain(request: Request):
    """Plain-English preview of what the (unsaved) keyword text will match after tokenisation
    — for the keyword box on the form. Body: {"keywords": "<textarea contents>"}."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    body = await request.json()
    terms = cat.parse_list(str(body.get("keywords") or ""))
    conn = connect()
    try:
        return JSONResponse({"terms": keyword_explain.explain_terms(conn, terms)})
    finally:
        conn.close()


@app.get("/api/categories/{cid}/results")
def api_results(request: Request, cid: int, limit: int = 10, offset: int = 0):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    limit = limit if limit in (10, 50, 100) else 10
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        return JSONResponse({"error": "not found"}, status_code=404)
    filters = _effective_filters(conn, category)
    total = cached_count(conn, **filters)
    offset = max(0, min(offset, max(0, total - 1)))   # clamp within range
    rows = shortlist.detail_rows(conn, limit=limit, offset=offset, **filters) if total else []
    conn.close()
    return JSONResponse({"total": total, "shown": len(rows),
                         "offset": offset, "limit": limit, "rows": rows})


# ---- AI evaluation (inclusion pass over the shortlist), tracked per run --
def _cfg_for(conn, provider: str, model: str) -> dict:
    """Build an AI config for a run's fixed provider + model; prices from ai_models."""
    return llm.cfg_for(conn, provider, model, key=_provider_key(provider))


def _active_run(conn, cid: int) -> str:
    return settings.get_setting(conn, f"active_run_{cid}", "") or ""


# The A/B trial knobs, the stamped prompt spec and its readers live in govuk_corpus.evaluate
# (shared with the benchmark CLI); aliases keep the call sites below unchanged.
_norm_trial = evaluate.norm_trial
_prompt_spec_json = evaluate.prompt_spec_json
_run_trial = evaluate.run_trial


def _run_prompt_inputs(conn, category, run_id: str, phase: str):
    """(template, inclusion, exclusion, name, keep_hints, drop_hints) for building a run's prompts:
    the run's stamped spec when present (exact reproduction), else resolved live (legacy runs)."""
    spec = evaluate.run_prompt_spec(conn, run_id)
    if spec:
        return (spec.get("template"), spec.get("inclusion_context", ""), spec.get("exclusion_context", ""),
                spec.get("name") or "the topic", spec.get("keep_hints", ""), spec.get("drop_hints", ""))
    is_excl = phase == evaluate.PHASE_EXCLUSION
    return (prompts.default_text("exclusion" if is_excl else "inclusion"),
            category.get("inclusion_context") or "", category.get("exclusion_context") or "",
            (category.get("description") or "").strip() or cat.prettify(category.get("slug")) or "the topic",
            category.get("adjudication_hints_keep") or "", category.get("adjudication_hints_drop") or "")


def _ensure_run(conn, cid: int) -> str:
    """Return the active run for the category, creating one (current provider/model) if none."""
    run_id = _active_run(conn, cid)
    if run_id and evaluate.get_run(conn, run_id):
        return run_id
    cfg = _ai_config_for_phase(conn, evaluate.PHASE_INCLUSION)
    run_id = evaluate.create_run(conn, cid, cfg["model"], cfg["provider"],
                                 phase=evaluate.PHASE_INCLUSION,
                                 prompt_spec=_prompt_spec_json(conn, cid, evaluate.PHASE_INCLUSION))
    settings.set_setting(conn, f"active_run_{cid}", run_id)
    return run_id


def _evaluate_one_page(cfg, is_exclusion, variant, caching, tmpl, inclusion, exclusion,
                       name, keep_hints, drop_hints, r, sampling=None):
    """Build one page's prompt and call the model (govuk_corpus.evaluate_driver). Touches NO DB,
    so it's safe to run concurrently across pages; the caller persists the result. The model
    hooks are looked up here at call time so the mocked tests can patch `_ai_reply`/`_ai_chat`."""
    return evaluate_driver.evaluate_one_page(
        cfg, is_exclusion, variant, caching, tmpl, inclusion, exclusion, name, keep_hints,
        drop_hints, r, sampling=sampling, reply_fn=_ai_reply, chat_fn=_ai_chat)


def _run_evaluation(cid: int, limit: int) -> dict:
    conn = connect()
    try:
        category = cat.get_category(conn, cid)
        if not category:
            return {"error": "Shortlist not found."}
        budget = _budget(conn)
        spent = _daily_spend(conn)
        if budget > 0 and spent >= budget:
            return {"error": f"Daily AI budget of ${budget:.2f} reached (${spent:.4f} spent today). "
                    f"Raise it in Settings or try again tomorrow.",
                    "spent_today": round(spent, 4), "budget": budget, "stopped": "budget"}

        run_id = _ensure_run(conn, cid)
        evaluate.mark_run_running(conn, run_id)   # this process is now driving the run (pid + heartbeat)
        run = evaluate.get_run(conn, run_id)
        phase = run.get("phase") or evaluate.PHASE_INCLUSION
        cfg = _cfg_for(conn, run["provider"], run["model"])
        if not cfg["key"]:
            evaluate.mark_run_stopped(conn, run_id)
            return {"error": f"No API key for {cfg['label']} (this run's provider) — set one in Settings."}

        limit = min(limit, _max_docs(conn))
        filters = _effective_filters(conn, category)
        is_exclusion = phase == evaluate.PHASE_EXCLUSION
        # Use the run's stamped prompt spec (exact + fixed for the run); legacy runs resolve live.
        prompt_tmpl, inclusion, exclusion, name, keep_hints, drop_hints = _run_prompt_inputs(
            conn, category, run_id, phase)
        # A scoped run (the benchmark) evaluates a fixed url list stamped on the run instead of
        # the category's shortlist; a bench-tagged run is never auto-advanced by the app.
        scope = evaluate.run_scope(conn, run_id)
        bench = (evaluate.run_prompt_spec(conn, run_id) or {}).get("bench")
        if is_exclusion:
            rows = evaluate.exclusion_candidates(conn, run_id, run.get("source_run_id"), limit)
        elif scope:
            rows = evaluate.scoped_candidates(conn, run_id, scope["urls"], limit)
        else:
            rows = evaluate.run_candidates(conn, run_id, cid, limit,
                                           min_es_score=_GOVUK_MIN_ES_SCORE, **filters)
        trial = _run_trial(conn, run_id)
        concurrency, variant, caching = trial["concurrency"], trial["prompt_variant"], trial["caching"]
        sampling = trial["sampling"]
        done = 0
        cost = 0.0
        stopped = None
        skipped = 0
        consec_err = 0        # AI errors in a row -> likely a provider outage, so bail out
        fatal_return = None
        # Pages are evaluated in waves of `concurrency` (1 = the original sequential path). Each page
        # is still its own call + JSON verdict; only the API round-trips overlap. Every DB write stays
        # on this worker thread, after each wave completes.
        idx = 0
        while idx < len(rows):
            if budget > 0 and spent >= budget:
                stopped = "budget"
                break
            wave = rows[idx:idx + concurrency]
            idx += len(wave)
            evone = lambda rr: (rr,) + _evaluate_one_page(
                cfg, is_exclusion, variant, caching, prompt_tmpl,
                inclusion, exclusion, name, keep_hints, drop_hints, rr, sampling)
            if concurrency > 1 and len(wave) > 1:
                with concurrent.futures.ThreadPoolExecutor(max_workers=len(wave)) as ex:
                    packed = list(ex.map(evone, wave))
            else:
                packed = [evone(rr) for rr in wave]
            for r, res, ms, _prompt in packed:
                if res.get("error"):
                    # Fatal (no key / bad config) -> stop the run and report it.
                    if res.get("fatal"):
                        evaluate.mark_run_stopped(conn, run_id)
                        fatal_return = {"error": res["error"], "fatal": True, "run_id": run_id,
                                        "evaluated_this_run": done, "skipped": skipped,
                                        "cost_usd": round(cost, 6), "spent_today": round(spent, 4),
                                        "budget": budget}
                        break
                    # Transient (the SDK has already retried): if it keeps failing the provider is
                    # likely down — stop cleanly; the user re-executes to resume where it left off.
                    consec_err += 1
                    if consec_err >= MAX_CONSEC_EVAL_ERRORS:
                        evaluate.mark_run_stopped(conn, run_id)
                        fatal_return = {"error": f"Stopped after {consec_err} evaluation errors in a row "
                                        f"(last: {res['error']}). The provider may be down or rate-limiting — "
                                        f"re-execute to resume where it left off.",
                                        "run_id": run_id, "evaluated_this_run": done, "skipped": skipped,
                                        "cost_usd": round(cost, 6), "spent_today": round(spent, 4),
                                        "budget": budget, "stopped": "errors"}
                        break
                    # Isolated failure: record the page as unscored (so it's excluded next time and
                    # shows on the Run detail 'Not parsed' list) and carry on.
                    evaluate.save_page(conn, run_id, cid, r["url"],
                                       {"keep": None, "score": None,
                                        "reason": f"skipped after AI error: {str(res['error'])[:300]}"}, ms,
                                       raw_reply=res.get("reply") or res.get("error"))
                    skipped += 1
                    done += 1
                    continue
                consec_err = 0
                _decision, c = evaluate_driver.persist_result(
                    conn, run_id, cid, r, res, ms, is_exclusion=is_exclusion, kind="evaluate")
                done += 1
                cost += c
                spent += c
            if fatal_return is not None:
                break
        if fatal_return is not None:
            return fatal_return
        run = evaluate.get_run(conn, run_id)
        if is_exclusion:
            total = evaluate.kept_count(conn, run.get("source_run_id")) if run.get("source_run_id") else 0
        elif scope:
            total = len(scope["urls"])
        else:
            total = cached_count(conn, **filters)
        remaining = max(0, total - (run["pages"] or 0))
        advanced = None
        if stopped == "budget":                 # driver will step away this batch — mark it idle
            evaluate.mark_run_stopped(conn, run_id)
        if remaining == 0 and stopped is None:
            evaluate.finish_run(conn, run_id)
            # Auto-advance: when an inclusion run completes, start Phase 2 (Exclusion)
            # over the pages it kept, on the exclusion phase's model. A benchmark run drives its
            # own Phase 2 (both models off the same keeps), so it is never advanced here; any
            # other scoped run inherits its Phase-1 model so the pair stays comparable.
            if not is_exclusion and not bench:
                keeps = evaluate.kept_count(conn, run_id)
                if keeps > 0:
                    excfg = ({"model": run["model"], "provider": run["provider"]} if scope
                             else _ai_config_for_phase(conn, evaluate.PHASE_EXCLUSION))
                    new_id = evaluate.create_run(conn, cid, excfg["model"], excfg["provider"],
                                                 phase=evaluate.PHASE_EXCLUSION, source_run_id=run_id,
                                                 name=run.get("name"),   # a run's name spans both phases
                                                 prompt_spec=_prompt_spec_json(conn, cid, evaluate.PHASE_EXCLUSION,
                                                                               _run_trial(conn, run_id),
                                                                               scope=scope))
                    settings.set_setting(conn, f"active_run_{cid}", new_id)
                    advanced = {"run_id": new_id, "phase": evaluate.PHASE_EXCLUSION, "remaining": keeps}
                    remaining = keeps
        warning = (f"Within 10% of the ${budget:.2f} daily budget (${spent:.4f} spent today)."
                   if budget > 0 and spent >= 0.9 * budget else None)
        return {"run_id": run_id, "model": run["model"], "provider": run["provider"],
                "phase": phase, "evaluated_this_run": done, "skipped": skipped,
                "cost_usd": round(cost, 6), "spent_today": round(spent, 4), "budget": budget,
                "advanced": advanced, "warning": warning, "stopped": stopped,
                "remaining": remaining, "run": run}
    finally:
        conn.close()


@app.post("/api/categories/{cid}/runs")
async def api_new_run(request: Request, cid: int):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    try:
        body = await request.json()
    except Exception:
        body = {}
    trial = {"prompt_variant": (body or {}).get("prompt_variant"),
             "concurrency": (body or {}).get("concurrency"),
             "caching": (body or {}).get("caching")}
    conn = connect()
    if not cat.get_category(conn, cid):
        conn.close()
        return JSONResponse({"error": "not found"}, status_code=404)
    cfg = _ai_config_for_phase(conn, evaluate.PHASE_INCLUSION)   # new manual run = a fresh inclusion run
    run_id = evaluate.create_run(conn, cid, cfg["model"], cfg["provider"],
                                 phase=evaluate.PHASE_INCLUSION,
                                 prompt_spec=_prompt_spec_json(conn, cid, evaluate.PHASE_INCLUSION, trial))
    settings.set_setting(conn, f"active_run_{cid}", run_id)
    run = evaluate.get_run(conn, run_id)
    conn.close()
    return JSONResponse({"run_id": run_id, "run": run})


@app.post("/api/categories/{cid}/evaluate")
async def api_evaluate(request: Request, cid: int, limit: int = 10):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    limit = max(1, min(limit, 50))
    result = await run_in_threadpool(_run_evaluation, cid, limit)
    return JSONResponse(result)


# ---- server-side background evaluation ----------------------------------
# A background thread drives _run_evaluation in chunks until the shortlist is
# exhausted (through the auto-advanced exclusion phase), the daily budget stops
# it, it's stopped, or it errors — so a run finishes without the browser open.
# Progress lives in-memory per category; it does NOT survive a service restart
# (the DB run does, and can be resumed from the page).
BG_EVAL_CHUNK = 25
_bg_evals: Dict[int, dict] = {}
_bg_lock = threading.Lock()


def _background_eval_loop(cid: int, stop_event: threading.Event, status: dict) -> None:
    try:
        conn = connect()
        try:
            cap = _max_docs(conn)          # pages-per-run limit (Settings), e.g. 600
            status["run_id"] = _active_run(conn, cid) or None   # so "In Progress" shows at once
        finally:
            conn.close()
        cleanups = 0
        MAX_CLEANUPS = 3
        while not stop_event.is_set():
            res = _run_evaluation(cid, BG_EVAL_CHUNK)
            if res.get("error"):
                status["error"] = res["error"]
                logging.getLogger("assistant").warning("background eval stopped (cid=%s): %s", cid, res["error"])
                break
            status["done"] += res.get("evaluated_this_run", 0)
            status["skipped"] = status.get("skipped", 0) + res.get("skipped", 0)
            status["cost"] += res.get("cost_usd") or 0.0
            status["phase"] = res.get("phase")
            status["remaining"] = res.get("remaining")
            status["run_id"] = res.get("run_id")
            status["spent_today"] = res.get("spent_today")
            status["budget"] = res.get("budget")
            if res.get("stopped") == "budget":
                status["stopped"] = "budget"
                break
            # No auto-advance and nothing left (or nothing progressed) -> the chain looks done.
            if not res.get("advanced") and (res.get("evaluated_this_run", 0) == 0
                                            or res.get("remaining") == 0):
                # End-of-run cleanup: retry any unparsable (keep IS NULL) rows once the pipeline
                # has otherwise finished. If that recovers pages — e.g. inclusion-keeps that were
                # unparsable and so never reached exclusion — loop again so exclusion folds them
                # in. Bounded (and stops when a pass recovers nothing) so it always terminates.
                if cleanups < MAX_CLEANUPS:
                    cleanups += 1
                    try:
                        cu = _retry_unparsable(cid)
                    except Exception as e:
                        cu = {"error": str(e)}
                    status["cost"] += cu.get("cost_usd") or 0.0
                    status["done"] += cu.get("retried", 0)
                    status["cleaned"] = status.get("cleaned", 0) + cu.get("now_parsed", 0)
                    if cu.get("now_parsed"):
                        continue   # recovered pages may create exclusion work — re-drive
                break
            # Per-run page cap: evaluate up to `cap` NEW pages this execution, then
            # stop cleanly so we actually run the full 600 (not nothing) and the user
            # re-executes to continue where it left off, rather than doing the whole
            # shortlist in one unbounded run.
            if cap and status["done"] >= cap:
                status["stopped"] = "cap"
                break
    except Exception as e:                      # never let the thread die silently
        status["error"] = f"{type(e).__name__}: {e}"
    finally:
        status["running"] = False
        status["finished_at"] = db.now_iso()
        # This driver is stepping away — clear the run's 'running' flag unless it completed
        # (mark_run_stopped is a no-op on a finished run), so a stalled run is detectable.
        rid = status.get("run_id")
        if rid:
            try:
                c = connect()
                try:
                    evaluate.mark_run_stopped(c, rid)
                finally:
                    c.close()
            except Exception:
                pass


def _spawn_bg_eval(cid: int) -> dict:
    """Start the background evaluation thread for a category (or return the one already
    running). Thread-safe. Returns {'already_running': bool, 'status': <status dict>}."""
    with _bg_lock:
        ent = _bg_evals.get(cid)
        if ent and ent["status"].get("running"):
            return {"already_running": True, "status": ent["status"]}
        stop_event = threading.Event()
        status = {"running": True, "done": 0, "skipped": 0, "cost": 0.0, "phase": None,
                  "remaining": None, "run_id": None, "stopped": None, "error": None,
                  "started_at": db.now_iso(), "finished_at": None}
        t = threading.Thread(target=_background_eval_loop, args=(cid, stop_event, status), daemon=True)
        _bg_evals[cid] = {"thread": t, "stop": stop_event, "status": status}
        t.start()
        return {"already_running": False, "status": status}


def _resume_stalled_runs() -> None:
    """Re-drive any category whose ACTIVE evaluation run is unfinished. The background
    driver lives only in memory, so a deploy/crash/restart strands an in-progress run —
    the DB row stays 'in progress' with nothing pushing it forward (and Phase 2 never
    auto-advances). Called once on startup. Best-effort: never blocks or fails startup.
    Disable by setting 'auto_resume_runs' = '0'."""
    log = logging.getLogger("assistant")
    try:
        conn = connect()
        try:
            if (settings.get_setting(conn, "auto_resume_runs", "1") or "1") == "0":
                return
            rows = conn.execute(
                "SELECT DISTINCT category_id FROM evaluation_runs "
                "WHERE finished_at IS NULL AND category_id IS NOT NULL").fetchall()
            cids = []
            for r in rows:
                cid = int(dict(r)["category_id"])
                active = _active_run(conn, cid)
                run = evaluate.get_run(conn, active) if active else None
                # Only re-drive a run that was genuinely interrupted mid-flight — i.e. its
                # ACTIVE run is unfinished AND already made progress. A pages=0 unfinished
                # run is indistinguishable from a created-but-never-started draft, so we
                # leave those for the user to start rather than resurrecting abandoned drafts.
                # Skip runs a driver is already working (run_is_alive) so we never double-drive
                # one — e.g. a detached CLI run, or another worker.
                if (run and not run.get("finished_at") and (run.get("pages") or 0) > 0
                        and not evaluate.run_is_alive(run)):
                    cids.append(cid)
        finally:
            conn.close()
    except Exception as e:
        log.warning("auto-resume: could not scan for stalled runs: %s", e)
        return
    for cid in cids:
        try:
            res = _spawn_bg_eval(cid)
            if not res.get("already_running"):
                log.info("auto-resumed stalled evaluation on startup (cid=%s)", cid)
        except Exception as e:
            log.warning("auto-resume failed for cid=%s: %s", cid, e)


@app.post("/api/categories/{cid}/evaluate-bg/start")
async def api_bg_eval_start(request: Request, cid: int):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    ok = cat.get_category(conn, cid) is not None
    conn.close()
    if not ok:
        return JSONResponse({"error": "not found"}, status_code=404)
    res = _spawn_bg_eval(cid)
    if res.get("already_running"):
        return JSONResponse({"already_running": True, "status": res["status"]})
    return JSONResponse({"started": True, "status": res["status"]})


@app.post("/api/categories/{cid}/evaluate-bg/resume")
async def api_bg_eval_resume(request: Request, cid: int):
    """Complete a stalled shortlist: a run left 'running' but whose driver process is gone.
    Makes that run the active run and starts a fresh background driver to finish it."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    try:
        if not cat.get_category(conn, cid):
            return JSONResponse({"error": "not found"}, status_code=404)
        stalled = next((r for r in evaluate.list_runs(conn, cid) if evaluate.run_is_stalled(r)), None)
        if not stalled:
            return JSONResponse({"error": "No stalled run to resume — nothing to complete."},
                                status_code=409)
        evaluate.mark_run_stopped(conn, stalled["run_id"])          # clear the stale 'running' flag
        settings.set_setting(conn, f"active_run_{cid}", stalled["run_id"])
        rid = stalled["run_id"]
    finally:
        conn.close()
    res = _spawn_bg_eval(cid)
    return JSONResponse({"resumed": True, "run_id": rid,
                         "already_running": bool(res.get("already_running")), "status": res["status"]})


def _retry_unparsable(cid: int, cap: int = 300) -> dict:
    """Re-evaluate the pages whose model reply couldn't be parsed (keep IS NULL) in the latest
    inclusion run and its exclusion run — in place, in each page's own run/phase — so you can see
    whether they clear on a retry (a transient blip) or fail again (a real problem). Deletes the
    stale unparsable rows, re-asks the model, and recomputes run totals. Respects the daily budget.
    """
    conn = connect()
    try:
        category = cat.get_category(conn, cid)
        if not category:
            return {"error": "Shortlist not found."}
        budget, spent = _budget(conn), _daily_spend(conn)
        if budget > 0 and spent >= budget:
            return {"error": f"Daily AI budget of ${budget:.2f} reached (${spent:.4f} today) — "
                    f"raise it in Settings or try again tomorrow."}
        inc = evaluate.latest_inclusion_run(conn, cid)
        if not inc:
            return {"error": "No AI run yet — nothing to retry."}
        exc = evaluate.latest_exclusion_run(conn, inc)
        inclusion = category.get("inclusion_context") or ""
        exclusion = category.get("exclusion_context") or ""
        name = (category.get("description") or "").strip() or cat.prettify(category.get("slug")) or "the topic"
        P = shortlist._P
        out = {"retried": 0, "now_parsed": 0, "still_unparsable": 0, "cost_usd": 0.0, "phases": []}
        left = cap
        for run_id, is_excl in ((inc, False), (exc, True)):
            if not run_id or left <= 0:
                continue
            run = evaluate.get_run(conn, run_id)
            cfg = _cfg_for(conn, run["provider"], run["model"])
            if not cfg["key"]:
                out["phases"].append({"phase": run.get("phase"), "error": f"no API key for {cfg['label']}"})
                continue
            # Reproduce THIS run's prompt exactly: its stamped spec (template + variant + caching),
            # else live for legacy runs.
            prompt_tmpl, inclusion, exclusion, name, keep_hints, drop_hints = _run_prompt_inputs(
                conn, category, run_id, evaluate.PHASE_EXCLUSION if is_excl else evaluate.PHASE_INCLUSION)
            rt = _run_trial(conn, run_id)
            urls = [dict(r)["url"] for r in conn.execute(
                f"SELECT url FROM evaluation_results WHERE run_id = {P} AND keep IS NULL LIMIT {int(left)}",
                (run_id,)).fetchall()]
            retried = parsed = 0
            for url in urls:
                c = conn.execute(
                    f"SELECT title, description, search_text AS body, content_hash FROM content WHERE url = {P}",
                    (url,)).fetchone()
                c = dict(c) if c else {"title": "", "description": "", "body": ""}
                pass1 = pass1_topic = ""
                if is_excl:
                    pr = conn.execute(
                        f"SELECT reason, primary_topic FROM evaluation_results WHERE run_id = {P} AND url = {P}",
                        (inc, url)).fetchone()
                    pr = dict(pr) if pr else {}
                    pass1 = pr.get("reason") or ""
                    pass1_topic = pr.get("primary_topic") or ""
                row = {"url": url, "title": c.get("title"), "description": c.get("description"),
                       "body": c.get("body"), "content_hash": c.get("content_hash"),
                       "pass1_reason": pass1, "pass1_topic": pass1_topic}
                res, ms, _prompt = _evaluate_one_page(cfg, is_excl, rt["prompt_variant"], rt["caching"],
                                                      prompt_tmpl, inclusion, exclusion, name,
                                                      keep_hints, drop_hints, row, rt["sampling"])
                if res.get("error"):
                    continue                                   # leave it unparsable; provider hiccup
                conn.execute(f"DELETE FROM evaluation_results WHERE run_id = {P} AND url = {P}", (run_id, url))
                decision, cost = evaluate_driver.persist_result(
                    conn, run_id, cid, row, res, ms, is_exclusion=is_excl, kind="retry-unparsable")
                out["cost_usd"] += cost
                retried += 1
                if decision is not None and decision.get("keep") is not None:
                    parsed += 1
            # save_page's incremental totals drift after our DELETEs — recompute from the rows.
            conn.execute(
                f"UPDATE evaluation_runs SET "
                f"pages = (SELECT COUNT(*) FROM evaluation_results WHERE run_id = {P}), "
                f"kept = (SELECT COUNT(*) FROM evaluation_results WHERE run_id = {P} AND keep = 1), "
                f"dropped = (SELECT COUNT(*) FROM evaluation_results WHERE run_id = {P} AND keep = 0), "
                f"unparseable = (SELECT COUNT(*) FROM evaluation_results WHERE run_id = {P} AND keep IS NULL) "
                f"WHERE run_id = {P}", (run_id, run_id, run_id, run_id, run_id))
            conn.commit()
            still = int(dict(conn.execute(
                f"SELECT COUNT(*) AS n FROM evaluation_results WHERE run_id = {P} AND keep IS NULL",
                (run_id,)).fetchone())["n"])
            out["retried"] += retried
            out["now_parsed"] += parsed
            out["phases"].append({"phase": run.get("phase"), "retried": retried,
                                  "now_parsed": parsed, "still_unparsable": still})
            left -= retried
        out["still_unparsable"] = sum(p.get("still_unparsable", 0) for p in out["phases"])
        out["cost_usd"] = round(out["cost_usd"], 6)
        return out
    finally:
        conn.close()


@app.post("/api/categories/{cid}/retry-unparsable")
async def api_retry_unparsable(request: Request, cid: int):
    """Re-evaluate the unparsable (keep IS NULL) pages of the latest run chain, in place."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    result = await run_in_threadpool(_retry_unparsable, cid)
    return JSONResponse(result)


@app.get("/api/categories/{cid}/evaluate-bg/status")
async def api_bg_eval_status(request: Request, cid: int):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    ent = _bg_evals.get(cid)
    return JSONResponse(ent["status"] if ent else {"running": False})


@app.post("/api/categories/{cid}/evaluate-bg/stop")
async def api_bg_eval_stop(request: Request, cid: int):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    ent = _bg_evals.get(cid)
    if ent:
        ent["stop"].set()
    return JSONResponse({"stopping": ent is not None})


# ---- background GOV.UK hybrid search (compare + fetch into the corpus) ----
# Can take minutes (up to 25 phrases x 10k pages), so it runs off-request in a thread,
# per category, with pollable progress. In-memory; does not survive a restart (the results
# it writes to category_search_pages / content do).
_bg_hybrid: Dict[int, dict] = {}


def _background_hybrid_loop(cid: int, stop_event: threading.Event, status: dict) -> None:
    try:
        conn = connect()
        try:
            category = cat.get_category(conn, cid)
            if not category:
                status["error"] = "shortlist not found"
                return
            filters = _effective_filters(conn, category)
            status["phase"] = "searching"

            def prog(done, total, pages):
                status["phrases_done"], status["phrases_total"], status["pages"] = done, total, pages

            def search_fn(phrases, org_slugs, doctypes, progress=None):
                return _govuk_search_multi(phrases, org_slugs, doctypes, progress=progress,
                                           should_stop=stop_event.is_set)
            summary = search_augment.compare(
                conn, cid, filters["keywords"], filters["organisations"], search_fn,
                document_types=filters["document_types"], progress=prog)
            status["counts"] = summary.get("counts")
            if stop_event.is_set():
                status["stopped"] = "user"
                return
            # Fetch GOV.UK-only pages into the corpus (bounded batches) so they're evaluable.
            status["phase"] = "fetching"
            while not stop_event.is_set():
                urls = search_augment.pending_fetch_urls(conn, cid)[:_GOVUK_FETCH_CAP]
                if not urls:
                    break
                run_id = db.start_run(conn, stage="govuk-augment", scope=str(cid))
                counters = stage_align.align_urls(conn, run_id, urls, source="govuk-search", stage="govuk-augment")
                db.finish_run(conn, run_id, counters)
                search_augment.refresh_content_ids(conn, cid)
                status["fetched"] = status.get("fetched", 0) + (
                    counters.get("new", 0) + counters.get("changed", 0) + counters.get("unchanged", 0))
                status["pending"] = len(search_augment.pending_fetch_urls(conn, cid))
            # Drop container rows whose html_publication attachment we already hold (re-tag it).
            status["resolved"] = search_augment.resolve_attachment_containers(conn, cid)
            status["evaluable"] = search_augment.evaluable_search_only(conn, cid)
        finally:
            conn.close()
    except Exception as e:
        status["error"] = f"{type(e).__name__}: {e}"
    finally:
        status["running"] = False
        status["finished_at"] = db.now_iso()


def _start_hybrid_bg(cid: int) -> dict:
    """Start (or report the running) background hybrid job for a category."""
    with _bg_lock:
        ent = _bg_hybrid.get(cid)
        if ent and ent["status"].get("running"):
            return {"already_running": True, "status": ent["status"]}
        stop_event = threading.Event()
        status = {"running": True, "phase": "starting", "phrases_done": 0, "phrases_total": 0,
                  "pages": 0, "fetched": 0, "pending": None, "evaluable": None, "counts": None,
                  "stopped": None, "error": None, "started_at": db.now_iso(), "finished_at": None}
        t = threading.Thread(target=_background_hybrid_loop, args=(cid, stop_event, status), daemon=True)
        _bg_hybrid[cid] = {"thread": t, "stop": stop_event, "status": status}
        t.start()
        return {"started": True, "status": status}


@app.post("/api/categories/{cid}/hybrid-bg/start")
async def api_hybrid_bg_start(request: Request, cid: int):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    ok = cat.get_category(conn, cid) is not None
    conn.close()
    if not ok:
        return JSONResponse({"error": "not found"}, status_code=404)
    return JSONResponse(_start_hybrid_bg(cid))


@app.get("/api/categories/{cid}/hybrid-bg/status")
async def api_hybrid_bg_status(request: Request, cid: int):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    ent = _bg_hybrid.get(cid)
    return JSONResponse(ent["status"] if ent else {"running": False})


@app.post("/api/categories/{cid}/hybrid-bg/stop")
async def api_hybrid_bg_stop(request: Request, cid: int):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    ent = _bg_hybrid.get(cid)
    if ent:
        ent["stop"].set()
    return JSONResponse({"stopping": ent is not None})


@app.get("/categories/{cid}/performance", response_class=HTMLResponse)
def performance_page(request: Request, cid: int):
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    category["display_name"] = cat.display_name(category)
    resp = templates.TemplateResponse("performance.html", ctx(conn, request, category=category))
    conn.close()
    return resp


# ---- Run-vs-baseline disagreement drilldown (guc-0028) -------------------
# "Why did the runs differ?" — the per-page detail behind the performance page's
# "Decisions disagree with Run 1" count, per phase, with an on-demand LLM verdict.
_DIFF_NOTES_READY = False


def _ensure_diff_notes(conn) -> None:
    """Create the cache table for LLM difference verdicts (idempotent, PG + sqlite)."""
    global _DIFF_NOTES_READY
    conn.execute(
        "CREATE TABLE IF NOT EXISTS run_diff_notes ("
        "base_run text, other_run text, url text, phase text, "
        "verdict text, model text, created_at text, "
        "PRIMARY KEY (base_run, other_run, url, phase))")
    conn.commit()
    _DIFF_NOTES_READY = True


def _chain_of(conn, head_run_id: str):
    """(inclusion_row, exclusion_row) for a chain identified by its Phase-1 run_id."""
    P = shortlist._P
    if not head_run_id:
        return None, None
    r = conn.execute(f"SELECT * FROM evaluation_runs WHERE run_id = {P}", (head_run_id,)).fetchone()
    incl = dict(r) if r else None
    # The exclusion (Phase 2) run is the child that points back at this head. The stored phase
    # label varies ("Phase 2 - Exclusion"), so identify it by source_run_id, not the label.
    r = conn.execute(
        f"SELECT * FROM evaluation_runs WHERE source_run_id = {P} "
        f"ORDER BY started_at DESC LIMIT 1", (head_run_id,)).fetchone()
    excl = dict(r) if r else None
    return incl, excl


def _reasons_map(conn, run_id) -> dict:
    """{url: {keep, score, reason, raw_reply}} for one phase-run."""
    if not run_id:
        return {}
    P = shortlist._P
    rows = conn.execute(
        f"SELECT url, keep, score, reason, raw_reply FROM evaluation_results WHERE run_id = {P}",
        (run_id,)).fetchall()
    return {r["url"]: dict(r) for r in rows}


def _diff_notes_map(conn, base: str, other: str) -> dict:
    """{(url, phase): verdict} for cached difference verdicts of this run pair."""
    _ensure_diff_notes(conn)
    P = shortlist._P
    rows = conn.execute(
        f"SELECT url, phase, verdict FROM run_diff_notes WHERE base_run = {P} AND other_run = {P}",
        (base, other)).fetchall()
    return {(r["url"], r["phase"]): r["verdict"] for r in rows}


def _diff_judge_config(conn, run_row) -> dict:
    """Config for the LLM that adjudicates why runs differ — the SAME model the given run used
    (its Phase-2/exclusion model), falling back to the configured exclusion-phase model."""
    if run_row and run_row.get("model"):
        m = ai_models.find(conn, run_row.get("provider"), run_row.get("model"))
        if m:
            return _config_from_model(conn, m)
    return _ai_config_for_phase(conn, "exclusion")


@app.get("/categories/{cid}/performance/diff", response_class=HTMLResponse)
def performance_diff_page(request: Request, cid: int, base: str = "", other: str = ""):
    """Drilldown: pages where a run disagrees with the baseline (Run 1), per phase, with both
    runs' inclusion & exclusion reasons side by side and an on-demand LLM difference verdict."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    category["display_name"] = cat.display_name(category)

    b_incl, b_excl = _chain_of(conn, base)
    o_incl, o_excl = _chain_of(conn, other)
    maps = {
        "b_incl": _reasons_map(conn, b_incl and b_incl["run_id"]),
        "b_excl": _reasons_map(conn, b_excl and b_excl["run_id"]),
        "o_incl": _reasons_map(conn, o_incl and o_incl["run_id"]),
        "o_excl": _reasons_map(conn, o_excl and o_excl["run_id"]),
    }
    notes = _diff_notes_map(conn, base, other)
    adj = _diff_adjud_map(conn, base, other)

    def card(url, phase):
        return {"url": url, "phase": phase,
                "b_incl": maps["b_incl"].get(url), "b_excl": maps["b_excl"].get(url),
                "o_incl": maps["o_incl"].get(url), "o_excl": maps["o_excl"].get(url),
                "verdict": notes.get((url, phase)), "choice": adj.get((url, phase))}

    p1 = (evaluate.run_disagreements(conn, b_incl["run_id"], o_incl["run_id"])
          if (b_incl and o_incl) else [])
    p2 = (evaluate.run_disagreements(conn, b_excl["run_id"], o_excl["run_id"])
          if (b_excl and o_excl) else [])
    p1_cards = [card(r["url"], "inclusion") for r in p1]
    p2_cards = [card(r["url"], "exclusion") for r in p2]

    # Pages whose FINAL shortlist result differs between the two chains (end of pipeline),
    # which is a subset/superset-independent view of the per-phase disagreements above.
    final_diffs = evaluate.final_outcome_diffs(
        conn, b_incl and b_incl["run_id"], b_excl and b_excl["run_id"],
        o_incl and o_incl["run_id"], o_excl and o_excl["run_id"])

    trials = {"base": _run_trial(conn, b_incl["run_id"]) if b_incl else {},
              "other": _run_trial(conn, o_incl["run_id"]) if o_incl else {}}
    judge = _diff_judge_config(conn, o_excl or o_incl)
    ctxd = ctx(conn, request, category=category, base_head=base, other_head=other,
               base_incl=b_incl, base_excl=b_excl, other_incl=o_incl, other_excl=o_excl,
               base_has_excl=bool(b_excl), other_has_excl=bool(o_excl),
               p1_cards=p1_cards, p2_cards=p2_cards, final_diffs=final_diffs,
               trials=trials, judged_count=len(adj),
               judge_label=f"{judge.get('provider','')} / {judge.get('model','')}")
    resp = templates.TemplateResponse("performance_diff.html", ctxd)
    conn.close()
    return resp


@app.post("/api/categories/{cid}/performance/diff/explain")
async def api_performance_diff_explain(request: Request, cid: int):
    """Generate (and cache) an LLM verdict on why the baseline and other run disagreed on one
    page in one phase. Uses the same model the other run used, so it 'speaks the run's language'."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    body = await request.json()
    base = str(body.get("base") or ""); other = str(body.get("other") or "")
    url = str(body.get("url") or ""); phase = str(body.get("phase") or "")
    if phase not in ("inclusion", "exclusion") or not (base and other and url):
        return JSONResponse({"error": "bad request"}, status_code=400)
    conn = connect()
    try:
        _ensure_diff_notes(conn)
        P = shortlist._P
        row = conn.execute(
            f"SELECT verdict FROM run_diff_notes WHERE base_run = {P} AND other_run = {P} "
            f"AND url = {P} AND phase = {P}", (base, other, url, phase)).fetchone()
        if row and row["verdict"]:
            return JSONResponse({"verdict": row["verdict"], "cached": True})

        category = cat.get_category(conn, cid) or {}
        b_incl, b_excl = _chain_of(conn, base)
        o_incl, o_excl = _chain_of(conn, other)
        b_run = b_incl if phase == "inclusion" else b_excl
        o_run = o_incl if phase == "inclusion" else o_excl
        if not (b_run and o_run):
            return JSONResponse({"error": "that phase is not present in both runs"}, status_code=400)

        br = _reasons_map(conn, b_run["run_id"]).get(url) or {}
        orr = _reasons_map(conn, o_run["run_id"]).get(url) or {}
        # cross-phase context (the other phase's reason, when the page reached it)
        b_incl_r = (_reasons_map(conn, b_incl["run_id"]).get(url) if b_incl else None) or {}
        o_incl_r = (_reasons_map(conn, o_incl["run_id"]).get(url) if o_incl else None) or {}
        b_excl_r = (_reasons_map(conn, b_excl["run_id"]).get(url) if b_excl else None) or {}
        o_excl_r = (_reasons_map(conn, o_excl["run_id"]).get(url) if o_excl else None) or {}

        kd = lambda k: "KEEP" if k == 1 else ("DROP" if k == 0 else "—")
        t_base = _run_trial(conn, b_incl["run_id"]) if b_incl else {}
        t_other = _run_trial(conn, o_incl["run_id"]) if o_incl else {}

        def cfg_line(t):
            return (f"variant={t.get('variant', 'current')}, concurrency={t.get('concurrency', 1)}, "
                    f"caching={'on' if t.get('caching') else 'off'}, "
                    f"prompt_version={t.get('template_version', '?')}, "
                    f"body_limit={t.get('body_limit', '?')}")

        phase_word = "INCLUSION (Phase 1)" if phase == "inclusion" else "EXCLUSION (Phase 2)"
        system = (
            "You are an evaluation analyst comparing two automated runs that each decided whether a "
            "GOV.UK page belongs in a shortlist. Explain, concisely and specifically, WHY the two runs "
            "reached different decisions on this page for the phase in question. Attribute the "
            "difference to a concrete cause: a genuinely borderline page, a difference in how each run "
            "read the criteria, or a run-configuration difference (prompt variant/version, body limit). "
            "Do NOT re-judge the page or say which run is 'correct'. 2–4 sentences, plain English.")
        prompt = (
            f"TOPIC: {category.get('display_name') or ''}\n"
            f"INCLUDE CRITERIA: {category.get('inclusion_context') or '(none)'}\n"
            f"EXCLUDE CRITERIA: {category.get('exclusion_context') or '(none)'}\n\n"
            f"PAGE: {url}\n\n"
            f"PHASE IN QUESTION: {phase_word}\n\n"
            f"RUN 1 (baseline) — config: {cfg_line(t_base)}\n"
            f"  {phase_word} decision: {kd(br.get('keep'))}\n"
            f"  {phase_word} reason: {br.get('reason') or '(none recorded)'}\n"
            f"  Inclusion reason: {b_incl_r.get('reason') or '(n/a)'}\n"
            f"  Exclusion reason: {b_excl_r.get('reason') or '(n/a — not reached)'}\n\n"
            f"RUN 2 — config: {cfg_line(t_other)}\n"
            f"  {phase_word} decision: {kd(orr.get('keep'))}\n"
            f"  {phase_word} reason: {orr.get('reason') or '(none recorded)'}\n"
            f"  Inclusion reason: {o_incl_r.get('reason') or '(n/a)'}\n"
            f"  Exclusion reason: {o_excl_r.get('reason') or '(n/a — not reached)'}\n\n"
            "Why did Run 1 and Run 2 differ on this page in this phase?")

        judge = _diff_judge_config(conn, o_excl or o_incl)
        budget, spent = _budget(conn), _daily_spend(conn)
        if budget > 0 and spent >= budget:
            return JSONResponse({"error": f"Daily AI budget of ${budget:.2f} reached "
                                 f"(${spent:.4f} spent today)."}, status_code=402)
        res = await run_in_threadpool(_ai_chat, judge, system,
                                      [{"role": "user", "content": prompt}], 700)
        if res.get("error") or res.get("fatal"):
            return JSONResponse({"error": res.get("error") or "AI error"}, status_code=502)
        _log_ai_usage(conn, res.get("cost_usd") or 0.0, res.get("input_tokens"),
                      res.get("output_tokens"), "judge")
        verdict = (res.get("reply") or "").strip()
        conn.execute(
            f"INSERT INTO run_diff_notes (base_run, other_run, url, phase, verdict, model, created_at) "
            f"VALUES ({P}, {P}, {P}, {P}, {P}, {P}, {P}) "
            f"ON CONFLICT (base_run, other_run, url, phase) DO UPDATE SET "
            f"verdict = EXCLUDED.verdict, model = EXCLUDED.model, created_at = EXCLUDED.created_at",
            (base, other, url, phase, verdict, res.get("actual_model") or judge.get("model"),
             db.now_iso()))
        conn.commit()
        return JSONResponse({"verdict": verdict, "model": res.get("actual_model"),
                             "cost": res.get("cost_usd")})
    finally:
        conn.close()


# ---- Gold labels from a run's outcomes (guc-0029) ------------------------
# Rate one run chain page by page — "the pipeline was right / wrong here", with a reason — and
# the verdict becomes a gold label (in / out / borderline) in category_gold_labels, the ground
# truth the benchmark (govuk_corpus.bench) scores against. Partial sets are fine.

def _gold_targets(conn, cid: int, run_id: str, stage_counts: dict) -> dict:
    """Per-stage sampling targets for (category, run): the saved ones, else the defaults
    (whole stage when it is small, otherwise gold.SAMPLE_DEFAULT_TARGET)."""
    raw = settings.get_setting(conn, f"gold_sample_{cid}_{run_id}", "") if run_id else ""
    targets = gold.default_targets(stage_counts)
    if raw:
        try:
            saved = json.loads(raw)
            for k, v in (saved.get("targets") or {}).items():
                if k in stage_counts and v is not None:
                    targets[k] = max(0, min(int(v), stage_counts[k]))
        except (ValueError, TypeError):
            pass
    return targets


def _gold_pages_sampled(conn, category, run_id: str):
    """(pages, selection, targets, stage_counts) for a run with the sampling applied."""
    pages = gold.run_pages(conn, category, run_id)
    stage_counts = {k: sum(1 for p in pages if p["stage"] == k) for k in gold.STAGES}
    targets = _gold_targets(conn, int(category["id"]), run_id, stage_counts)
    selection = gold.sample_selection(pages, targets)
    gold.apply_sampling(pages, selection)
    return pages, selection, targets, stage_counts


@app.get("/categories/{cid}/gold", response_class=HTMLResponse)
def gold_page(request: Request, cid: int, run: str = ""):
    return _gold_page(request, cid, run, wizard=False)


@app.get("/categories/{cid}/gold/wizard", response_class=HTMLResponse)
def gold_wizard_page(request: Request, cid: int, run: str = ""):
    """guc-0031: the same labelling page one card at a time, with Prev / Next that save."""
    return _gold_page(request, cid, run, wizard=True)


def _gold_page(request: Request, cid: int, run: str, wizard: bool):
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    try:
        category = cat.get_category(conn, cid)
        if not category:
            return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
        category["display_name"] = cat.display_name(category)
        head = run or evaluate.latest_inclusion_run(conn, cid) or ""
        incl = evaluate.get_run(conn, head) if head else None
        if incl and str(incl.get("category_id")) != str(cid):
            incl = None
        excl_id = evaluate.latest_exclusion_run(conn, head) if incl else None
        excl = evaluate.get_run(conn, excl_id) if excl_id else None
        if incl:
            pages, selection, targets, stage_counts = _gold_pages_sampled(conn, category, head)
            progress = gold.sample_progress(pages, selection)
        else:
            pages, selection, targets, stage_counts, progress = [], {}, {}, {}, {}
        # Inclusion runs (chain heads) of this category, for the run picker.
        heads = [r for r in evaluate.list_runs(conn, cid) if not r.get("source_run_id")]
        st = gold.status(conn, cid)
        cu = current_user(request)
        ctxd = ctx(conn, request, category=category, run=incl, excl=excl, pages=pages, heads=heads,
                   stages=gold.STAGES, stage_counts=stage_counts, gold_status=st,
                   targets=targets, progress=progress,
                   sample_total=sum(s["target"] for s in selection.values()),
                   labelled_here=sum(1 for p in pages if p["label"]),
                   my_labeller=(cu.get("email") if cu else "") or "",
                   labellers=gold.labellers(conn, cid), wizard=wizard)
        return templates.TemplateResponse("gold_labels.html", ctxd)
    finally:
        conn.close()


@app.post("/api/categories/{cid}/gold")
async def api_gold_label(request: Request, cid: int):
    """Save (or clear, with an empty verdict) the gold label for one page, judged against the
    outcome of the given inclusion run's chain."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    body = await request.json()
    run_id = str(body.get("run") or ""); url = str(body.get("url") or "")
    verdict = str(body.get("verdict") or ""); rationale = str(body.get("rationale") or "")
    label = str(body.get("label") or "")            # blind mode: the label itself, no verdict
    blind = bool(body.get("blind"))
    name = str(body.get("labeller") or "").strip()[:80]
    if not (run_id and url):
        return JSONResponse({"error": "bad request"}, status_code=400)
    conn = connect()
    try:
        category = cat.get_category(conn, cid)
        run_row = evaluate.get_run(conn, run_id)
        if not category or not run_row or str(run_row.get("category_id")) != str(cid):
            return JSONResponse({"error": "not found"}, status_code=404)
        pages, selection, _t, _c = _gold_pages_sampled(conn, category, run_id)
        page = next((p for p in pages if p["url"] == url), None)
        if not page:
            return JSONResponse({"error": "that page is not in this run"}, status_code=404)
        cu = current_user(request)
        who = (cu.get("email") if cu else None) or name     # signed-in email wins; else the typed name
        try:
            consensus = gold.save_verdict(conn, category, page, verdict, rationale, who,
                                          sample_run_id=run_id, label=label or None, blind=blind)
        except ValueError as e:
            return JSONResponse({"error": str(e)}, status_code=400)
        page["label"] = consensus or ""
        votes = gold.load_votes(conn, cid, url)
        mine = next((v for v in votes if v["labeller"] == who), None)
        row = next((r for r in gold.load_gold(conn, cid) if r["url"] == url), None)
        st = gold.status(conn, cid)
        return JSONResponse({"ok": True, "label": consensus, "my_label": (mine or {}).get("label"),
                            "labeller": who, "agreement": (row or {}).get("agreement"),
                            "votes": [{"labeller": v["labeller"], "label": v["label"], "rationale": v["rationale"],
                                       "blind": bool(v.get("blind"))} for v in votes],
                            "counts": st["counts"], "labelled": st["labelled"],
                            "forwarded": st["forwarded"], "gold_sha": st["gold_sha"],
                            "progress": gold.sample_progress(pages, selection)})
    finally:
        conn.close()


@app.get("/categories/{cid}/gold/agreement", response_class=HTMLResponse)
def gold_agreement_page(request: Request, cid: int):
    """Inter-labeller agreement (guc-0030): who labelled what, pairwise and Fleiss' κ, the pages
    the labellers disagree on with an adjudicate control, and the adjudications so far."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    try:
        category = cat.get_category(conn, cid)
        if not category:
            return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
        category["display_name"] = cat.display_name(category)
        ag = gold.agreement_stats(conn, cid)
        rows = {r["url"]: r for r in gold.load_gold(conn, cid)}
        titles = {}
        urls = list(rows)
        P = shortlist._P
        for i in range(0, len(urls), 500):
            chunk = urls[i:i + 500]
            for r in conn.execute(f"SELECT url, title FROM content WHERE url IN ({','.join([P] * len(chunk))})",
                                  tuple(chunk)).fetchall():
                titles[r["url"]] = r["title"]
        for d in ag["disagreements"]:
            d["title"] = titles.get(d["url"]) or ""
            d["consensus"] = rows.get(d["url"]) or {}
        adjudicated = [dict(r, title=titles.get(r["url"]) or "") for r in rows.values() if r.get("adjudicated_at")]
        st = gold.status(conn, cid)
        cu = current_user(request)
        ctxd = ctx(conn, request, category=category, ag=ag, adjudicated=adjudicated, gold_status=st,
                   my_labeller=(cu.get("email") if cu else "") or "")
        return templates.TemplateResponse("gold_agreement.html", ctxd)
    finally:
        conn.close()


@app.post("/api/categories/{cid}/gold/adjudicate")
async def api_gold_adjudicate(request: Request, cid: int):
    """Set (label + note) or clear (empty label) the adjudicated consensus for one page."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    body = await request.json()
    url = str(body.get("url") or ""); label = str(body.get("label") or "")
    note = str(body.get("note") or ""); name = str(body.get("adjudicator") or "").strip()[:80]
    if not url:
        return JSONResponse({"error": "bad request"}, status_code=400)
    conn = connect()
    try:
        if not cat.get_category(conn, cid):
            return JSONResponse({"error": "not found"}, status_code=404)
        cu = current_user(request)
        who = (cu.get("email") if cu else None) or name or "adjudicator"
        try:
            result = gold.adjudicate(conn, cid, url, label or None, note, who)
        except ValueError as e:
            return JSONResponse({"error": str(e)}, status_code=400)
        row = next((r for r in gold.load_gold(conn, cid) if r["url"] == url), None)
        return JSONResponse({"ok": True, "label": result, "agreement": (row or {}).get("agreement"),
                            "counts": gold.status(conn, cid)["counts"]})
    finally:
        conn.close()


@app.post("/api/categories/{cid}/gold/sample")
async def api_gold_sample(request: Request, cid: int):
    """Save the per-stage sampling targets for a run (how many pages of each stage make up the
    minimum gold set). Returns the resulting selection so the page can repaint."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    body = await request.json()
    run_id = str(body.get("run") or "")
    targets_in = body.get("targets") or {}
    conn = connect()
    try:
        category = cat.get_category(conn, cid)
        run_row = evaluate.get_run(conn, run_id) if run_id else None
        if not category or not run_row or str(run_row.get("category_id")) != str(cid):
            return JSONResponse({"error": "not found"}, status_code=404)
        clean = {}
        for k, v in targets_in.items():
            if k in gold.STAGES:
                try:
                    clean[k] = max(0, int(v))
                except (TypeError, ValueError):
                    return JSONResponse({"error": f"bad target for {k}"}, status_code=400)
        settings.set_setting(conn, f"gold_sample_{cid}_{run_id}",
                             json.dumps({"targets": clean, "seed": gold.DEFAULT_SEED, "updated_at": db.now_iso()}))
        pages, selection, targets, _c = _gold_pages_sampled(conn, category, run_id)
        return JSONResponse({"ok": True, "targets": targets,
                            "sampled": {s: sorted(v["urls"]) for s, v in selection.items()},
                            "progress": gold.sample_progress(pages, selection),
                            "sample_total": sum(s["target"] for s in selection.values())})
    finally:
        conn.close()


@app.get("/categories/{cid}/gold/sheet.csv")
def gold_sheet_download(request: Request, cid: int):
    """The labelling sheet (govuk_corpus.gold export) pre-filled with the labels so far."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    try:
        category = cat.get_category(conn, cid)
        if not category:
            return PlainTextResponse("not found", status_code=404)
        rows = gold.build_sheet(conn, cid)
    finally:
        conn.close()
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=gold.SHEET_COLUMNS, extrasaction="ignore")
    w.writeheader()
    for r in rows:
        w.writerow({k: ("" if r.get(k) is None else r.get(k)) for k in gold.SHEET_COLUMNS})
    fname = f"{category.get('slug') or cid}-gold-sheet.csv"
    return PlainTextResponse(buf.getvalue(), media_type="text/csv",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


# ---- Adjudication + prompt-improvement suggestions (guc-0028) ------------
# The user marks which run got a disagreement right; those judgements feed an LLM that
# proposes edits to THIS category's include/exclude criteria to reduce future divergence.
def _ensure_diff_adjud(conn) -> None:
    conn.execute(
        "CREATE TABLE IF NOT EXISTS run_diff_adjud ("
        "base_run text, other_run text, url text, phase text, "
        "choice text, created_at text, "
        "PRIMARY KEY (base_run, other_run, url, phase))")
    conn.commit()


def _diff_adjud_map(conn, base: str, other: str) -> dict:
    """{(url, phase): choice} — choice in {run1, run2, neither}."""
    _ensure_diff_adjud(conn)
    P = shortlist._P
    rows = conn.execute(
        f"SELECT url, phase, choice FROM run_diff_adjud WHERE base_run = {P} AND other_run = {P}",
        (base, other)).fetchall()
    return {(r["url"], r["phase"]): r["choice"] for r in rows}


@app.post("/api/categories/{cid}/performance/diff/adjudicate")
async def api_performance_diff_adjudicate(request: Request, cid: int):
    """Record (or clear) which run the user judges correct for one disagreement."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    body = await request.json()
    base = str(body.get("base") or ""); other = str(body.get("other") or "")
    url = str(body.get("url") or ""); phase = str(body.get("phase") or "")
    choice = str(body.get("choice") or "")
    if phase not in ("inclusion", "exclusion") or not (base and other and url):
        return JSONResponse({"error": "bad request"}, status_code=400)
    if choice and choice not in ("run1", "run2", "neither"):
        return JSONResponse({"error": "bad choice"}, status_code=400)
    conn = connect()
    try:
        _ensure_diff_adjud(conn)
        P = shortlist._P
        if not choice:   # empty choice clears the judgement
            conn.execute(
                f"DELETE FROM run_diff_adjud WHERE base_run = {P} AND other_run = {P} "
                f"AND url = {P} AND phase = {P}", (base, other, url, phase))
        else:
            conn.execute(
                f"INSERT INTO run_diff_adjud (base_run, other_run, url, phase, choice, created_at) "
                f"VALUES ({P}, {P}, {P}, {P}, {P}, {P}) "
                f"ON CONFLICT (base_run, other_run, url, phase) DO UPDATE SET "
                f"choice = EXCLUDED.choice, created_at = EXCLUDED.created_at",
                (base, other, url, phase, choice, db.now_iso()))
        conn.commit()
        judged = conn.execute(
            f"SELECT COUNT(*) AS n FROM run_diff_adjud WHERE base_run = {P} AND other_run = {P}",
            (base, other)).fetchone()["n"]
        return JSONResponse({"ok": True, "judged": judged})
    finally:
        conn.close()


@app.post("/api/categories/{cid}/performance/diff/suggest")
async def api_performance_diff_suggest(request: Request, cid: int):
    """From the adjudicated disagreements, ask the model for concrete edits to this category's
    include/exclude criteria that would reduce divergence. Suggest-only — nothing is applied."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    body = await request.json()
    base = str(body.get("base") or ""); other = str(body.get("other") or "")
    if not (base and other):
        return JSONResponse({"error": "bad request"}, status_code=400)
    conn = connect()
    try:
        category = cat.get_category(conn, cid) or {}
        adj = _diff_adjud_map(conn, base, other)
        if not adj:
            return JSONResponse({"error": "Judge at least one disagreement first "
                                 "(mark which run got it right)."}, status_code=400)
        b_incl, b_excl = _chain_of(conn, base)
        o_incl, o_excl = _chain_of(conn, other)
        run_of = {"inclusion": (b_incl, o_incl), "exclusion": (b_excl, o_excl)}
        maps = {}
        for ph, (br, orr) in run_of.items():
            maps[("b", ph)] = _reasons_map(conn, br and br["run_id"])
            maps[("o", ph)] = _reasons_map(conn, orr and orr["run_id"])

        kd = lambda k: "KEEP" if k == 1 else ("DROP" if k == 0 else "—")
        cases = []
        for (url, ph), choice in sorted(adj.items()):
            b = maps[("b", ph)].get(url) or {}
            o = maps[("o", ph)].get(url) or {}
            correct = {"run1": "Run 1", "run2": "Run 2", "neither": "NEITHER (both wrong)"}.get(choice, choice)
            cases.append(
                f"- PAGE: {url}\n"
                f"  PHASE: {ph}\n"
                f"  Run 1: {kd(b.get('keep'))} (score {b.get('score')}) — {b.get('reason') or '(none)'}\n"
                f"  Run 2: {kd(o.get('keep'))} (score {o.get('score')}) — {o.get('reason') or '(none)'}\n"
                f"  HUMAN SAYS CORRECT: {correct}")
        cases_block = "\n".join(cases)

        system = (
            "You are a prompt engineer improving the INCLUDE and EXCLUDE criteria a team uses to shortlist "
            "GOV.UK pages with an LLM. You are given the current criteria and a set of pages where two runs "
            "DISAGREED, each labelled by a human with the correct decision. Propose concrete, minimal edits "
            "to the INCLUDE and/or EXCLUDE criteria that would make future runs match the human labels and "
            "reduce divergence — generalise, do not overfit to single pages or hard-code URLs. Prefer "
            "clarifying ambiguous wording and adding explicit rules for the patterns you see. Only if a "
            "problem clearly comes from the shared BASE TEMPLATE (not these criteria) may you note it "
            "separately as advisory — do not rewrite the base template. "
            "Reply in this structure, plain text (no code fences):\n"
            "DIAGNOSIS: 1–3 sentences on the pattern behind the disagreements.\n"
            "INCLUDE CRITERIA — suggested: the full revised include criteria, ready to paste (or 'no change').\n"
            "EXCLUDE CRITERIA — suggested: the full revised exclude criteria, ready to paste (or 'no change').\n"
            "WHY: bullet points tying each edit to the judged cases.\n"
            "BASE TEMPLATE NOTE (optional): advisory only.")
        prompt = (
            f"TOPIC: {category.get('display_name') or ''}\n\n"
            f"CURRENT INCLUDE CRITERIA:\n{category.get('inclusion_context') or '(none)'}\n\n"
            f"CURRENT EXCLUDE CRITERIA:\n{category.get('exclusion_context') or '(none)'}\n\n"
            f"JUDGED DISAGREEMENTS ({len(cases)}):\n{cases_block}\n\n"
            "Propose the criteria edits.")

        cfg = _ai_config_for_phase(conn, "exclusion")
        budget, spent = _budget(conn), _daily_spend(conn)
        if budget > 0 and spent >= budget:
            return JSONResponse({"error": f"Daily AI budget of ${budget:.2f} reached "
                                 f"(${spent:.4f} spent today)."}, status_code=402)
        res = await run_in_threadpool(_ai_chat, cfg, system,
                                      [{"role": "user", "content": prompt}], 1500)
        if res.get("error") or res.get("fatal"):
            return JSONResponse({"error": res.get("error") or "AI error"}, status_code=502)
        _log_ai_usage(conn, res.get("cost_usd") or 0.0, res.get("input_tokens"),
                      res.get("output_tokens"), "suggest")
        return JSONResponse({"suggestion": (res.get("reply") or "").strip(),
                             "judged": len(cases), "model": res.get("actual_model"),
                             "cost": res.get("cost_usd"),
                             "current_include": category.get("inclusion_context") or "",
                             "current_exclude": category.get("exclusion_context") or ""})
    finally:
        conn.close()


@app.get("/categories/{cid}/url-check", response_class=HTMLResponse)
def url_check_page(request: Request, cid: int):
    """Check a set of pasted URLs against this category: in the corpus, active, and in the
    final shortlist. An Advanced-level diagnostic page (guc-0018)."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    category["display_name"] = cat.display_name(category)
    resp = templates.TemplateResponse("url_check.html", ctx(conn, request, category=category))
    conn.close()
    return resp


def _clean_urls(raw) -> list:
    """One-per-line or list -> a de-duplicated list of trimmed URLs."""
    lines = raw.splitlines() if isinstance(raw, str) else (raw if isinstance(raw, list) else [])
    out, seen = [], set()
    for u in lines:
        u = str(u or "").strip()
        if u and u not in seen:
            seen.add(u)
            out.append(u)
    return out


@app.post("/api/categories/{cid}/url-check")
async def api_url_check(request: Request, cid: int):
    """Check this category's "Should be in" and "Should not be in" URL lists. For each URL
    return its `kind` ("Should be in" | "Should not be in"), `corpus` (a row exists), `active` (a live,
    fetched page — not a redirect or withdrawn), and `final` (kept in this category's final
    shortlist, after inclusion + exclusion). Lists are de-duplicated and capped."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    body = await request.json()
    expected = _clean_urls(body.get("expected"))
    unexpected = _clean_urls(body.get("unexpected"))
    # One ordered, de-duplicated list — a URL in both lists is treated as "Should be in".
    tagged, seen = [], set()
    for u in expected:
        if u not in seen:
            seen.add(u); tagged.append((u, "Should be in"))
    for u in unexpected:
        if u not in seen:
            seen.add(u); tagged.append((u, "Should not be in"))
    tagged = tagged[:1000]
    if not tagged:
        return JSONResponse({"results": [], "has_run": False})
    urls = [u for u, _ in tagged]
    conn = connect()
    try:
        category = cat.get_category(conn, cid)
        if not category:
            return JSONResponse({"error": "not found"}, status_code=404)
        ph = ",".join([shortlist._P] * len(urls))
        # In the corpus, and active (live/fetched, not a redirect or withdrawn).
        info = {dict(r)["url"]: dict(r) for r in conn.execute(
            f"SELECT url, is_redirect, withdrawn, content_hash FROM content WHERE url IN ({ph})",
            tuple(urls)).fetchall()}
        # In this category's final shortlist (the exclusion run's keeps, else inclusion's).
        final, inc = set(), evaluate.latest_inclusion_run(conn, category["id"])
        if inc:
            run_id = evaluate.latest_exclusion_run(conn, inc) or inc
            final = {dict(r)["url"] for r in conn.execute(
                f"SELECT url FROM evaluation_results WHERE run_id = {shortlist._P} AND keep = 1 "
                f"AND url IN ({ph})", tuple([run_id, *urls])).fetchall()}
        results = []
        for u, kind in tagged:
            row = info.get(u)
            results.append({
                "url": u, "kind": kind,
                "corpus": row is not None,
                "active": bool(row and not row.get("is_redirect") and not row.get("withdrawn")
                               and row.get("content_hash")),
                "final": u in final,
            })
        return JSONResponse({"results": results, "has_run": bool(inc)})
    finally:
        conn.close()


@app.post("/categories/{cid}/url-check/save")
async def save_url_checklist(request: Request, cid: int):
    """Store this category's URL-check lists: URLs that should be included in the final
    shortlist, and URLs that should be excluded (kept for later checks)."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    body = await request.json()
    include = str(body.get("should_include_urls") or "")[:20000]
    exclude = str(body.get("should_exclude_urls") or "")[:20000]
    conn = connect()
    try:
        if not cat.get_category(conn, cid):
            return JSONResponse({"error": "not found"}, status_code=404)
        cat.set_url_checklist(conn, cid, include, exclude)
        return JSONResponse({"ok": True})
    finally:
        conn.close()


@app.post("/feedback")
async def submit_feedback(request: Request):
    """Store feedback left from the per-page widget: the page id + title, a comment, and the
    star questions. Records who left it (accounts mode) and when."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    body = await request.json()
    text = str(body.get("feedback_text") or "").strip()
    stars = {f: body.get(f) for f in feedback._STAR_FIELDS}
    if not text and not any(feedback._star(v) for v in stars.values()):
        return JSONResponse({"error": "Please give at least one rating or a comment."},
                            status_code=400)
    u = current_user(request)
    conn = connect()
    try:
        feedback.add_feedback(
            conn, page_id=str(body.get("page_id") or ""),
            page_title=str(body.get("page_title") or ""), feedback_text=text, stars=stars,
            created_by=(u.get("id") if u else None),
            created_by_email=(u.get("email") if u else None))
        return JSONResponse({"ok": True})
    finally:
        conn.close()


# ---- GOV.UK search (find gov.uk pages related to a set of phrases) --------
_GOVUK_SEARCH_URL = "https://www.gov.uk/api/search.json"
_GOVUK_SEARCH_FIELDS = ("title", "link", "description", "content_store_document_type")
_GOVUK_UA = {"User-Agent": "gov-uk-corpus-shortlist-builder", "Accept": "application/json"}
# Guardrails on a Compare run (env-configurable). GOV.UK caps rows at 100 per request, so
# per-phrase paging is always in 100s regardless; that hard limit isn't ours to change. It
# also can't paginate past from+size = 10000 (its deep-pagination ceiling), so 10000 is the
# most retrievable per phrase — we default to it, favouring completeness over speed.
_GOVUK_MAX_PHRASES = int(os.getenv("GOVUK_MAX_PHRASES", "25"))     # phrase lines searched per run
_GOVUK_PER_PHRASE = int(os.getenv("GOVUK_PER_PHRASE", "10000"))    # pages listed per phrase (GOV.UK ceiling)
_GOVUK_DOCTYPE_AGG = int(os.getenv("GOVUK_DOCTYPE_AGG", "200"))    # doc-type facet buckets (complete set)
# Relevance floor: GOV.UK-Search-only pages below this es_score are NOT fed to the AI (they're
# likely false positives). A page with no es_score is kept (unknown, not low). 0 disables it.
_GOVUK_MIN_ES_SCORE = float(os.getenv("GOVUK_MIN_ES_SCORE", "0.005"))


def _govuk_get(params) -> tuple:
    url = _GOVUK_SEARCH_URL + "?" + urlencode(params)
    req = urllib.request.Request(url, headers=_GOVUK_UA)
    with urllib.request.urlopen(req, timeout=15) as r:
        return url, json.loads(r.read().decode("utf-8"))


def _phrase_q(phrase: str) -> str:
    """Quote a line so the GOV.UK Search API matches it as an exact phrase (e.g. "SPS
    agreement" -> 162 pages, not 76,000 for the loose terms)."""
    return '"' + phrase.strip().strip('"') + '"'


def _govuk_filter_params(organisations, document_types=()) -> list:
    """GOV.UK Search API filter params (repeatable): organisations + document types. Corpus
    org/doc-type slugs are the gov.uk slugs (content_store_document_type), so they pass through.
    An empty list means no filter on that facet."""
    return ([("filter_organisations", s) for s in (organisations or [])]
            + [("filter_content_store_document_type", s) for s in (document_types or [])])


def _govuk_doctypes(phrase: str, organisations=(), document_types=()) -> tuple:
    """The complete document-type distribution for a phrase (via the API's aggregate), as
    {slug: count}, plus the total number of matching pages. Optionally org/doc-type scoped."""
    url, data = _govuk_get([("q", _phrase_q(phrase)), ("count", "0"),
                            ("aggregate_content_store_document_type", str(_GOVUK_DOCTYPE_AGG))]
                           + _govuk_filter_params(organisations, document_types))
    opts = data.get("aggregates", {}).get("content_store_document_type", {}).get("options", [])
    out = {}
    for o in opts:
        v = o.get("value")
        slug = (v.get("slug") if isinstance(v, dict) else v) or ""
        if slug:
            out[slug] = o.get("documents", 0)
    return out, int(data.get("total") or 0), url


def _govuk_search_multi(phrases, organisations=(), document_types=(), progress=None, should_stop=None) -> dict:
    """Search GOV.UK for several phrases (one per line). Returns the combined page list
    (title, link, document type, which phrases matched), the full document-type set with
    counts, per-phrase totals, and the query URLs used. Optionally scoped to org + doc-type slugs."""
    filter_params = _govuk_filter_params(organisations, document_types)
    # Human-readable scope for the annotated query list (expert box / query rolldown).
    scope = (f"{len(organisations)} organisation(s)" if organisations else "all organisations")
    scope += (f", {len(document_types)} document type(s)" if document_types else ", all document types")
    pages, doctypes, per_phrase, query_urls = {}, {}, [], []
    phrase_list = phrases[:_GOVUK_MAX_PHRASES]
    for idx, phrase in enumerate(phrase_list):
        if should_stop and should_stop():
            break
        dt, total, agg_url = _govuk_doctypes(phrase, organisations, document_types)
        query_urls.append({"comment": f'Phrase "{phrase}": total matches + the full document-type '
                                      f'breakdown ({scope}); exact-phrase match, no rows fetched.',
                           "url": agg_url})
        for slug, n in dt.items():
            doctypes[slug] = doctypes.get(slug, 0) + n
        per_phrase.append({"phrase": phrase, "total": total})
        # GOV.UK can't paginate past from+size = 10000, so never request beyond that,
        # even if GOVUK_PER_PHRASE is set higher.
        cap = min(_GOVUK_PER_PHRASE, 10000)
        start = 0
        while start < cap:
            if should_stop and should_stop():
                break
            count = min(100, cap - start)
            url, data = _govuk_get([("q", _phrase_q(phrase)), ("count", str(count)), ("start", str(start))]
                                   + [("fields", f) for f in _GOVUK_SEARCH_FIELDS] + filter_params)
            query_urls.append({"comment": f'Phrase "{phrase}": fetch matching pages '
                                          f'(rows {start + 1}–{start + count}, {scope}).',
                               "url": url})
            res = data.get("results", [])
            if not res:
                break
            for item in res:
                link = str(item.get("link") or "")
                if link.startswith("/"):
                    link = "https://www.gov.uk" + link
                try:
                    es = float(item.get("es_score")) if item.get("es_score") is not None else None
                except (TypeError, ValueError):
                    es = None
                if link in pages:
                    if phrase not in pages[link]["phrases"]:
                        pages[link]["phrases"].append(phrase)
                    if es is not None:            # keep the best relevance across phrases/pages
                        prev = pages[link].get("es_score")
                        pages[link]["es_score"] = es if prev is None else max(prev, es)
                    continue
                pages[link] = {"title": (item.get("title") or link).strip(), "link": link,
                               "document_type": item.get("content_store_document_type") or "",
                               "description": (item.get("description") or "").strip(),
                               "es_score": es,
                               "phrases": [phrase]}
            start += len(res)
            if start >= total:
                break
        if progress:
            progress(idx + 1, len(phrase_list), len(pages))
    doctype_list = sorted(({"type": t, "count": c} for t, c in doctypes.items()),
                          key=lambda x: (-x["count"], x["type"]))
    # Threshold check: did any configured cap truncate the GOV.UK result set? If so, the
    # "GOV.UK Search only" gap may be a limit artefact rather than a true search difference.
    per_cap = min(_GOVUK_PER_PHRASE, 10000)
    capped_phrases = [{"phrase": p["phrase"], "total": p["total"]}
                      for p in per_phrase if (p.get("total") or 0) > per_cap]
    phrases_truncated = max(0, len(phrases) - len(phrase_list))
    thresholds = {
        "max_phrases": _GOVUK_MAX_PHRASES,
        "per_phrase_cap": per_cap,
        "govuk_ceiling": 10000,
        "phrases_submitted": len(phrases),
        "phrases_searched": len(per_phrase),
        "phrases_truncated": phrases_truncated,      # beyond GOVUK_MAX_PHRASES, never searched
        "capped_phrases": capped_phrases,            # had more hits than we fetched per phrase
        "any_hit": bool(phrases_truncated or capped_phrases),
    }
    return {"results": list(pages.values()), "doctypes": doctype_list,
            "per_phrase": per_phrase, "query_urls": query_urls,
            "per_phrase_cap": _GOVUK_PER_PHRASE, "thresholds": thresholds}


@app.get("/govuk-search", response_class=HTMLResponse)
def govuk_search_page(request: Request):
    """Search GOV.UK for pages matching a set of phrases (guc-0019). Shown only at the
    Admin UI-complexity level (profile setting) — see the data-level="admin" nav link.
    Uses the official Search API."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    resp = templates.TemplateResponse("govuk_search.html", ctx(conn, request, active_nav="govuk_search"))
    conn.close()
    return resp


@app.get("/sustainability", response_class=HTMLResponse)
def sustainability_page(request: Request):
    """Modelled sustainability dashboard for the AI runs (guc-0022): energy / water / CO₂."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    try:
        data = sustainability.summary(conn)
        resp = templates.TemplateResponse("sustainability.html", ctx(
            conn, request, active_nav="sustainability", s=data))
    finally:
        conn.close()
    return resp


@app.post("/api/govuk-search")
async def api_govuk_search(request: Request):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    body = await request.json()
    raw = body.get("phrases")
    lines = raw.splitlines() if isinstance(raw, str) else (raw if isinstance(raw, list) else [])
    phrases, seen = [], set()
    for p in lines:
        p = str(p or "").strip()
        if p and p.lower() not in seen:
            seen.add(p.lower())
            phrases.append(p)
    if not phrases:
        return JSONResponse({"results": [], "doctypes": [], "per_phrase": [], "query_urls": []})
    try:
        data = await run_in_threadpool(_govuk_search_multi, phrases)
    except Exception as e:
        logging.getLogger("govuk_search").warning("search failed: %s", e)
        return JSONResponse({"error": "Couldn't reach GOV.UK search — try again."}, status_code=502)
    return JSONResponse(data)


def _doc_type_counts(conn, org_slugs) -> list:
    """Corpus document types (effective type) with page counts, restricted to the given
    organisations, newest-largest first. Empty org list = the whole corpus.

    For an org filter we start from page_organisations (the DISTINCT matching page_urls off
    the org index) and join content, rather than an EXISTS per content row — ~7x faster for
    broad orgs with child departments (e.g. DEFRA+children: ~0.9s vs ~6.5s), so the Page
    Types picker on the edit form loads its full list promptly instead of appearing stuck."""
    if org_slugs:
        ph = ",".join([shortlist._P] * len(org_slugs))
        eff = shortlist.EFFECTIVE_DOCTYPE_EXPR      # html_publication -> parent type
        sql = (f"SELECT {eff} AS dt, COUNT(*) AS n FROM "
               f"(SELECT DISTINCT po.page_url FROM page_organisations po "
               f" WHERE po.organisation_slug IN ({ph})) pp "
               f"JOIN content c ON c.url = pp.page_url "
               f"WHERE c.is_redirect = 0 AND c.content_hash IS NOT NULL "
               f"GROUP BY {eff} ORDER BY n DESC")
        params = list(org_slugs)
    else:
        # Whole corpus: raw type keeps the grouping fast (no org join).
        sql = ("SELECT c.document_type AS dt, COUNT(*) AS n FROM content c "
               "WHERE c.is_redirect = 0 AND c.content_hash IS NOT NULL "
               "GROUP BY c.document_type ORDER BY n DESC")
        params = []
    rows = conn.execute(sql, tuple(params)).fetchall()
    return [{"type": r["dt"], "count": r["n"]} for r in rows if r["dt"]]


@app.get("/api/doc-type-counts")
def api_doc_type_counts(request: Request, orgs: str = "", children: str = "0"):
    """Document types (with counts) available in the corpus for the given organisation slugs
    — feeds the Page Types multi-select on the category form."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    slugs = cat.parse_list(orgs)   # `orgs` here is the CSV query param, not the orgs module
    want_children = str(children) in ("1", "true", "on")
    key = (tuple(sorted(slugs)), want_children)
    if _COUNT_TTL > 0:
        with _DOCTYPE_LOCK:
            hit = _DOCTYPE_CACHE.get(key)
            if hit and hit[1] > time.time():
                return JSONResponse({"types": hit[0]})
    conn = connect()
    try:
        eff_slugs = orgs_mod.expand_with_children(conn, slugs) if (slugs and want_children) else slugs
        types = _doc_type_counts(conn, eff_slugs)
    finally:
        conn.close()
    if _COUNT_TTL > 0:
        with _DOCTYPE_LOCK:
            _DOCTYPE_CACHE[key] = (types, time.time() + _COUNT_TTL)
    return JSONResponse({"types": types})


@app.post("/api/org-options")
async def api_org_options(request: Request):
    """All organisations for the picker, plus the slugs matched from a brain-dump (via
    orgs.search) so the category assistant can pre-tick them — no LLM needed."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    body = await request.json()
    text = str(body.get("text") or "")
    conn = connect()
    try:
        options = orgs.all_orgs(conn)
        opt_slugs = {o["slug"] for o in options}
        matched = []
        if text.strip():
            matched = [m["slug"] for m in orgs.search(conn, text, limit=30) if m["slug"] in opt_slugs]
        return JSONResponse({"options": options, "matched": matched})
    finally:
        conn.close()


@app.get("/api/page-links")
def api_page_links(request: Request, url: str = ""):
    """All hyperlinks in a page's body, from the stored corpus content — for the per-row
    "links off this page" roll-down on the audit shortlist."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    url = (url or "").strip()
    if not url:
        return JSONResponse({"links": []})
    conn = connect()
    try:
        row = conn.execute(f"SELECT content FROM content WHERE url = {shortlist._P}", (url,)).fetchone()
    finally:
        conn.close()
    raw = row["content"] if row else None
    if not raw:
        return JSONResponse({"links": [], "no_content": True})
    try:
        payload = json.loads(raw) if isinstance(raw, str) else raw   # sqlite TEXT vs PG JSONB
        links = extract.page_body_links(payload)
    except Exception:
        return JSONResponse({"links": [], "no_content": True})
    return JSONResponse({"links": links})


@app.get("/api/categories/{cid}/runs")
def api_list_runs(request: Request, cid: int):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    runs = evaluate.list_runs(conn, cid)
    category = cat.get_category(conn, cid)
    # Annotate each run with the total it evaluates toward, for the "X / N" display:
    # inclusion runs work through the whole shortlist; exclusion runs only re-check
    # the pages their source inclusion run kept.
    _shortlist_total = [None]   # memoised (one count per request, only if needed)

    def shortlist_total():
        if _shortlist_total[0] is None and category:
            try:
                _shortlist_total[0] = cached_count(conn, **_effective_filters(conn, category))
            except Exception:
                _shortlist_total[0] = None
        return _shortlist_total[0]

    # Per-run count of evaluated pages that came from the GOV.UK-Search-only top-up (source='search'),
    # so an inclusion run's Phase-1 pages can be split into corpus-keyword vs GOV.UK-sourced.
    P = shortlist._P
    govuk_by_run = {}
    for row in conn.execute(
            f"SELECT er.run_id AS run_id, COUNT(*) AS n FROM evaluation_results er "
            f"WHERE er.category_id = {P} AND er.url IN "
            f"(SELECT url FROM category_search_pages WHERE category_id = {P} AND source = 'search') "
            f"GROUP BY er.run_id", (cid, cid)).fetchall():
        d = dict(row)
        govuk_by_run[d["run_id"]] = d.get("n") or 0

    by_id = {r["run_id"]: r for r in runs}
    stalled = None
    for r in runs:
        if "exclusion" in (r.get("phase") or "").lower():
            src = by_id.get(r.get("source_run_id"))
            r["target"] = src.get("kept") if src else None   # exclusion re-checks the inclusion's keeps
        else:
            # A scoped (benchmark) run works toward its fixed url list, not the shortlist.
            scope_n = (r.get("trial") or {}).get("scope_n")
            r["target"] = scope_n if scope_n else shortlist_total()
            g = govuk_by_run.get(r["run_id"], 0)              # GOV.UK-Search-only pages this run evaluated
            r["src_govuk"] = g
            r["src_corpus"] = max(0, (r.get("pages") or 0) - g)   # the rest = corpus keyword shortlist
        # Live driver state for the UI: is a process actually working this run, and is it a
        # stalled run (marked running but its driver is gone) that could be resumed?
        r["alive"] = evaluate.run_is_alive(r)
        r["stalled"] = evaluate.run_is_stalled(r)
        if r["stalled"] and stalled is None:
            stalled = {"run_id": r["run_id"], "phase": r.get("phase"), "name": r.get("name")}
    rsql, rparams = evaluate.list_runs_query(cid)
    out = {"runs": runs, "active": _active_run(conn, cid), "stalled": stalled,
           "any_running": any(r.get("alive") for r in runs), "sql": _display_sql(rsql, rparams)}
    conn.close()
    return JSONResponse(out)


@app.get("/api/categories/{cid}/active-run")
def api_active_run_summary(request: Request, cid: int):
    """Summary of the active run for the Active Run tab: its phase chain
    (inclusion → exclusion) plus a status of complete / incomplete / in_progress."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    try:
        run_id = _active_run(conn, cid)
        run = evaluate.get_run(conn, run_id) if run_id else None
        if not run:
            return JSONResponse({"run": None})
        chain = evaluate.run_chain(conn, run_id)
        category = cat.get_category(conn, cid)
        try:
            shortlist_total = cached_count(conn, **_effective_filters(conn, category)) if category else None
        except Exception:
            shortlist_total = None
        continue_reason = evaluate.continuable_reason(chain, shortlist_total)
        chain_ids = {r["run_id"] for r in chain}
        ent = _bg_evals.get(cid)
        in_progress = bool(ent and ent["status"].get("running")
                           and ent["status"].get("run_id") in chain_ids)
        status = "in_progress" if in_progress else ("complete" if not continue_reason else "incomplete")
        rows = [{"phase": r.get("phase"), "provider": r.get("provider"),
                 "model": r.get("actual_model") or r.get("model"),
                 "pages": r.get("pages") or 0, "kept": r.get("kept") or 0,
                 "dropped": r.get("dropped") or 0, "in_tokens": r.get("in_tokens") or 0,
                 "out_tokens": r.get("out_tokens") or 0, "cost": r.get("cost") or 0}
                for r in chain]
        return JSONResponse({"run": {"run_id": run["run_id"], "name": run.get("name")},
                             "status": status, "chain": rows})
    finally:
        conn.close()


@app.get("/categories/{cid}/runs/{run_id}", response_class=HTMLResponse)
def run_detail_page(request: Request, cid: int, run_id: str):
    """Per-run detail: the LLM phase chain (inclusion → exclusion), model used,
    tokens and cost per phase, with rename."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    run = evaluate.get_run(conn, run_id)
    if not category or not run or str(run["category_id"]) != str(cid):
        conn.close()
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    category["display_name"] = cat.display_name(category)
    chain = evaluate.run_chain(conn, run_id)
    totals = {k: sum((r.get(k) or 0) for r in chain)
              for k in ("cost", "in_tokens", "out_tokens", "hit_tokens", "miss_tokens", "pages")}
    try:
        shortlist_total = cached_count(conn, **_effective_filters(conn, category))
    except Exception:
        shortlist_total = None
    commentary = evaluate.run_commentary(chain, shortlist_total, opened_run_id=run_id)
    continue_reason = evaluate.continuable_reason(chain, shortlist_total)
    unparsed = evaluate.unparsed_results(conn, [r["run_id"] for r in chain])
    phase_by_id = {r["run_id"]: r.get("phase") for r in chain}
    for u in unparsed:
        u["phase"] = phase_by_id.get(u["run_id"])
    # The settings that actually controlled each phase of this run (from its stamp).
    run_configs = []
    for rr in chain:
        spec = evaluate.run_prompt_spec(conn, rr["run_id"]) or {}
        tr = _norm_trial(spec)
        run_configs.append({
            "phase": rr.get("phase"), "provider": rr.get("provider"),
            "model": rr.get("actual_model") or rr.get("model"),
            "prompt_variant": tr["prompt_variant"], "concurrency": tr["concurrency"],
            "caching": tr["caching"], "template_version": spec.get("template_version"),
            # Prompt identity: the stamped fingerprint, or one computed from the stored template
            # for runs stamped before the field existed (compute-on-read — no backfill needed).
            "template_hash": spec.get("template_hash") or evaluate.template_fingerprint(spec.get("template")),
            "body_limit": spec.get("body_limit"), "stamped": bool(spec)})
    # Run lifecycle state for the action button: fresh (never run) | partial (started, unfinished) | complete.
    started = (totals.get("pages") or 0) > 0
    run_state = "complete" if not continue_reason else ("partial" if started else "fresh")
    resp = templates.TemplateResponse("run_detail.html", ctx(
        conn, request, category=category, run=run, chain=chain, totals=totals,
        commentary=commentary, unparsed=unparsed, continue_reason=continue_reason,
        run_trial=_run_trial(conn, run_id), run_configs=run_configs, run_state=run_state))
    conn.close()
    return resp


def _example_phase_prompts(conn, category, run_id: str = "") -> dict:
    """Each phase's prompt for this run, rebuilt for one sample page. When a phase's run is
    stamped (evaluation_runs.prompt_spec) the rebuild is EXACT — its own template, version and
    Include/Exclude context — so only the page body reflects the current corpus. Legacy phases,
    or a phase not yet run, fall back to the current active template (clearly flagged)."""
    cid = category["id"]
    P = shortlist._P
    chain = evaluate.run_chain(conn, run_id) if run_id else []
    incl_run = next((r for r in chain if "inclusion" in (r.get("phase") or "").lower()), None)
    excl_run = next((r for r in chain if "exclusion" in (r.get("phase") or "").lower()), None)

    # Sample page: one THIS run evaluated (a kept page first), else the shortlist's first page.
    url = title = description = body = ""
    sample_from_run = False
    if run_id:
        row = conn.execute(
            f"SELECT c.url AS url, c.title AS title, c.description AS description, "
            f"c.search_text AS body FROM evaluation_results er JOIN content c ON c.url = er.url "
            f"WHERE er.run_id = {P} AND er.keep IS NOT NULL "
            f"ORDER BY (CASE WHEN er.keep = 1 THEN 0 ELSE 1 END), c.url LIMIT 1",
            (run_id,)).fetchone()
        if row:
            d = dict(row)
            url, title = d.get("url") or "", d.get("title") or ""
            description, body = d.get("description") or "", d.get("body") or ""
            sample_from_run = bool(url)
    if not url:
        f = _effective_filters(conn, category)
        sql, params = shortlist.build_query(
            select_expr="c.url AS url, c.title AS title, c.description AS description, c.search_text AS body",
            organisations=f["organisations"], document_types=f["document_types"],
            keywords=f["keywords"], match=f["match"], limit=1)
        row = conn.execute(sql, tuple(params)).fetchone()
        if row:
            d = dict(row)
            url, title = d.get("url") or "", d.get("title") or ""
            description, body = d.get("description") or "", d.get("body") or ""
    if not url:
        title = "(example page title)"
        description, body = "(example page description)", "(the page's body text goes here)"

    # PASS1_REASON / PASS1_TOPIC for the exclusion prompt: the inclusion run's stored note and
    # primary topic for that page.
    pass1 = "(the first pass's 1-2 sentence note for this page)"
    pass1_topic = "(the first pass's primary topic for this page)"
    if url:
        ir = incl_run["run_id"] if incl_run else evaluate.latest_inclusion_run(conn, cid)
        if ir:
            r = conn.execute(f"SELECT reason, primary_topic FROM evaluation_results WHERE run_id = {P} AND url = {P}",
                             (ir, url)).fetchone()
            r = dict(r) if r else {}
            if (r.get("reason") or "").strip():
                pass1 = r["reason"]
            if (r.get("primary_topic") or "").strip():
                pass1_topic = r["primary_topic"]

    # The raw template each phase used (from the run's stamp, else the current active version) —
    # shown alongside the completed prompt for comparison.
    incl_spec = evaluate.run_prompt_spec(conn, incl_run["run_id"]) if incl_run else None
    incl_template = (incl_spec or {}).get("template") or prompts.default_text("inclusion")
    if incl_spec:
        inclusion_prompt = evaluate.build_prompt(
            incl_spec.get("inclusion_context", ""), incl_spec.get("exclusion_context", ""),
            title, description, body, body_limit=incl_spec.get("body_limit", evaluate.BODY_CHAR_LIMIT),
            template=incl_template)
    else:
        inclusion_prompt = evaluate.build_prompt(
            category.get("inclusion_context") or "", category.get("exclusion_context") or "",
            title, description, body, template=incl_template)

    excl_spec = evaluate.run_prompt_spec(conn, excl_run["run_id"]) if excl_run else None
    excl_template = (excl_spec or {}).get("template") or prompts.default_text("exclusion")
    if excl_spec:
        exclusion_prompt = evaluate.build_exclusion_prompt(
            excl_spec.get("name") or "the topic", excl_spec.get("inclusion_context", ""),
            excl_spec.get("exclusion_context", ""), excl_spec.get("keep_hints", ""),
            excl_spec.get("drop_hints", ""), title, body, pass1,
            body_limit=excl_spec.get("body_limit", evaluate.BODY_CHAR_LIMIT), template=excl_template,
            pass1_topic=pass1_topic)
    else:
        exclusion_prompt = evaluate.build_exclusion_prompt(
            (category.get("description") or "").strip() or cat.prettify(category.get("slug")) or "the topic",
            category.get("inclusion_context") or "", category.get("exclusion_context") or "",
            category.get("adjudication_hints_keep") or "", category.get("adjudication_hints_drop") or "",
            title, body, pass1, template=excl_template, pass1_topic=pass1_topic)

    return {
        "sample_url": url,
        "sample_title": title or url,
        "sample_from_run": sample_from_run,
        "inclusion_template": incl_template,
        "inclusion": inclusion_prompt,
        "inclusion_exact": bool(incl_spec),
        "inclusion_version": (incl_spec or {}).get("template_version"),
        "exclusion_template": excl_template,
        "exclusion": exclusion_prompt,
        "exclusion_exact": bool(excl_spec),
        "exclusion_version": (excl_spec or {}).get("template_version"),
    }


@app.get("/api/categories/{cid}/example-prompts")
def api_example_prompts(request: Request, cid: int, run_id: str = ""):
    """The reconstructed inclusion + exclusion prompts for a page this run evaluated —
    fed to the advanced-only rolldown on the run-detail page (guc-0006)."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    try:
        category = cat.get_category(conn, cid)
        if not category:
            return JSONResponse({"error": "not found"}, status_code=404)
        return JSONResponse(_example_phase_prompts(conn, category, run_id))
    finally:
        conn.close()


@app.get("/api/categories/{cid}/runs/{run_id}/results")
def api_run_results(request: Request, cid: int, run_id: str, keep: Optional[int] = None,
                    limit: int = 2000):
    """Per-page results for a run (guc-0006): url, verdict, score + confidence label, reason, and
    the inclusion pass's grounding fields (primary_topic, where_hit, evidence). On an exclusion run
    the grounding fields come from its inclusion run — only Phase 1 writes them."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    try:
        run = evaluate.get_run(conn, run_id)
        if not run or str(run.get("category_id")) != str(cid):
            return JSONResponse({"error": "not found"}, status_code=404)
        rows = evaluate.run_results(conn, run_id, keep=keep, limit=max(1, min(int(limit), 5000)),
                                    source_run_id=run.get("source_run_id"))
        for r in rows:
            r["confidence"] = evaluate.confidence_label(r.get("score"))
            r["where_hit"] = evaluate.json_list_display(r.get("where_hit"))
            r["evidence"] = evaluate.json_list_display(r.get("evidence"))
        return JSONResponse({"run_id": run_id, "phase": run.get("phase"), "rows": rows, "total": len(rows)})
    finally:
        conn.close()


@app.post("/api/categories/{cid}/runs/{run_id}/activate")
async def api_activate_run(request: Request, cid: int, run_id: str):
    """Make this run the active run for the category — Evaluate / background runs
    then add to it."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    try:
        run = evaluate.get_run(conn, run_id)
        if not run or str(run["category_id"]) != str(cid):
            return JSONResponse({"error": "not found"}, status_code=404)
        settings.set_setting(conn, f"active_run_{cid}", run_id)
        return JSONResponse({"ok": True, "active": run_id})
    finally:
        conn.close()


@app.post("/api/categories/{cid}/runs/{run_id}/continue")
async def api_continue_run(request: Request, cid: int, run_id: str):
    """Make the phase that still needs work the active run so a background run
    finishes the LLM evaluation: resume an in-progress / short phase, or start the
    pending Phase-2 exclusion over the inclusion run's keeps. Returns the run that
    is now active, or {done: true} if there is nothing left to do."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    try:
        category = cat.get_category(conn, cid)
        run = evaluate.get_run(conn, run_id)
        if not category or not run or str(run["category_id"]) != str(cid):
            return JSONResponse({"error": "not found"}, status_code=404)
        chain = evaluate.run_chain(conn, run_id)
        try:
            shortlist_total = cached_count(conn, **_effective_filters(conn, category))
        except Exception:
            shortlist_total = None
        by_id = {r["run_id"]: r for r in chain}

        def target_of(r):
            if "exclusion" in (r.get("phase") or "").lower():
                src = by_id.get(r.get("source_run_id"))
                return src.get("kept") if src else None
            sc = evaluate.run_scope(conn, r["run_id"])
            return len(sc["urls"]) if sc else shortlist_total

        # 1) an unfinished phase, or a finished phase that stopped below its input.
        for r in chain:
            if not r.get("finished_at"):
                settings.set_setting(conn, f"active_run_{cid}", r["run_id"])
                return JSONResponse({"active": r["run_id"], "action": "resume"})
        for r in chain:
            tgt = target_of(r)
            if tgt is not None and (r.get("pages") or 0) < tgt:
                settings.set_setting(conn, f"active_run_{cid}", r["run_id"])
                return JSONResponse({"active": r["run_id"], "action": "resume"})

        # 2) inclusion finished with keeps but no exclusion yet → start Phase 2.
        have_excl = any("exclusion" in (r.get("phase") or "").lower() for r in chain)
        incl = next((r for r in chain if "inclusion" in (r.get("phase") or "").lower()), None)
        if incl and incl.get("finished_at") and (incl.get("kept") or 0) > 0 and not have_excl:
            spec = evaluate.run_prompt_spec(conn, incl["run_id"]) or {}
            if spec.get("bench"):
                return JSONResponse({"done": True, "note": "benchmark run: Phase 2 is driven by bench, not the app"})
            scope = evaluate.run_scope(conn, incl["run_id"])
            excfg = ({"model": incl["model"], "provider": incl["provider"]} if scope
                     else _ai_config_for_phase(conn, evaluate.PHASE_EXCLUSION))
            new_id = evaluate.create_run(conn, cid, excfg["model"], excfg["provider"],
                                         phase=evaluate.PHASE_EXCLUSION, source_run_id=incl["run_id"],
                                         name=incl.get("name"),   # a run's name spans both phases
                                         prompt_spec=_prompt_spec_json(conn, cid, evaluate.PHASE_EXCLUSION,
                                                                       _run_trial(conn, incl["run_id"]),
                                                                       scope=scope))
            settings.set_setting(conn, f"active_run_{cid}", new_id)
            return JSONResponse({"active": new_id, "action": "start_exclusion"})

        return JSONResponse({"done": True})
    finally:
        conn.close()


@app.post("/api/categories/{cid}/runs/{run_id}/rename")
async def api_rename_run(request: Request, cid: int, run_id: str):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    form = await request.form()
    name = (form.get("name") or "").strip()
    if not name:
        return JSONResponse({"error": "empty name"}, status_code=400)
    conn = connect()
    try:
        run = evaluate.get_run(conn, run_id)
        if not run or str(run["category_id"]) != str(cid):
            return JSONResponse({"error": "not found"}, status_code=404)
        evaluate.rename_run(conn, run_id, name[:120])
        return JSONResponse({"ok": True, "name": name[:120]})
    finally:
        conn.close()


@app.post("/api/categories/{cid}/runs/{run_id}/delete")
def api_delete_run(request: Request, cid: int, run_id: str):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    try:
        run = evaluate.get_run(conn, run_id)
        if not run or str(run["category_id"]) != str(cid):
            return JSONResponse({"error": "not found"}, status_code=404)
        evaluate.delete_run(conn, run_id)
        # If this was the category's active run, forget it so a fresh one starts next time.
        if _active_run(conn, cid) == run_id:
            settings.set_setting(conn, f"active_run_{cid}", "")
        return JSONResponse({"ok": True})
    finally:
        conn.close()


@app.get("/api/categories/{cid}/compare")
def api_compare(request: Request, cid: int, base: str = "", other: str = ""):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    out = evaluate.compare(conn, base, other) if base and other else {}
    conn.close()
    return JSONResponse(out)


@app.get("/api/categories/{cid}/final-diff")
def api_final_diff(request: Request, cid: int, base: str = "", other: str = ""):
    """Variability in the FINAL shortlist between two run-chains (base=Run 1, other=Run 2), over
    pages both evaluated. `base`/`other` are the inclusion run ids; the chains' exclusion phases
    are resolved here. Returns the total that differ plus the two directions."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    try:
        b_incl, b_excl = _chain_of(conn, base)
        o_incl, o_excl = _chain_of(conn, other)
        diffs = evaluate.final_outcome_diffs(
            conn, b_incl and b_incl["run_id"], b_excl and b_excl["run_id"],
            o_incl and o_incl["run_id"], o_excl and o_excl["run_id"])
        base_only = sum(1 for d in diffs if d["base_final"] and not d["other_final"])
        other_only = sum(1 for d in diffs if not d["base_final"] and d["other_final"])
        return JSONResponse({"count": len(diffs), "base_only": base_only, "other_only": other_only})
    finally:
        conn.close()


def _dl_stamp() -> str:
    """UTC timestamp for download filenames: yy-mm-dd-hr-min."""
    return time.strftime("%y-%m-%d-%H-%M", time.gmtime())


def _safe_filename(name: str, default: str) -> str:
    """A safe download base filename from user input: drop any extension they typed,
    keep letters/digits/space/dot/dash/underscore, spaces -> dashes, cap the length."""
    base = re.sub(r"\.(csv|xlsx|json)$", "", (name or "").strip(), flags=re.I)
    base = re.sub(r"[^A-Za-z0-9._ -]", "", base).strip().replace(" ", "-")
    return base[:120] or default


# ---- download page (choose format + fields) -----------------------------
# Grouped into sections for the download page. Each field: (key, label, default_on,
# disabled, warning). 'url' is mandatory (disabled + on).
# Category/run-scoped columns that aren't plain content fields — added to each row in Python
# (see _enrich_audit_rows), not fetched by the shortlist SQL. key -> label.
_VIRTUAL_FIELDS = {
    "run_id": "Run ID",
    "matched_keywords": "Matched keywords",
    "inclusion_reason": "Inclusion reason",
    "exclusion_reason": "Exclusion reason",
    "confidence": "Confidence",
    "primary_topic": "Primary topic",
    "where_hit": "Matched in",
    "evidence": "Evidence quotes",
    "inclusion_raw_reply": "Inclusion raw reply",
    "exclusion_raw_reply": "Exclusion raw reply",
    "stage": "Stage",
    "decision": "Stage decision",
    "funnel_stage": "Furthest stage reached",
}

_DOWNLOAD_SECTIONS = [
    ("Content", [
        ("url", "URL", True, True, None),
        ("title", "Title", True, False, None),
        ("parent_document_type", "Parent document type", False, False,
         "for html_publication pages: the parent publication's type"),
        ("search_text", "Body text (search_text)", False, False,
         "warning: the full page text — can make the download very large"),
        ("content", "Content (raw JSON)", False, False, "warning: may make the download very large"),
    ]),
    ("Provenance", [
        ("source", "Source / provenance", False, False, None),
    ]),
    ("Popularity", [
        ("view_count", "View count", True, False,
         "GOV.UK Search pageviews (~last 14 days); for attachment / guide-part sub-pages this is "
         "the parent publication's count"),
        ("view_count_source", "View count source", False, False,
         "page = the page's own count · parent = inherited from the parent publication"),
        ("view_count_updated", "View count collected", False, False,
         "the date the view count was collected"),
    ]),
    ("Matching & AI decision", [
        ("run_id", "Run ID", True, True,
         "the inclusion run the AI columns come from — always included, so rows from several runs "
         "can be told apart in one file"),
        ("organisations", "Organisations", True, False, "all linked organisation slugs"),
        ("document_type", "Document type", True, False, None),
        ("matched_keywords", "Matched keywords", True, False,
         "the shortlist keywords this page matched"),
        ("inclusion_reason", "Inclusion reason", True, False,
         "from the latest AI run; blank for pages not evaluated"),
        ("exclusion_reason", "Exclusion reason", True, False,
         "from the latest AI run's exclusion pass"),
        ("inclusion_raw_reply", "Inclusion raw reply", False, False,
         "warning: the model's verbatim inclusion reply — can be long"),
        ("exclusion_raw_reply", "Exclusion raw reply", False, False,
         "warning: the model's verbatim exclusion reply — can be long"),
        ("stage", "Stage", False, False,
         "the funnel stage / phase these rows are shown at (tags each row for the download)"),
        ("decision", "Stage decision", False, False,
         "the AI verdict for the page — Keep / Drop / Unscored (from the run relevant to the stage)"),
        ("funnel_stage", "Furthest stage reached", True, False,
         "how far each page got in the funnel — Keyword search / Reached LLM inclusion / Reached "
         "exclusion / Final shortlist — independent of which stage is selected above"),
        ("confidence", "Confidence", False, False,
         "how central the topic is to the page, from the inclusion score (Wrong sense / Mentioned in "
         "passing / Discussed a moderate amount / Major focus)"),
        ("primary_topic", "Primary topic", False, False,
         "from the latest inclusion run: what the page is mainly about (<=10-word noun phrase), "
         "grounded in its title/description"),
        ("where_hit", "Matched in", False, False,
         "from the latest inclusion run: which fields carried the topic — title, description, body"),
        ("evidence", "Evidence quotes", False, False,
         "from the latest inclusion run: the verbatim passages quoted as grounding for the decision"),
    ]),
    ("Ownership", [
        ("primary_org", "Primary publishing organisation", False, False, None),
    ]),
    ("Quality attributes", [
        ("size", "Size", True, False, None),
        ("readability", "Readability score", False, False, None),
        ("gds_issues", "GDS number of issues", False, False, None),
        ("gds_findings", "GDS issues text", False, False, None),
        ("gds_stars", "GDS stars", False, False, None),
    ]),
    ("Freshness", [
        ("first_published_at", "First published at", False, False, None),
        ("public_updated_at", "Public updated at", False, False, None),
        ("last_seen_at", "Last seen (crawl)", False, False, None),
    ]),
]


def _all_audit_fields() -> set:
    return set(shortlist.EXPORT_FIELDS) | set(_VIRTUAL_FIELDS)


def _field_label(k: str) -> str:
    return _VIRTUAL_FIELDS[k] if k in _VIRTUAL_FIELDS else shortlist.EXPORT_FIELDS[k][1]


def _enrich_audit_rows(conn, cid: int, rows: list, keys: list, stage: str = None,
                       run: str = None) -> None:
    """Add the category/run-scoped virtual columns (matched keywords, inclusion/exclusion
    reason, stage, decision) to each row in place, for whichever are in `keys`. Reasons come
    from the inclusion run and its exclusion run; pages not evaluated get ''. `run` (an inclusion
    run id) scopes those to that specific run-chain; None uses the category's latest/active run.
    `stage` is the audit stage KEY (drives the stage label and which run the decision reads)."""
    need = [k for k in keys if k in _VIRTUAL_FIELDS]
    if "stage" in need:                          # constant per view — tags each row with its stage
        stage_label = dict(_AUDIT_STAGES).get(stage, stage) or ""
        for r in rows:
            r["stage"] = stage_label
    if "run_id" in need:                         # constant per run — which inclusion run the AI columns come from
        rid = run or evaluate.latest_inclusion_run(conn, cid) or ""
        for r in rows:
            r["run_id"] = rid
    urls = [r.get("url") for r in rows if r.get("url")]
    if not need or not urls:
        for r in rows:                       # still populate empty cells so columns render
            for k in need:
                r.setdefault(k, "")
        return
    P = shortlist._P

    def _map(sql_head, id_val):
        out = {}
        for i in range(0, len(urls), 500):
            ch = urls[i:i + 500]
            ph = ",".join([P] * len(ch))
            for row in conn.execute(sql_head + f" AND url IN ({ph})", tuple([id_val] + ch)).fetchall():
                d = dict(row)
                out[d["url"]] = d.get("val")
        return out

    if "matched_keywords" in need:
        # Keyword hits are stored once per page, under its representative URL (MIN over all the
        # page's alias URLs). But a row here can carry a *different* alias — e.g. the AI kept one
        # chapter of a multi-URL guide, so the row's URL is that chapter, not the guide's base.
        # Joining on the exact URL then misses the stored hits. Resolve by content_id instead
        # (the same COALESCE(content_id, url) key membership dedupes on), so every alias of a
        # page inherits its hits.
        url2key = {}
        for i in range(0, len(urls), 500):
            ch = urls[i:i + 500]
            ph = ",".join([P] * len(ch))
            for row in conn.execute(
                    f"SELECT url, COALESCE(content_id, url) AS k FROM content WHERE url IN ({ph})",
                    tuple(ch)).fetchall():
                d = dict(row)
                url2key[d["url"]] = d["k"]
        keys = list({v for v in url2key.values()})
        key2kw = {}
        for i in range(0, len(keys), 500):
            ch = keys[i:i + 500]
            ph = ",".join([P] * len(ch))
            for row in conn.execute(
                    f"SELECT content_id AS k, matched_keywords AS val FROM category_shortlist_pages "
                    f"WHERE category_id = {P} AND content_id IN ({ph})", tuple([cid] + ch)).fetchall():
                d = dict(row)
                key2kw[d["k"]] = d.get("val")
        for r in rows:
            key = url2key.get(r.get("url"), r.get("url"))
            r["matched_keywords"] = ", ".join(search_augment._load_list(key2kw.get(key)))
    # The inclusion pass's grounding fields (only Phase 1 writes them) — always read from the
    # inclusion run, whatever the field name starts with.
    _GROUNDING = ("primary_topic", "where_hit", "evidence")
    _phase_cols = {"inclusion_reason": "reason", "exclusion_reason": "reason",
                   "inclusion_raw_reply": "raw_reply", "exclusion_raw_reply": "raw_reply",
                   "primary_topic": "primary_topic", "where_hit": "where_hit", "evidence": "evidence"}
    if any(k in need for k in _phase_cols):
        inc = run or evaluate.latest_inclusion_run(conn, cid)
        exc = evaluate.latest_exclusion_run(conn, inc) if inc else None
        for field, col in _phase_cols.items():
            if field not in need:
                continue
            run_id = inc if (field.startswith("inclusion") or field in _GROUNDING) else exc
            m = _map(f"SELECT url, {col} AS val FROM evaluation_results WHERE run_id = {P}", run_id) if run_id else {}
            for r in rows:
                val = m.get(r.get("url")) or ""
                # Prefix the exclusion reason with its exclusion_hit label header (single source
                # for the table + every export); other virtual fields pass through unchanged.
                r[field] = (evaluate.exclusion_display(val) if field == "exclusion_reason"
                            else evaluate.json_list_display(val) if field in ("where_hit", "evidence")
                            else val)
    if "decision" in need:
        # The AI verdict for each page — Keep / Drop / Unscored (a row with keep = NULL) — from
        # the run relevant to the stage: the inclusion run for include/excl_include, the exclusion
        # run for final/excl_final, and the page's furthest disposition on a deterministic stage.
        inc = run or evaluate.latest_inclusion_run(conn, cid)
        exc = evaluate.latest_exclusion_run(conn, inc) if inc else None
        inc_keep = _map(f"SELECT url, keep AS val FROM evaluation_results WHERE run_id = {P}", inc) if inc else {}
        exc_keep = _map(f"SELECT url, keep AS val FROM evaluation_results WHERE run_id = {P}", exc) if exc else {}
        def _verdict(km, u):
            if u not in km:
                return ""                        # page not evaluated at this stage
            k = km[u]
            return "Keep" if k == 1 else "Drop" if k == 0 else "Unscored"
        for r in rows:
            u = r.get("url")
            if stage in ("include", "excl_include"):
                r["decision"] = _verdict(inc_keep, u)
            elif stage in ("final", "excl_final"):
                r["decision"] = _verdict(exc_keep if exc else inc_keep, u)
            else:                                # deterministic stage → furthest AI disposition
                r["decision"] = _verdict(exc_keep, u) if u in exc_keep else _verdict(inc_keep, u)
    if "confidence" in need:
        # Confidence label from the inclusion score (pass 1 is the only phase that scores).
        inc = run or evaluate.latest_inclusion_run(conn, cid)
        sc = _map(f"SELECT url, score AS val FROM evaluation_results WHERE run_id = {P}", inc) if inc else {}
        for r in rows:
            s = sc.get(r.get("url"))
            lab = evaluate.confidence_label(s)
            r["confidence"] = f"{lab} ({float(s):.1f})" if lab and s is not None else ""
    if "funnel_stage" in need:
        # The furthest funnel stage each page actually reached — independent of the selected
        # view — read from the inclusion/exclusion keep decisions. Cumulative ladder:
        # Final shortlist ⊃ Reached exclusion ⊃ Reached inclusion ⊃ Keyword only.
        inc = run or evaluate.latest_inclusion_run(conn, cid)
        exc = evaluate.latest_exclusion_run(conn, inc) if inc else None
        inc_keep = _map(f"SELECT url, keep AS val FROM evaluation_results WHERE run_id = {P}", inc) if inc else {}
        exc_keep = _map(f"SELECT url, keep AS val FROM evaluation_results WHERE run_id = {P}", exc) if exc else {}
        for r in rows:
            u = r.get("url")
            if exc_keep.get(u) == 1:
                r["funnel_stage"] = "Final shortlist"          # kept by exclusion (Phase 2)
            elif inc_keep.get(u) == 1:
                r["funnel_stage"] = "Reached exclusion"        # inclusion kept it → entered Phase 2
            elif u in inc_keep:
                r["funnel_stage"] = "Reached LLM inclusion"    # scored by Phase 1, not kept
            else:
                r["funnel_stage"] = "Keyword search"           # got no further than keyword search


# ---- results table (browse the shortlist with configurable columns) -----
# (key, label, default_on). 'content' is intentionally excluded (too big to show).
_RESULTS_SECTIONS = [
    ("Content", [
        ("title", "Title", True),                        # rendered hyperlinked to URL
        ("effective_document_type", "Document type", True),   # html_publication -> parent type
        ("url", "URL", False),
    ]),
    ("Ownership", [
        ("primary_org", "Primary publishing organisation", True),
        ("organisations", "Organisations", False),
    ]),
    ("Freshness", [
        ("public_updated_at", "Last update", True),
        ("last_update_band", "Last update band", True),
        ("first_published_at", "First published at", False),
    ]),
    ("Quality attributes", [
        ("readability", "Readability age", True),
        ("gds_issues", "GDS issue count", True),
        ("gds_findings", "GDS issues text", False),
        ("size", "Size", False),
    ]),
]
# All fields the results endpoint fetches (url always, for the Title link).
_RESULTS_FIELDS = ["url"] + [k for _, fs in _RESULTS_SECTIONS for k, _, _ in fs if k != "url"]


@app.get("/categories/{cid}/results", response_class=HTMLResponse)
def results_table_page(request: Request, cid: int):
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    category["display_name"] = cat.display_name(category)
    resp = templates.TemplateResponse("results_table.html", ctx(
        conn, request, category=category, sections=_RESULTS_SECTIONS))
    conn.close()
    return resp


@app.get("/api/categories/{cid}/results-table")
def api_results_table(request: Request, cid: int, limit: int = 50, offset: int = 0, q: str = ""):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    limit = max(1, min(limit, 200))
    offset = max(0, offset)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return JSONResponse({"error": "not found"}, status_code=404)
    filters = _effective_filters(conn, category)
    q = (q or "").strip()
    extra = {}
    if q:
        # Title contains `q`, over the WHOLE shortlist (not just the page). The wildcards
        # live in the parameter value, so there is no literal '%' in the SQL text.
        extra = {"extra_where": f"LOWER(c.title) LIKE LOWER({shortlist._P})",
                 "extra_params": ["%" + q + "%"]}
    try:
        # Fetch the page of rows first — the important part. A single page is cheap even
        # when the whole-shortlist count is slow.
        _keys, esql, eparams = shortlist.export_query(_RESULTS_FIELDS, limit=limit, offset=offset,
                                                      **filters, **extra)
        rows = [dict(r) for r in conn.execute(esql, tuple(eparams)).fetchall()]
        # Attach the corpus keyword hits stored on the membership (single source of truth,
        # shared with the charts). One IN-list lookup for this page of rows — no N+1.
        page_urls = [r["url"] for r in rows if r.get("url")]
        if page_urls:
            ph = ",".join([shortlist._P] * len(page_urls))
            hits = {dict(h)["url"]: dict(h)["matched_keywords"] for h in conn.execute(
                f"SELECT url, matched_keywords FROM category_shortlist_pages "
                f"WHERE category_id = {shortlist._P} AND url IN ({ph})",
                tuple([cid] + page_urls)).fetchall()}
            for r in rows:
                r["corpus_keywords"] = search_augment._load_list(hits.get(r.get("url")))
    except Exception as e:
        conn.close()
        return JSONResponse({"error": f"{type(e).__name__}: {e}"}, status_code=500)
    try:
        # A title search bypasses the count cache (its key ignores extra_where).
        total = (shortlist.count(conn, **filters, **extra) if q else cached_count(conn, **filters))
    except Exception:
        total = None
    conn.close()
    return JSONResponse({"total": total, "rows": rows, "limit": limit, "offset": offset, "q": q,
                         "sql": _display_sql(esql, eparams)})


# ---- audit shortlist (browse/extract the pages at any funnel stage) -----
# Stage key -> label. Deterministic stages progressively narrow the corpus; the two
# LLM stages restrict to the pages a run kept.
# The four funnel tiers offered in the Shortlist Stage dropdown, each the SUPERSET of pages
# that reached that tier. Underlying keys reuse the _stage_query filters:
#   doctype          → org + document-type set (the pages keyword search runs on)
#   keyword          → keyword shortlist        (the pages fed to LLM inclusion)
#   reached_exclusion→ inclusion keeps          (the pages fed to exclusion)
#   final            → exclusion keeps          (the final shortlist)
# Other keys (dept/include/excl_include/excl_final/reached_inclusion) remain handled by
# _stage_query for backward-compatible ?stage= links but are no longer listed here.
_AUDIT_STAGES = [
    ("doctype", "Made it to Keyword Search (org + doc type)"),
    ("keyword", "Made it to LLM Inclusion (keyword shortlist)"),
    ("reached_exclusion", "Made it to Exclusion (inclusion keeps)"),
    ("final", "Made it to final shortlist (exclusion keeps)"),
]
# Every key _stage_query understands stays valid for direct ?stage= URLs, even the ones the
# dropdown no longer shows.
_AUDIT_STAGE_KEYS = [k for k, _ in _AUDIT_STAGES] + [
    "dept", "include", "reached_inclusion", "excl_include", "excl_final"]
_AUDIT_DEFAULT_FIELDS = [k for _, fs in _DOWNLOAD_SECTIONS for (k, l, d, dis, w) in fs if d]


def _stage_query(conn, category, stage, run=None):
    """Filters for export_rows/count at an audit stage, or None if an LLM stage has
    no run yet. Deterministic stages drop the later filters; LLM stages restrict the
    content rows to the URLs their run kept. `run` (an inclusion run id) scopes the LLM
    stages to that specific run-chain; None uses the category's latest/active run."""
    eff = _effective_filters(conn, category)
    if stage == "dept":
        return dict(organisations=eff["organisations"], document_types=(), keywords=(), match=eff["match"])
    if stage == "doctype":
        return dict(organisations=eff["organisations"], document_types=eff["document_types"],
                    keywords=(), match=eff["match"])
    if stage == "keyword":
        return dict(eff)
    inc = run or evaluate.latest_inclusion_run(conn, category["id"])
    if not inc:
        return None
    exc = evaluate.latest_exclusion_run(conn, inc)
    # Pick which run and which decision (kept vs dropped) the stage shows. The two excl_*
    # stages surface the pages a phase EXCLUDED (keep = 0) — the phase the page was dropped at.
    # Superset stages: everything that REACHED a phase, not just what it kept. "reached_inclusion"
    # is every page the inclusion run scored (keep 1/0/NULL); "reached_exclusion" is every page fed
    # into exclusion = the inclusion keeps (whether or not exclusion has decided them yet).
    if stage == "reached_inclusion":
        return dict(organisations=(), document_types=(), keywords=(), match="any",
                    extra_where=f"c.url IN (SELECT url FROM evaluation_results "
                                f"WHERE run_id = {shortlist._P})",
                    extra_params=[inc])
    if stage == "reached_exclusion":
        return dict(organisations=(), document_types=(), keywords=(), match="any",
                    extra_where=f"c.url IN (SELECT url FROM evaluation_results "
                                f"WHERE run_id = {shortlist._P} AND keep = 1)",
                    extra_params=[inc])
    if stage == "final":
        run_id, keep = (exc or inc), 1                # final shortlist = exclusion keeps
    elif stage == "excl_include":
        run_id, keep = inc, 0                         # dropped by the inclusion phase
    elif stage == "excl_final":
        if not exc:
            return None                               # no exclusion run yet → nothing dropped there
        run_id, keep = exc, 0                         # dropped by the exclusion phase
    else:                                             # "include"
        run_id, keep = inc, 1                         # inclusion keeps
    return dict(organisations=(), document_types=(), keywords=(), match="any",
                extra_where=f"c.url IN (SELECT url FROM evaluation_results "
                            f"WHERE run_id = {shortlist._P} AND keep = {int(keep)})",
                extra_params=[run_id])


def _merge_extra(filters, q):
    """Combine a stage filter's extra_where with an optional title search (q)."""
    f = dict(filters)
    clauses, params = [], []
    ew = f.pop("extra_where", None)
    if ew:
        clauses.append(ew)
        params.extend(f.pop("extra_params", []))
    else:
        f.pop("extra_params", None)
    if q:
        clauses.append(f"LOWER(c.title) LIKE LOWER({shortlist._P})")
        params.append("%" + q + "%")
    if clauses:
        f["extra_where"] = " AND ".join(clauses)
        f["extra_params"] = params
    return f


@app.get("/api/categories/{cid}/audit-shortlist")
def api_audit_shortlist(request: Request, cid: int, stage: str = "keyword",
                        limit: int = 50, offset: int = 0, q: str = "",
                        fields: List[str] = Query(default=[]), run: str = ""):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    limit = max(1, min(limit, 200))
    offset = max(0, offset)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return JSONResponse({"error": "not found"}, status_code=404)
    run_row = evaluate.get_run(conn, run) if run else None
    if run and (not run_row or str(run_row.get("category_id")) != str(cid)):
        run = ""
    if stage not in _AUDIT_STAGE_KEYS:
        stage = "keyword"
    sq = _stage_query(conn, category, stage, run=run or None)
    if sq is None:
        conn.close()
        return JSONResponse({"stage": stage, "no_data": True, "rows": [], "keys": [], "total": 0,
                             "message": "No AI evaluation run yet — run it on the Semantic Match tab first."})
    keys_wanted = [f for f in fields if f in _all_audit_fields()] or list(_AUDIT_DEFAULT_FIELDS)
    real_keys = [k for k in keys_wanted if k in shortlist.EXPORT_FIELDS]
    filters = _merge_extra(sq, (q or "").strip())
    try:
        _keys, esql, eparams = shortlist.export_query(real_keys, limit=limit, offset=offset, **filters)
        rows = [dict(r) for r in conn.execute(esql, tuple(eparams)).fetchall()]
        _enrich_audit_rows(conn, cid, rows, keys_wanted, stage=stage, run=run or None)   # kw + reasons + stage + decision
    except Exception as e:
        conn.close()
        return JSONResponse({"error": f"{type(e).__name__}: {e}"}, status_code=500)
    try:
        total = shortlist.count(conn, **filters)
    except Exception:
        total = None
    conn.close()
    # Return the full selection (real + virtual) as the column order for the display.
    return JSONResponse({"stage": stage, "keys": keys_wanted, "rows": rows, "total": total,
                         "limit": limit, "offset": offset, "q": (q or "").strip(),
                         "sql": _display_sql(esql, eparams)})


def _valid_runs(conn, cid: int, runs: List[str]) -> List[str]:
    """Keep only inclusion runs of THIS category, in the order given, without duplicates."""
    out = []
    for rid in runs:
        rid = (rid or "").strip()
        if not rid or rid in out:
            continue
        row = evaluate.get_run(conn, rid)
        if row and str(row.get("category_id")) == str(cid):
            out.append(rid)
    return out


@app.get("/categories/{cid}/download", response_class=HTMLResponse)
def download_page(request: Request, cid: int, stage: str = "keyword",
                  fields: List[str] = Query(default=[]), run: List[str] = Query(default=[])):
    """`run` may repeat: every run given is pre-selected in the (multi-select) Run picker."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    category["display_name"] = cat.display_name(category)
    runs_selected = _valid_runs(conn, cid, run)
    if stage not in _AUDIT_STAGE_KEYS:
        stage = "keyword"
    # Page count: summed over the selected runs (each run contributes its own rows to the file).
    total = 0
    for rid in (runs_selected or [""]):
        sq = _stage_query(conn, category, stage, run=rid or None)
        if sq is None:
            total = None
            break
        try:
            total += shortlist.count(conn, **_merge_extra(sq, ""))
        except Exception:
            total = None
            break
    # Columns come from the Audit shortlist selection (passed as ?fields=…); fall back
    # to the default set if the page is opened directly with none. URL and Run ID are
    # always included (and first), so every download keys back to the page and its run.
    preselect = [f for f in fields if f in _all_audit_fields()] or list(_AUDIT_DEFAULT_FIELDS)
    preselect = ["url", "run_id"] + [f for f in preselect if f not in ("url", "run_id")]
    options = _incl_run_options(conn, cid)
    names = {o["run_id"]: o["label"] for o in options}
    resp = templates.TemplateResponse("download.html", ctx(
        conn, request, category=category, total=total,
        stage=stage, stage_label=dict(_AUDIT_STAGES).get(stage, stage),
        runs_selected=runs_selected,
        run_name=", ".join(names.get(r, r) for r in runs_selected),
        runs=options,
        preselect=preselect, audit_sections=_DOWNLOAD_SECTIONS,
        default_filename=f"gov-uk-audit-shortlist-{_dl_stamp()}"))
    conn.close()
    return resp


@app.get("/categories/{cid}/export")
def export_category(request: Request, cid: int, format: str = "csv", stage: str = "keyword",
                    filename: str = "", fields: List[str] = Query(default=[]), run: List[str] = Query(default=[])):
    """Build the chosen-format, chosen-field export for an audit stage. Sync route ->
    runs in a threadpool so a large export doesn't block the event loop. `run` may repeat:
    each inclusion run of this category contributes its own block of rows (scoping the LLM
    stages + AI-decision columns to that run-chain) to the same file, told apart by the
    Run ID column, which is always included. No run = the category's active run."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    runs_selected = _valid_runs(conn, cid, run) or [""]
    if stage not in _AUDIT_STAGE_KEYS:
        stage = "keyword"
    wanted = [f for f in fields if f in _all_audit_fields()] or list(_AUDIT_DEFAULT_FIELDS)
    wanted = ["url", "run_id"] + [k for k in wanted if k not in ("url", "run_id")]   # always included
    real = [k for k in wanted if k in shortlist.EXPORT_FIELDS]
    rows, keys = [], None
    for rid in runs_selected:
        sq = _stage_query(conn, category, stage, run=rid or None)
        if sq is None:
            conn.close()
            return RedirectResponse(url=str(request.url_for("download_page", cid=cid)), status_code=303)
        keys, run_rows = shortlist.export_rows(conn, real, **_merge_extra(sq, ""))   # keys: real, url-first
        _enrich_audit_rows(conn, cid, run_rows, wanted, stage=stage, run=rid or None)  # kw + reasons + stage + decision
        rows.extend(run_rows)
    conn.close()
    # Columns: url, then Run ID, then the other real fields, then the selected virtual fields.
    keys = ["url", "run_id"] + [k for k in keys if k != "url"] + [k for k in wanted if k in _VIRTUAL_FIELDS and k not in keys and k != "run_id"]
    labels = [_field_label(k) for k in keys]
    name = _safe_filename(filename, f"gov-uk-audit-shortlist-{_dl_stamp()}")

    if format == "json":
        payload = [{k: r.get(k) for k in keys} for r in rows]
        return Response(json.dumps(payload, indent=2, default=str), media_type="application/json",
                        headers={"Content-Disposition": f'attachment; filename="{name}.json"'})
    if format == "xlsx":
        try:
            import openpyxl
        except Exception:
            return JSONResponse({"error": "Excel export needs openpyxl (pip install -r requirements.txt)."},
                                status_code=500)
        from openpyxl.styles import Alignment
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "shortlist"
        ws.append(labels)
        # exclusion_reason is baked as 'label\nreason' (see _enrich_audit_rows); wrap so the
        # exclusion_hit label header shows on the line above the reason in Excel.
        excl_col = keys.index("exclusion_reason") + 1 if "exclusion_reason" in keys else None
        for r in rows:
            ws.append([r.get(k) for k in keys])
            if excl_col and r.get("exclusion_reason"):
                ws.cell(row=ws.max_row, column=excl_col).alignment = Alignment(wrap_text=True, vertical="top")
        bio = io.BytesIO()
        wb.save(bio)
        return Response(bio.getvalue(),
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'})
    # default CSV
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(labels)
    for r in rows:
        w.writerow([r.get(k) for k in keys])   # exclusion_reason already carries its label header
    return Response(buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="{name}.csv"'})


# ---- Data Analysis (guc-0004c4) + the generalised review list (guc-0032) ------------------
ANALYSIS_CONTEXTS = {
    "phase1": "Inclusion Phase (Recall)",
}
_VERDICT_LABEL = {"keep": "Keep", "reject": "Reject", "all": "All"}
_VERDICT_KEEP = {"keep": 1, "reject": 0}


def _run_label(run_row: dict) -> str:
    if not run_row:
        return "—"
    dt = (run_row.get("started_at") or "")[:19].replace("T", " ")
    return f"{run_row.get('name') or run_row['run_id']} · {dt} · {run_row.get('actual_model') or run_row.get('model') or ''}".strip(" ·")


def _own_run(conn, cid: int, run_id: str) -> Optional[dict]:
    """The run row if it is an inclusion run of THIS category, else None."""
    if not run_id:
        return None
    row = evaluate.get_run(conn, run_id)
    if not row or str(row.get("category_id")) != str(cid) or row.get("source_run_id"):
        return None
    return row


def _phase1_compare(conn, cid: int, baseline: str, comparison: str) -> dict:
    """Phase-1 (inclusion) decisions of two runs over the pages BOTH evaluated."""
    base = {r["url"]: r for r in evaluate.run_results(conn, baseline)}
    comp = {r["url"]: r for r in evaluate.run_results(conn, comparison)}
    shared = sorted(set(base) & set(comp))

    def counts(m):
        acc = sum(1 for u in shared if m[u].get("keep") == 1)
        rej = sum(1 for u in shared if m[u].get("keep") == 0)
        return {"accepted": acc, "rejected": rej, "unscored": len(shared) - acc - rej, "total": len(shared)}

    diff = [u for u in shared if base[u].get("keep") != comp[u].get("keep")]
    a2r = sum(1 for u in diff if base[u].get("keep") == 1 and comp[u].get("keep") == 0)
    r2a = sum(1 for u in diff if base[u].get("keep") == 0 and comp[u].get("keep") == 1)
    return {
        "shared": len(shared), "only_baseline": len(set(base) - set(comp)), "only_comparison": len(set(comp) - set(base)),
        "baseline": {"run_id": baseline, "label": _run_label(evaluate.get_run(conn, baseline)), **counts(base)},
        "comparison": {"run_id": comparison, "label": _run_label(evaluate.get_run(conn, comparison)), **counts(comp)},
        "differences": {"total": len(diff), "accepted_to_rejected": a2r, "rejected_to_accepted": r2a,
                        "involving_unscored": len(diff) - a2r - r2a},
    }


def _band(s) -> str:
    s = s or 0
    return "0 (wrong sense)" if s <= 0 else "0.01–0.35 (in passing)" if s <= 0.35 \
        else "0.36–0.65 (moderate)" if s <= 0.65 else "0.66–1 (major)"

_BAND_ORDER = ["0 (wrong sense)", "0.01–0.35 (in passing)", "0.36–0.65 (moderate)", "0.66–1 (major)"]

_SCORE_BUCKETS = ["0.0", "0.1", "0.15–0.2", "0.25–0.35", "0.4–0.65", "0.7–1.0"]


def _score_bucket(s) -> str:
    s = s or 0
    return "0.0" if s <= 0 else "0.1" if s <= 0.12 else "0.15–0.2" if s <= 0.23 \
        else "0.25–0.35" if s <= 0.35 else "0.4–0.65" if s <= 0.65 else "0.7–1.0"


def _phase2_compare(conn, cid: int, baseline: str, comparison: str) -> dict:
    """Phase-2 (exclusion) decisions of two run-chains, and where the FINAL result diverges,
    stratified by Phase-1 confidence — the key to explaining run-to-run variance."""
    _, b_excl = _chain_of(conn, baseline)
    _, c_excl = _chain_of(conn, comparison)
    bi = {r["url"]: r for r in evaluate.run_results(conn, baseline)}
    ci = {r["url"]: r for r in evaluate.run_results(conn, comparison)}
    be = {r["url"]: r for r in evaluate.run_results(conn, b_excl["run_id"])} if b_excl else {}
    ce = {r["url"]: r for r in evaluate.run_results(conn, c_excl["run_id"])} if c_excl else {}
    fb = evaluate._final_keep_map(conn, baseline, b_excl and b_excl["run_id"])
    fc = evaluate._final_keep_map(conn, comparison, c_excl and c_excl["run_id"])
    shared = sorted(set(bi) & set(ci))
    excl_input = [u for u in shared if bi[u].get("keep") == 1 and ci[u].get("keep") == 1]

    def ecount(em):
        dropped = sum(1 for u in excl_input if em.get(u, {}).get("keep") == 0)
        return {"kept": len(excl_input) - dropped, "dropped": dropped}

    flips = [u for u in shared if fb.get(u) != fc.get(u)]
    excl_driven = sum(1 for u in flips if bi[u].get("keep") == ci[u].get("keep"))
    bands = {}
    hist = {b: {"pages": 0, "flipped": 0} for b in _SCORE_BUCKETS}
    for u in shared:
        sc = min(bi[u].get("score") or 0, ci[u].get("score") or 0)
        flip = 1 if fb.get(u) != fc.get(u) else 0
        d = bands.setdefault(_band(sc), {"pages": 0, "flipped": 0})
        d["pages"] += 1; d["flipped"] += flip
        h = hist[_score_bucket(sc)]; h["pages"] += 1; h["flipped"] += flip
    return {
        "shared": len(shared), "excl_input": len(excl_input),
        "baseline": {"label": _run_label(evaluate.get_run(conn, baseline)), **ecount(be)},
        "comparison": {"label": _run_label(evaluate.get_run(conn, comparison)), **ecount(ce)},
        "flips": {"total": len(flips), "exclusion_driven": excl_driven,
                  "inclusion_driven": len(flips) - excl_driven},
        "bands": [{"band": b, **bands[b]} for b in _BAND_ORDER if b in bands],
        "score_hist": [{"label": b, **hist[b]} for b in _SCORE_BUCKETS],
    }


def _overall_compare(conn, cid: int, baseline: str, comparison: str) -> dict:
    """Overall run-to-run reliability: the two FINAL shortlists compared, plus Phase-1 score
    stability (how repeatable the scoring itself is)."""
    _, b_excl = _chain_of(conn, baseline)
    _, c_excl = _chain_of(conn, comparison)
    bi = {r["url"]: r for r in evaluate.run_results(conn, baseline)}
    ci = {r["url"]: r for r in evaluate.run_results(conn, comparison)}
    fb = evaluate._final_keep_map(conn, baseline, b_excl and b_excl["run_id"])
    fc = evaluate._final_keep_map(conn, comparison, c_excl and c_excl["run_id"])
    shared = sorted(set(fb) & set(fc))
    both_in = sum(1 for u in shared if fb[u] and fc[u])
    both_out = sum(1 for u in shared if not fb[u] and not fc[u])
    base_only = sum(1 for u in shared if fb[u] and not fc[u])
    comp_only = sum(1 for u in shared if not fb[u] and fc[u])
    agree = both_in + both_out
    deltas = [abs((bi[u].get("score") or 0) - (ci[u].get("score") or 0))
              for u in shared if u in bi and u in ci]
    mean_delta = round(sum(deltas) / len(deltas), 3) if deltas else 0
    return {
        "shared": len(shared), "agree": agree,
        "agree_pct": round(100 * agree / len(shared), 1) if shared else 0,
        "sym_diff": len(shared) - agree,
        "confusion": {"both_in": both_in, "both_out": both_out, "base_only": base_only, "comp_only": comp_only},
        "score_mean_abs_delta": mean_delta,
        "baseline_label": _run_label(evaluate.get_run(conn, baseline)),
        "comparison_label": _run_label(evaluate.get_run(conn, comparison)),
    }


@app.get("/categories/{cid}/analysis", response_class=HTMLResponse)
def analysis_page(request: Request, cid: int):
    """Data Analysis sub-tab: cards explaining the difference between a baseline run and a
    comparison run. Pickers default to the two newest inclusion runs; ?baseline=&comparison= pre-select."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    category["display_name"] = cat.display_name(category)
    resp = templates.TemplateResponse("analysis.html", ctx(conn, request, category=category))
    conn.close()
    return resp


@app.get("/api/categories/{cid}/analysis/phase1")
def api_analysis_phase1(request: Request, cid: int, baseline: str = "", comparison: str = ""):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    try:
        if not (_own_run(conn, cid, baseline) and _own_run(conn, cid, comparison)):
            return JSONResponse({"error": "Pick two inclusion runs of this shortlist."}, status_code=400)
        return JSONResponse(_phase1_compare(conn, cid, baseline, comparison))
    finally:
        conn.close()


@app.get("/api/categories/{cid}/analysis/phase2")
def api_analysis_phase2(request: Request, cid: int, baseline: str = "", comparison: str = ""):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    try:
        if not (_own_run(conn, cid, baseline) and _own_run(conn, cid, comparison)):
            return JSONResponse({"error": "Pick two inclusion runs of this shortlist."}, status_code=400)
        return JSONResponse(_phase2_compare(conn, cid, baseline, comparison))
    finally:
        conn.close()


@app.get("/api/categories/{cid}/analysis/overall")
def api_analysis_overall(request: Request, cid: int, baseline: str = "", comparison: str = ""):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    try:
        if not (_own_run(conn, cid, baseline) and _own_run(conn, cid, comparison)):
            return JSONResponse({"error": "Pick two inclusion runs of this shortlist."}, status_code=400)
        return JSONResponse(_overall_compare(conn, cid, baseline, comparison))
    finally:
        conn.close()


_PROMPT_REVIEW_SYSTEM = (
    "You are a prompt engineer improving the REPEATABILITY of an LLM pipeline that shortlists "
    "GOV.UK pages in two phases: Phase 1 INCLUSION scores each page 0.0–1.0 and keeps any positive "
    "score; Phase 2 EXCLUSION re-checks the keeps and may only DROP. You are given the current "
    "INCLUDE / EXCLUDE criteria and KEEP / DROP examples the team edits (the base template around "
    "them is FIXED — do not rewrite it), a DATA DIAGNOSIS of how two identical-config runs diverged, "
    "and concrete FLIPPING pages (kept by one run, dropped by the other) with each run's reason. "
    "Explain what is driving the run-to-run variance and propose concrete, minimal, paste-ready edits "
    "that would make future runs repeatable. Generalise — never hard-code URLs or overfit single "
    "pages. Use the DIAGNOSIS to identify which phase is at fault: if the flips are in EXCLUSION on "
    "low-confidence pages, sharpen the EXCLUDE criteria with a single decisive, testable rule and add "
    "2–4 KEEP and DROP examples drawn from the flipping pages; only edit INCLUDE if the diagnosis "
    "blames Phase 1. If an edit changes what the shortlist MEANS (e.g. dropping incidental mentions), "
    "call it out as an explicit editorial choice for the user to decide. "
    "Reply in plain text, no code fences, in exactly this structure:\n"
    "DIAGNOSIS: 1–3 sentences on what is driving the variance.\n"
    "FAULT: which phase / prompt.\n"
    "EXCLUDE CRITERIA — suggested: full revised text ready to paste (or 'no change').\n"
    "KEEP EXAMPLES — suggested: (or 'no change').\n"
    "DROP EXAMPLES — suggested: (or 'no change').\n"
    "INCLUDE CRITERIA — suggested: (or 'no change').\n"
    "WHY: bullet points tying each edit to the diagnosis and the pages.\n"
    "EDITORIAL CHOICE: any decision the user must make (or 'none').\n"
    "VERIFY: one line on how to confirm it worked.")


@app.post("/api/categories/{cid}/analysis/prompt-review")
async def api_analysis_prompt_review(request: Request, cid: int):
    """Ask the model to critique the inclusion/exclusion prompts using the run-to-run divergence
    data + the flipping borderline pages, and propose paste-ready edits. Suggest-only."""
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    body = await request.json()
    base = str(body.get("baseline") or ""); comp = str(body.get("comparison") or "")
    if not (base and comp):
        return JSONResponse({"error": "bad request"}, status_code=400)
    conn = connect()
    try:
        if not (_own_run(conn, cid, base) and _own_run(conn, cid, comp)):
            return JSONResponse({"error": "Pick two inclusion runs of this shortlist."}, status_code=400)
        category = cat.get_category(conn, cid) or {}
        p1 = _phase1_compare(conn, cid, base, comp)
        p2 = _phase2_compare(conn, cid, base, comp)
        ov = _overall_compare(conn, cid, base, comp)
        _, b_excl = _chain_of(conn, base)
        _, o_excl = _chain_of(conn, comp)
        ib, io = _reasons_map(conn, base), _reasons_map(conn, comp)
        eb = _reasons_map(conn, b_excl and b_excl["run_id"])
        eo = _reasons_map(conn, o_excl and o_excl["run_id"])
        fb = evaluate._final_keep_map(conn, base, b_excl and b_excl["run_id"])
        fc = evaluate._final_keep_map(conn, comp, o_excl and o_excl["run_id"])
        shared = sorted(set(fb) & set(fc))
        flipped = [u for u in shared if fb[u] != fc[u]]
        flipped.sort(key=lambda u: min(ib.get(u, {}).get("score") or 0, io.get(u, {}).get("score") or 0))
        titles = {}
        head = flipped[:12]
        if head:
            ph = ",".join([shortlist._P] * len(head))
            for r in conn.execute(f"SELECT url, title FROM content WHERE url IN ({ph})", tuple(head)).fetchall():
                d = dict(r); titles[d["url"]] = d.get("title") or ""
        kd = lambda k: "KEEP" if k == 1 else ("DROP" if k == 0 else "—")
        cases = []
        for u in head:
            sb, so, xb, xo = ib.get(u, {}), io.get(u, {}), eb.get(u, {}), eo.get(u, {})
            cases.append(
                f"- {(titles.get(u, '') or u)[:80]}\n"
                f"  Phase-1 score: Run1 {sb.get('score')} / Run2 {so.get('score')}\n"
                f"  Exclusion: Run1 {kd(xb.get('keep'))} — {xb.get('reason') or '(none)'}\n"
                f"             Run2 {kd(xo.get('keep'))} — {xo.get('reason') or '(none)'}")
        diag = (
            f"Pages both runs evaluated: {ov['shared']}. Final agreement {ov['agree_pct']}%, "
            f"{ov['sym_diff']} pages end differently.\n"
            f"Phase 1 (inclusion): decisions differ on {p1['differences']['total']} pages; the scoring "
            f"itself is stable (mean per-page score change {ov['score_mean_abs_delta']}).\n"
            f"Phase 2 (exclusion): of {p2['excl_input']} pages both runs kept at Phase 1, "
            f"{p2['flips']['total']} end differently and {p2['flips']['exclusion_driven']} of those are "
            f"exclusion-driven.\nFlip rate by Phase-1 confidence band:\n"
            + "\n".join(f"  {b['band']}: {b['flipped']}/{b['pages']} flipped" for b in p2['bands']))
        prompt = (
            f"TOPIC: {category.get('display_name') or ''}\n\n"
            f"CURRENT INCLUDE CRITERIA:\n{category.get('inclusion_context') or '(none)'}\n\n"
            f"CURRENT EXCLUDE CRITERIA:\n{category.get('exclusion_context') or '(none)'}\n\n"
            f"CURRENT KEEP EXAMPLES:\n{category.get('adjudication_hints_keep') or '(none)'}\n\n"
            f"CURRENT DROP EXAMPLES:\n{category.get('adjudication_hints_drop') or '(none)'}\n\n"
            f"DATA DIAGNOSIS (two identical-config runs):\n{diag}\n\n"
            f"FLIPPING BORDERLINE PAGES ({len(flipped)} total; lowest-confidence shown):\n"
            f"{chr(10).join(cases) or '(no final-outcome flips between these runs)'}\n\n"
            "Explain the variance and propose the edits.")
        cfg = _ai_config_for_phase(conn, "exclusion")
        budget, spent = _budget(conn), _daily_spend(conn)
        if budget > 0 and spent >= budget:
            return JSONResponse({"error": f"Daily AI budget of ${budget:.2f} reached "
                                 f"(${spent:.4f} spent today)."}, status_code=402)
        res = await run_in_threadpool(_ai_chat, cfg, _PROMPT_REVIEW_SYSTEM,
                                      [{"role": "user", "content": prompt}], 1800)
        if res.get("error") or res.get("fatal"):
            return JSONResponse({"error": res.get("error") or "AI error"}, status_code=502)
        _log_ai_usage(conn, res.get("cost_usd") or 0.0, res.get("input_tokens"),
                      res.get("output_tokens"), "prompt-review")
        return JSONResponse({"review": (res.get("reply") or "").strip(), "model": res.get("actual_model"),
                             "cost": res.get("cost_usd"), "flips": len(flipped)})
    finally:
        conn.close()


@app.get("/categories/{cid}/analysis/list", response_class=HTMLResponse)
def analysis_list_page(request: Request, cid: int, context: str = "phase1", run: str = "", verdict: str = "all",
                       shared_with: str = "", baseline: str = "", comparison: str = ""):
    """Generalised review list, opened from the analysis cards (and, later, from elsewhere).
    Title = the context and the verdict, e.g. "Inclusion Phase (Recall) – Keep".
      context=phase1      one run's Phase-1 decisions; verdict=keep|reject|all; shared_with=<run>
                          restricts to pages that run also evaluated (what the cards count).
      context=phase1_diff the pages two runs decided differently (baseline=&comparison=)."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    category["display_name"] = cat.display_name(category)
    rows, runs_shown, note = [], [], ""
    if context == "phase1_diff":
        b, c = _own_run(conn, cid, baseline), _own_run(conn, cid, comparison)
        title = f"{ANALYSIS_CONTEXTS['phase1']} – Differences"
        if b and c:
            base = {r["url"]: r for r in evaluate.run_results(conn, baseline)}
            comp = {r["url"]: r for r in evaluate.run_results(conn, comparison)}
            for u in sorted(set(base) & set(comp)):
                if base[u].get("keep") != comp[u].get("keep"):
                    rows.append({"url": u, "decisions": [base[u], comp[u]],
                                 "reasons": [base[u].get("reason") or "", comp[u].get("reason") or ""]})
            runs_shown = [{"role": "Baseline", "label": _run_label(b)}, {"role": "Comparison", "label": _run_label(c)}]
            note = "pages both runs evaluated but decided differently"
    else:
        verdict = verdict if verdict in _VERDICT_LABEL else "all"
        title = f"{ANALYSIS_CONTEXTS.get(context, ANALYSIS_CONTEXTS['phase1'])} – {_VERDICT_LABEL[verdict]}"
        r = _own_run(conn, cid, run)
        if r:
            results = evaluate.run_results(conn, run, keep=_VERDICT_KEEP.get(verdict))
            other = _own_run(conn, cid, shared_with)
            if other:
                shared = {x["url"] for x in evaluate.run_results(conn, shared_with)}
                results = [x for x in results if x["url"] in shared]
                note = f"only pages also evaluated by {_run_label(other)}"
            rows = [{"url": x["url"], "decisions": [x], "reasons": [x.get("reason") or ""]} for x in results]
            runs_shown = [{"role": "Run", "label": _run_label(r)}]
        baseline, comparison = (run, shared_with)
    # Titles from the corpus, in batches.
    urls = [x["url"] for x in rows]
    titles = {}
    for i in range(0, len(urls), 500):
        chunk = urls[i:i + 500]
        ph = ",".join([shortlist._P] * len(chunk))
        for t in conn.execute(f"SELECT url, title FROM content WHERE url IN ({ph})", tuple(chunk)).fetchall():
            titles[t["url"]] = t["title"]
    for x in rows:
        x["title"] = titles.get(x["url"]) or ""
    resp = templates.TemplateResponse("analysis_list.html", ctx(
        conn, request, category=category, title=title, rows=rows, runs_shown=runs_shown, note=note,
        baseline=baseline, comparison=comparison))
    conn.close()
    return resp


# ---- funnel audit (per-page log) ----------------------------------------
def _build_audit(cid: int, filters: dict) -> dict:
    conn = connect()
    try:
        return audit.build_audit(conn, cid, filters["organisations"],
                                 filters["document_types"], filters["keywords"])
    finally:
        conn.close()


@app.post("/api/categories/{cid}/audit")
async def api_build_audit(request: Request, cid: int):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return JSONResponse({"error": "not found"}, status_code=404)
    filters = _effective_filters(conn, category)
    conn.close()
    if not filters["organisations"]:
        return JSONResponse({"error": "Set at least one organisation — the audit starts "
                             "at the organisation filter."}, status_code=400)
    counters = await run_in_threadpool(_build_audit, cid, filters)
    summary = [{"outcome": o, "count": n} for o, n in
               [(audit.OUTCOME_INCLUDED, counters[audit.OUTCOME_INCLUDED]),
                (audit.OUTCOME_DOCTYPE, counters[audit.OUTCOME_DOCTYPE]),
                (audit.OUTCOME_KEYWORD, counters[audit.OUTCOME_KEYWORD])]]
    return JSONResponse({"total": counters["total"], "summary": summary})


@app.get("/categories/{cid}/audit/download")
def download_audit(request: Request, cid: int, outcome: str = ""):
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    category = cat.get_category(conn, cid)
    if not category:
        conn.close()
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    # Built on demand (there is no separate build step) so the file always reflects the
    # current definition. Requires an organisation — the audit starts at the org filter.
    filters = _effective_filters(conn, category)
    if filters["organisations"]:
        audit.build_audit(conn, cid, filters["organisations"],
                          filters["document_types"], filters["keywords"])
    rows = audit.audit_rows(conn, cid, outcome=outcome or None)
    conn.close()
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["url", "outcome"])
    for r in rows:
        w.writerow([r["url"], r["outcome"]])
    tag = (outcome or "all").replace(":", "").replace(" ", "-")
    return Response(buf.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="audit-{cid}-{tag}.csv"'})


# ---- category transfer (download here, import into a test env) -----------
@app.get("/categories/{cid}/download-bundle")
def download_category_bundle(request: Request, cid: int):
    """Download a category + everything needed to reproduce it (definition, runs,
    results, membership, audit, and metadata-only content/orgs) as a gzipped JSON
    bundle. Import it into a test environment with `category_transfer import`."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    try:
        bundle = category_transfer.export_bundle(conn, cid)
    finally:
        conn.close()
    if bundle is None:
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    slug = (bundle["meta"].get("slug") or f"cat-{cid}")
    name = f"shortlist-{slug}-{_dl_stamp()}.json.gz"
    body = gzip.compress(category_transfer.dumps(bundle).encode("utf-8"))
    return Response(body, media_type="application/gzip",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.post("/admin/import-category")
async def admin_import_category(request: Request):
    """Admin-only: upload a category bundle and load it into THIS environment, reassigning
    the category to the importing user. Intended for a test environment, to reproduce and
    fix a live category."""
    if not authed(request):
        return login_redirect(request)
    users_url = str(request.url_for("admin_users_page"))
    if not _require_admin(request):
        return RedirectResponse(url=users_url, status_code=303)
    form = await request.form()
    upload = form.get("bundle")
    if upload is None or not hasattr(upload, "read"):
        return RedirectResponse(url=users_url + "?import_error=No+file+chosen", status_code=303)
    raw = await upload.read()
    try:
        if raw[:2] == b"\x1f\x8b":            # gzip magic
            raw = gzip.decompress(raw)
        bundle = category_transfer.loads(raw.decode("utf-8"))
    except Exception:
        return RedirectResponse(url=users_url + "?import_error=Could+not+read+the+bundle", status_code=303)
    cu = current_user(request)
    owner = cu.get("email") if cu else None
    conn = connect()
    try:
        summary = category_transfer.import_bundle(conn, bundle, owner_email=owner)
    except Exception as e:
        conn.close()
        return RedirectResponse(url=users_url + f"?import_error={quote(f'{type(e).__name__}: {e}')}",
                                status_code=303)
    conn.close()
    msg = f"Imported shortlist {summary['slug']} ({summary['content']} pages, {summary['runs']} runs)."
    return RedirectResponse(url=users_url + f"?import_ok={quote(msg)}", status_code=303)


@app.get("/assistant", response_class=HTMLResponse)
def assistant_page(request: Request):
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    cfg = _ai_config(conn)
    resp = templates.TemplateResponse("assistant.html", ctx(
        conn, request, active_nav="assistant", model=cfg["model"],
        provider_label=cfg["label"], has_key=cfg["has_key"],
        default_system=prompts.default_text("assistant"),
        base_url=cfg["base_url"] or "api.anthropic.com (Claude default)"))
    conn.close()
    return resp


@app.post("/api/assistant")
async def api_assistant(request: Request):
    if not authed(request):
        return JSONResponse({"error": "auth"}, status_code=401)
    form = await request.form()
    prompt = (form.get("prompt") or "").strip()
    system = (form.get("system") or "").strip()
    if not prompt:
        return JSONResponse({"error": "Enter a prompt."}, status_code=400)
    finding = guardrails.check(prompt)          # block personal data / prohibited language
    if finding:
        return JSONResponse({"reply": guardrails.refusal_message(finding)})
    conn = connect()
    cfg = _ai_config(conn)
    budget = _budget(conn)
    spent = _daily_spend(conn)
    conn.close()
    if budget > 0 and spent >= budget:
        return JSONResponse({"error": f"Daily AI budget of ${budget:.2f} reached "
                             f"(${spent:.4f} spent today). Raise it in Settings."})
    # Run the blocking SDK call off the event loop so one slow request can't stall the app.
    result = await run_in_threadpool(_ai_reply, cfg, system, prompt)
    if not result.get("error"):
        conn = connect()
        _log_ai_usage(conn, result.get("cost_usd"), result.get("input_tokens"),
                      result.get("output_tokens"), "assistant")
        conn.close()
    return JSONResponse(result)


@app.get("/profile", response_class=HTMLResponse)
def profile_page(request: Request, details_ok: int = 0, pw_ok: int = 0,
                 details_error: str = "", pw_error: str = ""):
    """The user's profile: their details + password (accounts mode) and the per-viewer
    UI display level. The display level is stored client-side (localStorage) and applied as
    a data-ui-level attribute; it needs no account. Higher levels reveal more detail."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    levels = [
        ("simple", "Simple",
         "Build shortlists, run the AI phase, and review and download the final results."),
        ("advanced", "Advanced",
         "Adds detailed data on the selection — almost log level."),
        ("expert", "Expert",
         "See everything except system configuration."),
        ("admin", "Admin",
         "See everything."),
    ]
    resp = templates.TemplateResponse("profile.html", ctx(
        conn, request, levels=levels, accounts_mode=(AUTH_MODE == "accounts"),
        me=current_user(request), details_ok=details_ok, pw_ok=pw_ok,
        details_error=details_error, pw_error=pw_error))
    conn.close()
    return resp


@app.post("/profile/details")
async def update_profile_details(request: Request):
    """Update the signed-in user's own first/last name (accounts mode). Email is the login
    identity and is not editable here."""
    if not authed(request):
        return login_redirect(request)
    profile_url = str(request.url_for("profile_page"))
    u = current_user(request)
    if not u:
        return RedirectResponse(url=profile_url, status_code=303)
    form = await request.form()
    first = (form.get("first_name") or "").strip()
    last = (form.get("last_name") or "").strip()
    conn = connect()
    try:
        if not first or not last:
            return RedirectResponse(url=profile_url + "?details_error=First+and+last+name+are+required.#details",
                                    status_code=303)
        accounts.update_user(conn, u["id"], first_name=first, last_name=last)
    finally:
        conn.close()
    return RedirectResponse(url=profile_url + "?details_ok=1#details", status_code=303)


@app.post("/profile/password")
async def change_own_password(request: Request):
    """Change the signed-in user's own password: requires the current password (accounts mode)."""
    if not authed(request):
        return login_redirect(request)
    profile_url = str(request.url_for("profile_page"))
    u = current_user(request)
    if not u:
        return RedirectResponse(url=profile_url, status_code=303)
    form = await request.form()
    old = form.get("current_password") or ""
    new = form.get("new_password") or ""
    confirm = form.get("confirm_password") or ""
    conn = connect()
    try:
        if new != confirm:
            msg = "The new passwords do not match."
            return RedirectResponse(url=profile_url + f"?pw_error={quote(msg)}#password", status_code=303)
        try:
            accounts.change_password(conn, u["id"], old, new)
        except ValueError as exc:
            return RedirectResponse(url=profile_url + f"?pw_error={quote(str(exc))}#password", status_code=303)
    finally:
        conn.close()
    return RedirectResponse(url=profile_url + "?pw_ok=1#password", status_code=303)


def _require_admin(request: Request):
    """The current user if they are an admin (accounts mode only), else None."""
    if AUTH_MODE != "accounts":
        return None
    u = current_user(request)
    return u if (u and roles.allows(u.get("role"), "admin")) else None


@app.get("/admin/users", response_class=HTMLResponse)
def admin_users_page(request: Request, user_ok: int = 0, user_error: str = "",
                     import_ok: str = "", import_error: str = ""):
    """Admin: view every user and open each one to edit their profile (guc-0020).
    Accounts mode + admin only."""
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    if not _require_admin(request):
        conn.close()
        return RedirectResponse(url=str(request.url_for("list_categories_page")), status_code=303)
    usql, uparams = accounts.list_users_query()
    resp = templates.TemplateResponse("admin_users.html", ctx(
        conn, request, active_nav="users", users=accounts.list_users(conn),
        account_roles=accounts.ROLES, user_ok=user_ok, user_error=user_error,
        import_ok=import_ok, import_error=import_error,
        list_sql=_display_sql(usql, uparams)))
    conn.close()
    return resp


@app.post("/admin/users")
async def admin_create_user(request: Request):
    """Admin creates an account (accounts mode). Same fields as registration, plus role."""
    if not authed(request):
        return login_redirect(request)
    users_url = str(request.url_for("admin_users_page"))
    if not _require_admin(request):
        return RedirectResponse(url=users_url, status_code=303)
    form = await request.form()
    role = form.get("role") if form.get("role") in accounts.ROLES else "User"
    conn = connect()
    try:
        accounts.create_user(conn, email=(form.get("email") or ""),
                             first_name=(form.get("first_name") or ""),
                             last_name=(form.get("last_name") or ""),
                             password=(form.get("password") or ""), role=role)
    except accounts.EmailTakenError:
        return RedirectResponse(url=users_url + "?user_error=exists", status_code=303)
    except ValueError as e:
        return RedirectResponse(url=users_url + "?user_error=" + quote(str(e)), status_code=303)
    finally:
        conn.close()
    return RedirectResponse(url=users_url + "?user_ok=1", status_code=303)


@app.get("/admin/users/{user_id}/edit", response_class=HTMLResponse)
def admin_edit_user_page(request: Request, user_id: str, pw_msg: str = "", reset_link: str = ""):
    if not authed(request):
        return login_redirect(request)
    if not _require_admin(request):
        return RedirectResponse(url=str(request.url_for("admin_users_page")), status_code=303)
    conn = connect()
    try:
        u = accounts.get_user(conn, user_id)
        if not u:
            return RedirectResponse(url=str(request.url_for("admin_users_page")), status_code=303)
        return templates.TemplateResponse("user_edit.html", ctx(
            conn, request, u=u, account_roles=accounts.ROLES, account_statuses=accounts.STATUSES,
            pw_msg=pw_msg, reset_link=reset_link))
    finally:
        conn.close()


@app.post("/admin/users/{user_id}/edit")
async def admin_update_user(request: Request, user_id: str):
    if not authed(request):
        return login_redirect(request)
    if not _require_admin(request):
        return RedirectResponse(url=str(request.url_for("admin_users_page")), status_code=303)
    form = await request.form()
    role = form.get("role") if form.get("role") in accounts.ROLES else None
    status = form.get("account_status") if form.get("account_status") in accounts.STATUSES else None
    conn = connect()
    try:
        accounts.update_user(conn, user_id, first_name=form.get("first_name"),
                             last_name=form.get("last_name"), role=role, account_status=status)
    finally:
        conn.close()
    return RedirectResponse(url=str(request.url_for("admin_users_page")) + "?user_ok=1", status_code=303)


@app.post("/admin/users/{user_id}/require-reset")
async def admin_require_reset(request: Request, user_id: str):
    """Admin 'dirties' the record: the user must choose a new password at next login."""
    if not authed(request):
        return login_redirect(request)
    if not _require_admin(request):
        return RedirectResponse(url=str(request.url_for("admin_users_page")), status_code=303)
    conn = connect()
    try:
        accounts.require_password_change(conn, user_id, on=True)
        actor = current_user(request)
        accounts.audit(conn, "password_change_required", user_id=user_id,
                       actor_id=(actor or {}).get("id"), ip=_client_ip(request))
    finally:
        conn.close()
    edit = str(request.url_for("admin_edit_user_page", user_id=user_id))
    return RedirectResponse(url=edit + "?pw_msg=" + quote("This user must set a new password at next sign-in."),
                            status_code=303)


@app.post("/admin/users/{user_id}/reset-link")
async def admin_reset_link(request: Request, user_id: str):
    """Admin generates a one-time reset link to hand to the user (until email is wired up)."""
    if not authed(request):
        return login_redirect(request)
    if not _require_admin(request):
        return RedirectResponse(url=str(request.url_for("admin_users_page")), status_code=303)
    conn = connect()
    try:
        token = accounts.create_reset_token(conn, user_id)
        actor = current_user(request)
        accounts.audit(conn, "password_reset_link_created", user_id=user_id,
                       actor_id=(actor or {}).get("id"), ip=_client_ip(request))
    finally:
        conn.close()
    link = str(request.url_for("reset_password_page", token=token))
    edit = str(request.url_for("admin_edit_user_page", user_id=user_id))
    return RedirectResponse(url=edit + "?reset_link=" + quote(link), status_code=303)


_PROMPT_NOTES = {
    "inclusion": "Sent once per page in an inclusion run — keep or drop, with a score and reason.",
    "exclusion": "Recall-priority re-check of the pages Phase 1 kept — only ever turns a keep into a drop.",
    "builder": "System prompt for the assistant that helps define a shortlist. When editing an existing "
               "shortlist a summary of its current fields is appended automatically.",
    "assistant": "Default system prompt pre-filled in the AI Assistant box (editable there per request).",
}


def _ai_prompts() -> list:
    """Read-only view for Settings → AI Prompts: each prompt's git default and the commit it
    comes from. Prompts live in git — the single source of truth — so there is no in-app
    editing and no saved-version history."""
    v = prompts.template_version()
    return [{
        "name": name,
        "title": prompts.LABELS[name],
        "note": _PROMPT_NOTES.get(name, ""),
        "placeholders": prompts.PLACEHOLDERS[name],
        "text": prompts.default_text(name),
        "version": v,
    } for name in prompts.NAMES]


def _settings_admin_ok(request: Request) -> bool:
    """Admin gate for settings mutations: admin in accounts mode; the single user in shared mode."""
    return AUTH_MODE != "accounts" or _require_admin(request) is not None


@app.get("/settings", response_class=HTMLResponse)
def settings_page(request: Request, saved: int = 0, user_ok: int = 0, user_error: str = ""):
    if not authed(request):
        return login_redirect(request)
    conn = connect()
    ai_models.seed_defaults(conn)
    models = ai_models.list_models(conn)
    active = settings.get_setting(conn, "active_model_id", "")
    if not active and models:
        active = str(models[0]["id"])
    for m in models:                      # is this model's provider usable?
        m["has_key"] = _provider_configured(m["provider"])
        m["active"] = str(m["id"]) == str(active)
    phase_models = [
        {"phase": ph, "key": PHASE_MODEL_KEYS[ph],
         "current": settings.get_setting(conn, PHASE_MODEL_KEYS[ph], ""),
         "mode_key": PHASE_MODE_KEYS[ph], "mode": _phase_mode(conn, ph),
         "provider": _ai_config_for_phase(conn, ph).get("provider")}
        for ph in (evaluate.PHASE_INCLUSION, evaluate.PHASE_EXCLUSION, evaluate.PHASE_ADJUDICATION)]
    accounts_mode = AUTH_MODE == "accounts"
    users = accounts.list_users(conn) if accounts_mode else None
    _msql, _mparams = ai_models.list_models_query()
    resp = templates.TemplateResponse("settings.html", ctx(
        conn, request, active_nav="settings", models=models,
        models_sql=_display_sql(_msql, _mparams),
        accounts_mode=accounts_mode, users=users, account_roles=accounts.ROLES,
        user_ok=user_ok, user_error=user_error,
        providers=list(PROVIDERS.keys()), phase_models=phase_models,
        provider_keys={k: _provider_configured(k) for k in PROVIDERS},
        daily_budget=_budget(conn), max_docs=_max_docs(conn),
        spent_today=round(_daily_spend(conn), 4), saved=saved,
        prompts=_ai_prompts()))
    conn.close()
    return resp


def _price_form(form, prefix: str) -> dict:
    """Pull the six tiered price fields (prefix + field name) out of a form."""
    out = {}
    for f in ai_models.PRICE_FIELDS:
        v = form.get(prefix + f)
        if v not in (None, ""):
            try:
                out[f] = float(v)
            except (TypeError, ValueError):
                pass
    return out


@app.post("/settings")
async def save_settings(request: Request):
    if not authed(request):
        return login_redirect(request)
    form = await request.form()
    conn = connect()
    if "active_role" in form:
        roles.set_role(conn, form.get("active_role"))
    active = form.get("active_model_id")
    if active:
        settings.set_setting(conn, "active_model_id", active)
    # Per-phase model selection (Inclusion / Exclusion / Adjudication).
    for ph, key in PHASE_MODEL_KEYS.items():
        if key in form:
            settings.set_setting(conn, key, (form.get(key) or "").strip())
    # Per-phase execution mode (synchronous | batch).
    for ph, key in PHASE_MODE_KEYS.items():
        if key in form:
            settings.set_setting(conn, key, evaluate.normalise_mode(form.get(key)))
    # Save edits to existing model rows.
    for m in ai_models.list_models(conn):
        rid = m["id"]
        prov = (form.get(f"m_{rid}_provider") or "").strip()
        mod = (form.get(f"m_{rid}_model_id") or "").strip()
        if prov in PROVIDERS and mod:
            try:
                ai_models.update_model(conn, rid, prov, mod,
                                       prices=_price_form(form, f"m_{rid}_"))
            except ValueError:
                pass
    # Guarded so saving another Settings tab (which doesn't post these) can't reset them.
    if "ai_daily_budget" in form:
        try:
            settings.set_setting(conn, "ai_daily_budget", str(float(form.get("ai_daily_budget") or DEFAULT_DAILY_BUDGET)))
        except ValueError:
            pass
    if "ai_max_docs_per_run" in form:
        try:
            settings.set_setting(conn, "ai_max_docs_per_run", str(int(float(form.get("ai_max_docs_per_run") or DEFAULT_MAX_DOCS))))
        except ValueError:
            pass
    conn.close()
    return RedirectResponse(url=str(request.url_for("settings_page")) + "?saved=1", status_code=303)


@app.post("/settings/models/add")
async def add_model_route(request: Request):
    if not authed(request):
        return login_redirect(request)
    form = await request.form()
    provider = (form.get("provider") or "").strip()
    model_id = (form.get("model_id") or "").strip()
    conn = connect()
    if provider in PROVIDERS and model_id:
        try:
            ai_models.add_model(conn, provider, model_id,
                                prices=_price_form(form, "add_"))
        except ValueError:
            pass
    conn.close()
    return RedirectResponse(url=str(request.url_for("settings_page")), status_code=303)


@app.post("/settings/models/delete")
async def delete_model_route(request: Request):
    if not authed(request):
        return login_redirect(request)
    form = await request.form()
    mid = form.get("model_id")
    conn = connect()
    if mid:
        try:
            ai_models.delete_model(conn, int(mid))
        except (TypeError, ValueError):
            pass
    conn.close()
    return RedirectResponse(url=str(request.url_for("settings_page")), status_code=303)


@app.post("/api/models/{mid}/test")
async def test_model_route(request: Request, mid: int):
    """Ping a model with a tiny prompt to confirm the supplier accepts the name."""
    if not authed(request):
        return JSONResponse({"ok": False, "error": "Not signed in."}, status_code=401)
    conn = connect()
    row = ai_models.get_model(conn, mid)
    if not row:
        conn.close()
        return JSONResponse({"ok": False, "error": "Model not found."}, status_code=404)
    cfg = _cfg_for(conn, row["provider"], row["model_id"])
    conn.close()
    if not _provider_configured(row["provider"]):   # API key, or AWS creds for Bedrock
        creds = ("AWS credentials" if row["provider"] == "bedrock" else "API key")
        return JSONResponse({"ok": False,
                             "error": f"No {creds} set for {cfg['label']}."})
    res = await run_in_threadpool(_ai_reply, cfg, "", "Reply with the single word: ok")
    if res.get("error"):
        return JSONResponse({"ok": False, "error": res["error"]})
    return JSONResponse({"ok": True, "actual_model": res.get("actual_model") or cfg["model"]})


@app.get("/settings/peak/{provider}", response_class=HTMLResponse)
def peak_schedule_page(request: Request, provider: str, saved: int = 0):
    if not authed(request):
        return login_redirect(request)
    if provider not in PROVIDERS:
        return RedirectResponse(url=str(request.url_for("settings_page")), status_code=303)
    conn = connect()
    resp = templates.TemplateResponse("peak_schedule.html", ctx(
        conn, request, active_nav="settings", provider=provider,
        provider_label=PROVIDERS[provider]["label"],
        days=peak_schedule.DAYS, hours=list(range(peak_schedule.HOURS)),
        grid=peak_schedule.get_grid(conn, provider),
        peak_count=peak_schedule.peak_hour_count(conn, provider), saved=saved))
    conn.close()
    return resp


@app.post("/settings/peak/{provider}")
async def save_peak_schedule(request: Request, provider: str):
    if not authed(request):
        return login_redirect(request)
    if provider not in PROVIDERS:
        return RedirectResponse(url=str(request.url_for("settings_page")), status_code=303)
    form = await request.form()
    grid = [[1 if form.get(f"h_{d}_{h}") else 0 for h in range(peak_schedule.HOURS)]
            for d in range(len(peak_schedule.DAYS))]
    conn = connect()
    peak_schedule.set_grid(conn, provider, grid)
    conn.close()
    return RedirectResponse(
        url=str(request.url_for("peak_schedule_page", provider=provider)) + "?saved=1",
        status_code=303)


@app.on_event("startup")
def _startup():
    conn = connect()
    db.init_db(conn)
    conn.close()
    # Apply auth-schema migrations too (accounts mode). Best-effort: on setups where the app role
    # doesn't own the `auth` schema (postgres created it), this is a no-op and the DDL is applied
    # out of band — see docs. Never let it block startup.
    if AUTH_MODE == "accounts":
        try:
            c = connect()
            try:
                accounts.init_auth_schema(c)
            finally:
                c.close()
        except Exception as e:
            logging.getLogger("assistant").warning("auth schema init skipped: %s", e)
    # Re-drive any evaluation run left mid-flight by the previous process (deploy/crash),
    # so a restart doesn't silently strand a run 'in progress'. Budget-capped; opt out via
    # the 'auto_resume_runs' setting.
    _resume_stalled_runs()
